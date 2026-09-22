"""FastAPI uç noktaları için entegrasyon testleri (TestClient + geçici SQLite).

NOT: Bu dosya fastapi + sqlalchemy gerektirir. Bu depo bu testleri yazan
oturumun kum havuzunda (dış paket erişimi organizasyon politikasıyla kapalı
olduğu için) çalıştırılamadı — kullanıcının kendi ortamında ya da CI'da
(bkz. .github/workflows/tests.yml) çalışır. Testler, `tests/conftest.py`
içinde ayarlanan geçici bir SQLite veritabanı ve sabit test secret'larıyla
tamamen izole çalışacak şekilde tasarlandı.
"""
import os
from datetime import datetime, timedelta

import pytest
from fastapi.testclient import TestClient

from backend import lisans
from backend import main as pts_main
from backend.main import app


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


@pytest.fixture(scope="module")
def admin_token(client):
    """İlk yönetici hesabını bir kez oluşturur ve token döner (modül genelinde paylaşılır)."""
    r = client.post("/auth/ilk-yonetici", json={"kullanici_adi": "admin", "parola": "GucluParola123!"})
    assert r.status_code == 200, r.text
    r2 = client.post("/auth/giris", json={"kullanici_adi": "admin", "parola": "GucluParola123!"})
    assert r2.status_code == 200, r2.text
    return r2.json()["token"]


@pytest.fixture
def yetkili_header(admin_token):
    return {"Authorization": f"Bearer {admin_token}"}


def _rol_ile_kullanici_olustur_ve_giris_yap(client, yetkili_header, kullanici_adi: str, rol: str, kisi_id: int = None) -> dict:
    """Verilen role sahip yeni bir kullanıcı oluşturur (yönetici gerektirir),
    onunla giriş yapar ve Authorization header'ını döner. RBAC testlerinde
    kullanılır. `kisi_id`, yalnızca rol="sakin" iken anlamlıdır (bkz.
    main.py::_sakin_kisi_id_dogrula)."""
    govde = {"kullanici_adi": kullanici_adi, "parola": "GucluParola123!", "rol": rol}
    if kisi_id is not None:
        govde["kisi_id"] = kisi_id
    r = client.post("/kullanicilar", json=govde, headers=yetkili_header)
    assert r.status_code == 200, r.text
    r2 = client.post("/auth/giris", json={"kullanici_adi": kullanici_adi, "parola": "GucluParola123!"})
    assert r2.status_code == 200, r2.text
    return {"Authorization": f"Bearer {r2.json()['token']}"}


@pytest.fixture(scope="module")
def sakin_kisi_id(client, yetkili_header):
    """'sakin' (site sakini öz-hizmet portalı) RBAC/işlevsellik testlerinde
    kullanılacak, sakin hesabına bağlanacak Kişi kaydı."""
    r = client.post("/kisiler", json={
        "ad_soyad": "Sakin Test Kişi", "plaka_no": "34 SKN 01", "tip": "abone",
        "daire_departman": "A Blok Daire 5",
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text
    return r.json()["id"]


@pytest.fixture(scope="module")
def sakin_header(client, admin_token, sakin_kisi_id):
    return _rol_ile_kullanici_olustur_ve_giris_yap(
        client, {"Authorization": f"Bearer {admin_token}"}, "rbac-sakin", "sakin", kisi_id=sakin_kisi_id
    )


@pytest.fixture(scope="module")
def izleyici_header(client, admin_token):
    # NOT: module-scope bir fixture, fonksiyon-scope'lu `yetkili_header`'a değil
    # doğrudan module-scope `admin_token`'a bağımlı olmalı (aksi halde pytest
    # "ScopeMismatch" hatası verir — daha dar scope'lu bir fixture'a bağımlı
    # olamaz).
    return _rol_ile_kullanici_olustur_ve_giris_yap(
        client, {"Authorization": f"Bearer {admin_token}"}, "rbac-izleyici", "izleyici"
    )


@pytest.fixture(scope="module")
def operator_header(client, admin_token):
    return _rol_ile_kullanici_olustur_ve_giris_yap(
        client, {"Authorization": f"Bearer {admin_token}"}, "rbac-operator", "operatör"
    )


# ------------------------------------------------------------------
# Sağlık / temel
# ------------------------------------------------------------------

def test_saglik_kontrolu(client):
    r = client.get("/saglik")
    assert r.status_code == 200


# ------------------------------------------------------------------
# Kimlik doğrulama
# ------------------------------------------------------------------

def test_ilk_yonetici_ve_giris_akisi(client, admin_token):
    assert admin_token  # fixture zaten başarılı giriş yaptı

    # İkinci kez ilk yönetici oluşturmaya çalışmak reddedilmeli
    r = client.post("/auth/ilk-yonetici", json={"kullanici_adi": "baska", "parola": "GucluParola123!"})
    assert r.status_code == 409


def test_yanlis_parola_ile_giris_reddedilir(client):
    r = client.post("/auth/giris", json={"kullanici_adi": "admin", "parola": "yanlis-parola"})
    assert r.status_code == 401


def test_tokensiz_korumali_uc_noktaya_erisim_reddedilir(client):
    r = client.get("/kisiler")
    assert r.status_code == 401


def test_giris_deneme_kilitlemesi(client):
    """Var olmayan bir kullanıcı adıyla 5 başarısız denemeden sonra kilitlenmeli (429)."""
    kullanici_adi = "hic-var-olmayan-kullanici"
    for _ in range(5):
        r = client.post("/auth/giris", json={"kullanici_adi": kullanici_adi, "parola": "yanlis"})
        assert r.status_code == 401
    r = client.post("/auth/giris", json={"kullanici_adi": kullanici_adi, "parola": "yanlis"})
    assert r.status_code == 429


# ------------------------------------------------------------------
# Plaka yetki kontrolü (yetkili / yetkisiz / kara liste / süresi dolmuş / saat kısıtlaması)
# ------------------------------------------------------------------

def test_taninmayan_plaka_yetkisiz_gelir(client, yetkili_header):
    r = client.post("/kayitlar", json={"plaka_no": "06 ZZZ 999", "kamera_id": "TEST", "yon": "giris"},
                     headers=yetkili_header)
    assert r.status_code == 200, r.text
    assert r.json()["yetki_durumu"] == "yetkisiz"


def test_kayitli_abone_yetkili_gelir(client, yetkili_header):
    r = client.post("/kisiler", json={
        "ad_soyad": "Test Abone",
        "plaka_no": "34 TEST 01",
        "tip": "abone",
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text

    r2 = client.post("/kayitlar", json={"plaka_no": "34 TEST 01", "kamera_id": "TEST", "yon": "giris"},
                      headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    assert r2.json()["yetki_durumu"] == "yetkili"


def test_kara_listedeki_plaka_engellenir(client, yetkili_header):
    r = client.post("/kara-listesi", json={"plaka_no": "34 KARA 02", "sebep": "test"}, headers=yetkili_header)
    assert r.status_code == 200, r.text

    r2 = client.post("/kayitlar", json={"plaka_no": "34 KARA 02", "kamera_id": "TEST", "yon": "giris"},
                      headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    assert r2.json()["yetki_durumu"] == "kara_liste"


def test_kara_liste_farkli_bosluklu_girişte_de_eslesir(client, yetkili_header):
    """GÜVENLİK KÖK NEDEN REGRESYON TESTİ (2026-09-17): main.py::_plaka_yetki_kontrol
    kara liste kontrolünü önceden boşluksuz normalize edilmiş bir değerle
    (`_plaka_normalize`), boşluklu SAKLANAN `KaraListesi.plaka_no` (bkz.
    schemas.py::plaka_normalize, boşlukları KORUR) arasında DOĞRUDAN SQL
    eşitliğiyle yapıyordu -- iki taraf da farklı biçimde olduğu için ASLA
    eşleşmiyordu, yani kara liste engeli FİİLEN HİÇ ÇALIŞMIYORDU (yalnızca
    yukarıdaki test gibi aynı boşluk biçimiyle tesadüfen "çalışıyormuş gibi"
    görünebilirdi -- oysa o test de aslında başarısız olurdu, çünkü sorgu
    hiçbir zaman boşluklu bir değerle eşleşmiyordu). Bu test, kara listeye
    BOŞLUKSUZ girilen bir plakanın, kamerada BOŞLUKLU okunan aynı plakayı da
    doğru şekilde engellediğini kanıtlar."""
    r = client.post("/kara-listesi", json={"plaka_no": "34KARA03", "sebep": "test"}, headers=yetkili_header)
    assert r.status_code == 200, r.text

    r2 = client.post("/kayitlar", json={"plaka_no": "34 KARA 03", "kamera_id": "TEST", "yon": "giris"},
                      headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    assert r2.json()["yetki_durumu"] == "kara_liste"


def test_ek_plaka_yetkili_gelir(client, yetkili_header):
    """GÜVENLİK KÖK NEDEN REGRESYON TESTİ (2026-09-17): bir kişinin ek/ikincil
    plakası (bkz. POST /kisiler/{id}/plakalar) da aynı boşluk-normalizasyonu
    hatasından etkileniyordu -- gerçek bir geçişte hiçbir zaman "yetkili"
    olarak tanınmıyordu (yalnızca Kisi.plaka_no ANA plaka üzerinden yapılan
    eşleştirme -- Python tarafında zaten doğru normalize ediliyordu --
    çalışıyordu)."""
    r = client.post("/kisiler", json={
        "ad_soyad": "Ek Plaka Testi", "plaka_no": "34 EKA 01", "tip": "abone",
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text
    kisi_id = r.json()["id"]

    r2 = client.post(f"/kisiler/{kisi_id}/plakalar", json={"plaka_no": "34 EKA 02"}, headers=yetkili_header)
    assert r2.status_code == 200, r2.text

    r3 = client.post("/kayitlar", json={"plaka_no": "34 EKA 02", "kamera_id": "TEST", "yon": "giris"},
                      headers=yetkili_header)
    assert r3.status_code == 200, r3.text
    assert r3.json()["yetki_durumu"] == "yetkili"
    assert r3.json()["kisi_id"] == kisi_id


def test_suresi_dolmus_ziyaretci_dogru_isaretlenir(client, yetkili_header):
    gecmis_tarih = (datetime.now() - timedelta(days=1)).isoformat()
    r = client.post("/kisiler", json={
        "ad_soyad": "Eski Ziyaretci",
        "plaka_no": "34 ESKI 03",
        "tip": "ziyaretci",
        "bitis_tarihi": gecmis_tarih,
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text

    r2 = client.post("/kayitlar", json={"plaka_no": "34 ESKI 03", "kamera_id": "TEST", "yon": "giris"},
                      headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    assert r2.json()["yetki_durumu"] == "suresi_dolmus"


def test_saat_disi_erisim_yetkisiz_sayilir(client, yetkili_header):
    """Bu regresyon testi özellikle önemli: schemas.py'deki tekrar tanımlanmış
    (duplicate) sınıflar yüzünden bu alanlar API'den API'ye hiç ulaşmıyordu.
    Burada saat penceresinin GERÇEKTEN uygulandığını uçtan uca doğruluyoruz."""
    simdi = datetime.now()
    # Şu andan 2 saat önce başlayıp 1 saat önce biten bir pencere: garanti olarak "şu an" dışında.
    baslangic = (simdi - timedelta(hours=2)).strftime("%H:%M")
    bitis = (simdi - timedelta(hours=1)).strftime("%H:%M")

    r = client.post("/kisiler", json={
        "ad_soyad": "Saat Kisitli Personel",
        "plaka_no": "34 SAAT 04",
        "tip": "personel",
        "giris_saati_baslangic": baslangic,
        "giris_saati_bitis": bitis,
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text
    olusturulan = r.json()
    # Alan gerçekten kaydedilmiş mi? (Sadece istek kabul edildi diye yetmez, geri de dönmeli.)
    assert olusturulan["giris_saati_baslangic"] == baslangic
    assert olusturulan["giris_saati_bitis"] == bitis

    r2 = client.post("/kayitlar", json={"plaka_no": "34 SAAT 04", "kamera_id": "TEST", "yon": "giris"},
                      headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    assert r2.json()["yetki_durumu"] == "yetkisiz", (
        "Saat penceresi dışında olmasına rağmen 'yetkili' döndü — "
        "muhtemelen schemas.py'deki alanlar API'ye ulaşmıyor"
    )


# ------------------------------------------------------------------
# Kamera: giriş doğrulama + lisans/limit kontrolü
# ------------------------------------------------------------------

def test_kamera_gecersiz_adres_semasi_reddedilir(client, yetkili_header):
    r = client.post("/kameralar", json={"ad": "Test Kamera", "rtsp_url": "file:///etc/passwd", "yon": "giris"},
                     headers=yetkili_header)
    assert r.status_code == 400


def test_kamera_gecersiz_yon_reddedilir(client, yetkili_header):
    r = client.post("/kameralar", json={"ad": "Test Kamera", "rtsp_url": "rtsp://127.0.0.1/test", "yon": "yukari"},
                     headers=yetkili_header)
    assert r.status_code == 400


def test_lisanssiz_kamera_eklenemez(client, yetkili_header):
    r = client.get("/lisans", headers=yetkili_header)
    assert r.status_code == 200
    if r.json().get("aktif"):
        pytest.skip("Bu test sırasında lisans zaten aktifleştirilmiş (test sırası bağımlılığı)")
    r2 = client.post("/kameralar", json={"ad": "Kamera 1", "rtsp_url": "rtsp://127.0.0.1/test", "yon": "giris"},
                      headers=yetkili_header)
    assert r2.status_code == 403


def test_lisans_aktivasyonu_ve_kamera_limiti(client, yetkili_header):
    # kamera_limiti=1 olan, cihaza kilitli olmayan bir lisans üret ve aktive et.
    anahtar = lisans.uret("Test Site", kamera_limiti=1, gun=30)
    r = client.post("/lisans/aktive-et", json={"anahtar": anahtar}, headers=yetkili_header)
    assert r.status_code == 200, r.text
    assert r.json()["aktif"] is True
    assert r.json()["kamera_limiti"] == 1

    r2 = client.post("/kameralar", json={"ad": "Kamera 1", "rtsp_url": "rtsp://127.0.0.1/kanal1", "yon": "giris"},
                      headers=yetkili_header)
    assert r2.status_code == 200, r2.text

    # Limit 1: ikinci kamera reddedilmeli
    r3 = client.post("/kameralar", json={"ad": "Kamera 2", "rtsp_url": "rtsp://127.0.0.1/kanal2", "yon": "cikis"},
                      headers=yetkili_header)
    assert r3.status_code == 403


def test_kamera_silme_loglanir(client, caplog, yetkili_header):
    """2026-09-20: bir kameranın kaldırılması (kazayla ya da kötü niyetle)
    önceden hiçbir yere loglanmıyordu (bkz. main.py::kamera_sil)."""
    anahtar = lisans.uret("Kamera Silme Test Site", kamera_limiti=5, gun=30)
    client.post("/lisans/aktive-et", json={"anahtar": anahtar}, headers=yetkili_header)
    r = client.post("/kameralar", json={
        "ad": "Silinecek Denetim Kamerası", "rtsp_url": "rtsp://127.0.0.1/silinecek", "yon": "giris",
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text
    kamera_id = r.json()["id"]
    with caplog.at_level("INFO", logger="pts"):
        r2 = client.delete(f"/kameralar/{kamera_id}", headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    assert kamera_id in caplog.text
    assert "Silinecek Denetim Kamerası" in caplog.text

    # 2026-09-20 (devam, kullanıcı talebi): log dosyasına EK olarak kalıcı
    # `denetim_kayitlari` tablosuna da yazılmalı (bkz. main.py::_denetim_kaydet).
    r3 = client.get("/denetim-kayitlari", params={"eylem": "kamera_sil"}, headers=yetkili_header)
    assert r3.status_code == 200, r3.text
    aciklamalar = " ".join(k["aciklama"] for k in r3.json())
    assert kamera_id in aciklamalar
    assert "Silinecek Denetim Kamerası" in aciklamalar


def test_lisans_aktivasyonu_loglanir(client, caplog, yetkili_header):
    """2026-09-20: bir lisansın aktive edilmesi (hangi müşteri, hangi lisans
    id'si, ne zamana kadar) önceden hiçbir yere loglanmıyordu."""
    anahtar = lisans.uret("Loglama Test Müşterisi", kamera_limiti=2, gun=30)
    with caplog.at_level("INFO", logger="pts"):
        r = client.post("/lisans/aktive-et", json={"anahtar": anahtar}, headers=yetkili_header)
    assert r.status_code == 200, r.text
    assert "Loglama Test Müşterisi" in caplog.text
    assert r.json()["lisans_id"] in caplog.text

    r2 = client.get("/denetim-kayitlari", params={"eylem": "lisans_aktivasyon"}, headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    assert any("Loglama Test Müşterisi" in k["aciklama"] for k in r2.json())


def test_lisans_kalan_gun_ve_yakinda_doluyor_bayragi(client, yetkili_header):
    """2026-09-20: lisans süresi dolmadan ÖNCE panelin uyarı gösterebilmesi
    için /lisans yanıtına eklenen kalan_gun/yakinda_doluyor alanlarını
    doğrular (bkz. main.py::_lisans_kalan_gun_ekle). ÖNCEKİ davranış tamamen
    ikiliydi (aktif/pasif) -- süre dolmadan hiçbir erken uyarı yoktu."""
    anahtar = lisans.uret("Kısa Süreli Müşteri", kamera_limiti=5, gun=5)
    r = client.post("/lisans/aktive-et", json={"anahtar": anahtar}, headers=yetkili_header)
    assert r.status_code == 200, r.text
    assert r.json()["kalan_gun"] == 5
    assert r.json()["yakinda_doluyor"] is True

    r2 = client.get("/lisans", headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    assert r2.json()["kalan_gun"] == 5
    assert r2.json()["yakinda_doluyor"] is True


def test_lisans_uzun_surede_yakinda_doluyor_false_doner(client, yetkili_header):
    anahtar = lisans.uret("Uzun Süreli Müşteri", kamera_limiti=5, gun=365)
    r = client.post("/lisans/aktive-et", json={"anahtar": anahtar}, headers=yetkili_header)
    assert r.status_code == 200, r.text
    assert r.json()["kalan_gun"] == 365
    assert r.json()["yakinda_doluyor"] is False


def test_gecersiz_imzali_lisans_reddedilir(client, yetkili_header):
    sahte = "PTS1.eW9rLWJvenVsbXVzLXZlcmk.eW9rLWltemE"
    r = client.post("/lisans/aktive-et", json={"anahtar": sahte}, headers=yetkili_header)
    assert r.status_code == 400


# ------------------------------------------------------------------
# Görüntü saklama / otomatik temizlik
# ------------------------------------------------------------------
# NOT: Bu testler `_goruntu_temizle_calistir` fonksiyonunu (hem manuel
# /sistem/goruntu-temizle uç noktası hem de arka planda periyodik çalışan
# otomatik saklama görevi tarafından paylaşılan mantık) uçtan uca doğrular.
# Amaç: disk sürekli dolup sistemi "donmuş" gibi göstermesin diye eklenen bu
# otomasyonun, YANLIŞLIKLA yeni/güncel kayıtların görüntülerini silmediğini
# ve gerçekten SADECE eski kayıtları temizlediğini garanti altına almak.

def _test_goruntu_dosyasi_olustur(client, dosya_adi: str) -> str:
    from backend.main import GORUNTU_KLASORU
    os.makedirs(GORUNTU_KLASORU, exist_ok=True)
    tam_yol = os.path.join(GORUNTU_KLASORU, dosya_adi)
    with open(tam_yol, "wb") as f:
        f.write(b"sahte-goruntu-verisi")
    return tam_yol


def test_eski_goruntulu_kayit_temizlenir_yeni_olan_korunur(client, yetkili_header):
    from backend.database import SessionLocal
    from backend import models

    eski_dosya_adi = "test-eski-goruntu.jpg"
    yeni_dosya_adi = "test-yeni-goruntu.jpg"
    eski_tam_yol = _test_goruntu_dosyasi_olustur(client, eski_dosya_adi)
    yeni_tam_yol = _test_goruntu_dosyasi_olustur(client, yeni_dosya_adi)

    db = SessionLocal()
    try:
        eski_kayit = models.Kayit(
            plaka_no="34 ESK 99", kamera_id="TEST",
            tarih_saat=datetime.now() - timedelta(days=45),
            goruntu_yolu=eski_tam_yol,
        )
        yeni_kayit = models.Kayit(
            plaka_no="34 YEN 99", kamera_id="TEST",
            tarih_saat=datetime.now() - timedelta(hours=1),
            goruntu_yolu=yeni_tam_yol,
        )
        db.add_all([eski_kayit, yeni_kayit])
        db.commit()
    finally:
        db.close()

    try:
        r = client.post("/sistem/goruntu-temizle", params={"gun": 30}, headers=yetkili_header)
        assert r.status_code == 200, r.text
        assert r.json()["silinen_goruntu"] >= 1

        assert not os.path.exists(eski_tam_yol), "30 günden eski görüntü silinmeliydi"
        assert os.path.exists(yeni_tam_yol), (
            "1 saat önceki kayıt YANLIŞLIKLA silindi — tarih filtresinde bir "
            "regresyon olabilir (otomatik saklama görevi de aynı fonksiyonu kullanıyor)"
        )
    finally:
        for yol in (eski_tam_yol, yeni_tam_yol):
            if os.path.exists(yol):
                os.remove(yol)


# ------------------------------------------------------------------
# GET /sistem/yedek — veritabanı yedeği indirme (RBAC), 2026-09-20
# ------------------------------------------------------------------
# Geniş kapsamlı bir kod denetiminde bu uç noktanın hiç test kapsamında
# olmadığı tespit edildi -- tam da "sessiz yetki sızıntısı" riski taşıyan
# türden bir uç nokta (tüm veritabanının HAM bir kopyasını indirir).

def test_sistem_yedek_yonetici_indirebilir(client, yetkili_header):
    r = client.get("/sistem/yedek", headers=yetkili_header)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/octet-stream"
    assert len(r.content) > 0


def test_sistem_yedek_operator_yetkisiz_403_doner(client, operator_header):
    r = client.get("/sistem/yedek", headers=operator_header)
    assert r.status_code == 403, r.text


def test_sistem_yedek_izleyici_yetkisiz_403_doner(client, izleyici_header):
    r = client.get("/sistem/yedek", headers=izleyici_header)
    assert r.status_code == 403, r.text


def test_sistem_yedek_girissiz_401_doner(client):
    r = client.get("/sistem/yedek")
    assert r.status_code == 401, r.text


# ------------------------------------------------------------------
# GET /sakin/goruntu/{kayit_id} — sakinin kendi kaydının fotoğrafı (IDOR), 2026-09-20
# ------------------------------------------------------------------
# Bu uç nokta da denetimde test kapsamı dışı bulundu -- tam da bir sakinin
# BAŞKA bir sakinin fotoğrafını görüp göremediğini (IDOR) doğrulaması gereken,
# güvenlik açısından hassas bir sınır.

def _test_goruntulu_kayit_olustur(kisi_id, plaka_no: str, dosya_adi: str):
    """`_test_goruntu_dosyasi_olustur`ın aksine, ORM'i atlamadan doğrudan
    `kisi_id`ye bağlı gerçek bir dosyalı Kayit satırı oluşturur (normal
    `/kayitlar` POST akışı `goruntu_yolu` alanını hiç doldurmaz -- bu yalnızca
    `/kayitlar/otomatik`nın kamera yüklemesiyle dolar)."""
    from backend.database import SessionLocal
    from backend import models
    from backend.main import GORUNTU_KLASORU

    os.makedirs(GORUNTU_KLASORU, exist_ok=True)
    tam_yol = os.path.join(GORUNTU_KLASORU, dosya_adi)
    with open(tam_yol, "wb") as f:
        f.write(b"sahte-sakin-goruntusu")

    db = SessionLocal()
    try:
        kayit = models.Kayit(
            plaka_no=plaka_no, kamera_id="TEST-SAKIN-GORUNTU", tarih_saat=datetime.now(),
            kisi_id=kisi_id, goruntu_yolu=tam_yol,
        )
        db.add(kayit)
        db.commit()
        db.refresh(kayit)
        return kayit.id, tam_yol
    finally:
        db.close()


def test_sakin_kendi_kaydinin_goruntusunu_gorebilir(client, sakin_header, sakin_kisi_id):
    kayit_id, tam_yol = _test_goruntulu_kayit_olustur(sakin_kisi_id, "34 SKG 01", "test-sakin-kendi-goruntusu.jpg")
    try:
        r = client.get(f"/sakin/goruntu/{kayit_id}", headers=sakin_header)
        assert r.status_code == 200, r.text
        assert r.content == b"sahte-sakin-goruntusu"
    finally:
        if os.path.exists(tam_yol):
            os.remove(tam_yol)


def test_sakin_baska_kisinin_goruntusunu_goremez_idor(client, sakin_header, yetkili_header):
    """IDOR regresyonu: bir sakin, kendi kisi_id'sine ait OLMAYAN bir kayıt
    id'sini tahmin ederek o kaydın fotoğrafını görememeli (404, dosyanın var
    olup olmadığından bağımsız olarak -- bkz. main.py::sakin_goruntu)."""
    rk = client.post("/kisiler", json={
        "ad_soyad": "Başka Sakin Görüntü", "plaka_no": "34 BSG 02", "tip": "abone",
    }, headers=yetkili_header)
    baska_kisi_id = rk.json()["id"]
    kayit_id, tam_yol = _test_goruntulu_kayit_olustur(baska_kisi_id, "34 BSG 02", "test-baska-sakinin-goruntusu.jpg")
    try:
        r = client.get(f"/sakin/goruntu/{kayit_id}", headers=sakin_header)
        assert r.status_code == 404, r.text
    finally:
        if os.path.exists(tam_yol):
            os.remove(tam_yol)


def test_sakin_goruntu_dosyasi_diskte_yoksa_404_doner(client, sakin_header, sakin_kisi_id):
    """Kayıt kendisine ait olsa bile, dosya diskten (örn. temizlik görevi
    tarafından) silinmişse hâlâ 404 dönmeli, 500 değil."""
    from backend.database import SessionLocal
    from backend import models

    db = SessionLocal()
    try:
        kayit = models.Kayit(
            plaka_no="34 SKG 03", kamera_id="TEST-SAKIN-GORUNTU", tarih_saat=datetime.now(),
            kisi_id=sakin_kisi_id, goruntu_yolu="/tmp/bu-dosya-hic-var-olmadi-pts-test.jpg",
        )
        db.add(kayit)
        db.commit()
        db.refresh(kayit)
        kayit_id = kayit.id
    finally:
        db.close()

    r = client.get(f"/sakin/goruntu/{kayit_id}", headers=sakin_header)
    assert r.status_code == 404, r.text


def test_sakin_goruntu_olmayan_kayit_id_404_doner(client, sakin_header):
    r = client.get("/sakin/goruntu/999999999", headers=sakin_header)
    assert r.status_code == 404, r.text


def test_sakin_goruntu_personel_rolune_kapali_403_doner(client, izleyici_header):
    """Bu uç nokta yalnızca 'sakin' rolüne açık -- personel (izleyici dahil)
    kendi ayrı uçlarını (`/goruntuler/{dosya_adi}`) kullanmalı."""
    r = client.get("/sakin/goruntu/1", headers=izleyici_header)
    assert r.status_code == 403, r.text


# ------------------------------------------------------------------
# Bilinen plakaya göre OCR düzeltmesi (veritabanı çapraz kontrolü)
# ------------------------------------------------------------------
# Bu, doğruluğu artırmak için eklenen bir tekniktir: OCR tek bir karakteri
# yanlış okusa bile (düşük güvenle), sahada kayıtlı bilinen bir plakayla tek
# karakter farkı varsa ve başka hiçbir aday bu kadar yakın değilse, o bilinen
# plakaya düzeltilir. GÜVENLİK AÇISINDAN KRİTİK: bu SADECE erişim vermek için
# çalışmalı, kara listeye asla uygulanmamalı (aksi halde alakasız bir araç
# yanlışlıkla engellenebilir).

def test_dusuk_guvenli_tek_karakter_hatasi_bilinen_plakaya_duzeltilir(client, yetkili_header):
    r = client.post("/kisiler", json={
        "ad_soyad": "Duzeltme Testi", "plaka_no": "34 DZT 123", "tip": "abone",
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text

    # OCR "3" yerine "8" okumuş (tek karakter hatası), düşük güvenle.
    r2 = client.post("/kayitlar", json={
        "plaka_no": "34 DZT 128", "kamera_id": "TEST", "yon": "giris", "guven_skoru": 0.55,
    }, headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    sonuc = r2.json()
    assert sonuc["plaka_no"] == "34 DZT 123", "Bilinen plakaya düzeltilmedi"
    assert sonuc["ham_plaka_metni"] == "34 DZT 128", "Ham OCR okuması denetim için saklanmalıydı"
    assert sonuc["yetki_durumu"] == "yetkili", (
        "Düzeltme uygulandığı halde 'yetkisiz' döndü — düzeltme yetki kontrolünden ÖNCE uygulanmalı"
    )


def test_yuksek_guvenli_okuma_duzeltilmez(client, yetkili_header):
    """Zaten yüksek güvenli (>= eşik) bir okuma, bilinen bir plakaya yakın olsa
    bile OLDUĞU GİBİ bırakılmalı — aksi halde doğru bir okuma yanlışlıkla
    'düzeltilerek' bozulabilir."""
    r = client.post("/kisiler", json={
        "ad_soyad": "Yuksek Guven Testi", "plaka_no": "34 YKG 123", "tip": "abone",
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text

    r2 = client.post("/kayitlar", json={
        "plaka_no": "34 YKG 128", "kamera_id": "TEST", "yon": "giris", "guven_skoru": 0.97,
    }, headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    sonuc = r2.json()
    assert sonuc["plaka_no"] == "34 YKG 128", "Yüksek güvenli okuma yanlışlıkla değiştirildi"
    assert sonuc["ham_plaka_metni"] is None
    assert sonuc["yetki_durumu"] == "yetkisiz"


def test_kara_listeye_yakinlik_duzeltmesi_uygulanmaz(client, yetkili_header):
    """Güvenlik regresyon testi: kara listedeki bir plakaya YAKIN ama farklı
    (ve bilinmeyen) bir plaka, düzeltme mekanizması yüzünden YANLIŞLIKLA
    engellenmemeli. Düzeltme sadece bilinen (abone/personel) plakalara karşı
    çalışır, kara listeye karşı hiç çalışmaz."""
    r = client.post("/kara-listesi", json={"plaka_no": "34 KRY 123", "sebep": "test"}, headers=yetkili_header)
    assert r.status_code == 200, r.text

    r2 = client.post("/kayitlar", json={
        "plaka_no": "34 KRY 128", "kamera_id": "TEST", "yon": "giris", "guven_skoru": 0.55,
    }, headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    sonuc = r2.json()
    assert sonuc["yetki_durumu"] != "kara_liste", (
        "Kara listedeki bir plakaya yakın farklı bir araç yanlışlıkla kara listeye düzeltildi/engellendi"
    )


# ------------------------------------------------------------------
# Rol bazlı yetkilendirme (RBAC)
# ------------------------------------------------------------------
# Daha önce sadece kullanıcı yönetimi / sistem ayarları / DB yedeği / görüntü
# temizleme rol kontrolü yapıyordu; kamera/kişi/kara liste/bariyer açma/webhook
# gibi geri kalan HER YAZMA işlemi sadece giriş yapmış olmayı yeterli
# sayıyordu — yani "izleyici" (salt okunur) rolündeki bir kullanıcı bile
# bariyer açabiliyor, kamera silebiliyor, lisans aktive edebiliyordu. Bu
# testler _rol_dogrula() ile eklenen kısıtlamaları uçtan uca doğrular.

def test_izleyici_hicbir_yazma_islemi_yapamaz(client, izleyici_header):
    """İzleyici (salt okunur) rolü GERÇEKTEN salt okunur olmalı."""
    denemeler = [
        ("post", "/kisiler", {"ad_soyad": "X", "plaka_no": "34 RBC 001", "tip": "abone"}),
        ("post", "/kameralar", {"ad": "X", "rtsp_url": "rtsp://127.0.0.1/x", "yon": "giris"}),
        ("post", "/kara-listesi", {"plaka_no": "34 RBC 002", "sebep": "test"}),
        ("post", "/bariyer/ayarlar", {"ad": "X", "mod": "simulate"}),
        ("post", "/lisans/aktive-et", {"anahtar": "PTS1.x.y"}),
        ("post", "/bildirim/ayarlar", {"ad": "X", "hedef": "https://example.com"}),
        ("put", "/sistem/ayarlar", {"supheli_esik": 5}),
        ("post", "/kullanicilar", {"kullanici_adi": "baska", "parola": "GucluParola123!", "rol": "izleyici"}),
        ("post", "/siteler", {"ad": "X"}),
    ]
    for metot, yol, govde in denemeler:
        r = getattr(client, metot)(yol, json=govde, headers=izleyici_header)
        assert r.status_code == 403, f"{metot.upper()} {yol}: izleyici 403 almalıydı, {r.status_code} aldı ({r.text})"


def test_izleyici_okuma_uc_noktalarina_erisebilir(client, izleyici_header):
    """Kısıtlama SADECE yazma işlemlerinde olmalı; izleyici okuyabilmeli."""
    for yol in ("/kisiler", "/kameralar", "/kara-listesi", "/kayitlar", "/bariyer/ayarlar"):
        r = client.get(yol, headers=izleyici_header)
        assert r.status_code == 200, f"GET {yol}: izleyici erişemedi ({r.status_code})"


def test_sistem_loglarina_izleyici_erisemez_operator_erisebilir(client, izleyici_header, operator_header):
    """Güvenlik regresyonu: /sistem/loglar kamera bağlantı adresleri (RTSP
    kimlik bilgileri dahil, maskelenmiş olsa da) ve dahili hata detayları
    içerir — salt-okunur 'izleyici' rolüne açık kalmamalı. En az operatör
    olmalı."""
    r = client.get("/sistem/loglar", headers=izleyici_header)
    assert r.status_code == 403, f"izleyici loglara erişebildi ({r.status_code})"
    r2 = client.get("/sistem/loglar", headers=operator_header)
    assert r2.status_code == 200, f"operatör loglara erişemedi ({r2.text})"


# ------------------------------------------------------------------
# /goruntuler — araç/sürücü görselleri (KVKK kapsamında kişisel veri)
# ------------------------------------------------------------------
# Güvenlik regresyonu: bu yol önceden kimliksiz bir StaticFiles mount'uydu.
# Artık giriş yapmamış hiç kimse (rol farketmeksizin) bu görsellere
# erişememeli, ve dosya adı yalnızca DÜZ bir ad olmalı (yol geçişi/"../"
# denemeleri reddedilmeli).

def test_goruntu_girissiz_erisilemez(client):
    r = client.get("/goruntuler/herhangi_bir_dosya.jpg")
    assert r.status_code in (401, 403), f"Girişsiz istek engellenmedi ({r.status_code})"


def test_goruntu_yol_gecisi_denemesi_reddedilir(client, izleyici_header):
    """NOT: Starlette'in varsayılan `{dosya_adi}` yol dönüştürücüsü zaten TEK
    bir segment içinde eşleşir (ham bir '/' içeren istekler bu uç noktaya HİÇ
    ulaşmaz, 404 döner) — bu yüzden burada özellikle TEK segment içinde
    kalan ama yine de ".." içeren bir değeri test ediyoruz: gorsel_getir'in
    kendi ".." kontrolünün (yalnızca yönlendirme katmanına güvenmeden)
    gerçekten çalıştığını doğrudan doğrular."""
    r = client.get("/goruntuler/..gizli_dosya.jpg", headers=izleyici_header)
    assert r.status_code == 400, f"'..' içeren dosya adı reddedilmedi ({r.status_code})"


def test_goruntu_girisli_kullanici_gercek_dosyayi_alabilir(client, izleyici_header, tmp_path_factory):
    from backend.main import GORUNTU_KLASORU

    dosya_adi = "pytest_gecici_test_gorseli.jpg"
    tam_yol = os.path.join(GORUNTU_KLASORU, dosya_adi)
    with open(tam_yol, "wb") as f:
        f.write(b"\xff\xd8\xff\xe0sahte-jpeg-icerigi")
    try:
        r = client.get(f"/goruntuler/{dosya_adi}", headers=izleyici_header)
        assert r.status_code == 200, r.text
        assert r.content.startswith(b"\xff\xd8\xff")
    finally:
        os.remove(tam_yol)


def test_operator_gunluk_islemleri_yapabilir_ama_yonetim_islemlerini_yapamaz(client, operator_header):
    """Operatör günlük operasyonu (kişi/kamera/kara liste/bariyer açma) yapabilmeli,
    ama yönetici'ye özel işlemleri (kullanıcı yönetimi, sistem ayarları, webhook
    yapılandırması, lisans) YAPAMAMALI."""
    r = client.post("/kisiler", json={
        "ad_soyad": "Operator Testi", "plaka_no": "34 RBC 003", "tip": "abone",
    }, headers=operator_header)
    assert r.status_code == 200, f"Operatör kişi ekleyemedi: {r.text}"

    r2 = client.post("/kara-listesi", json={"plaka_no": "34 RBC 004", "sebep": "test"}, headers=operator_header)
    assert r2.status_code == 200, f"Operatör kara listeye ekleyemedi: {r2.text}"

    yonetim_denemeleri = [
        ("post", "/kullanicilar", {"kullanici_adi": "baska2", "parola": "GucluParola123!", "rol": "izleyici"}),
        ("put", "/sistem/ayarlar", {"supheli_esik": 5}),
        ("post", "/bildirim/ayarlar", {"ad": "X", "hedef": "https://example.com"}),
        ("post", "/lisans/aktive-et", {"anahtar": "PTS1.x.y"}),
        # Site/Nokta yönetimi de sadece yönetici'ye açık — operatör günlük
        # operasyon yapar ama yerleşke topolojisini değiştiremez.
        ("post", "/siteler", {"ad": "Operatör Sitesi"}),
        ("post", "/noktalar", {"site_id": 1, "ad": "X"}),
    ]
    for metot, yol, govde in yonetim_denemeleri:
        r3 = getattr(client, metot)(yol, json=govde, headers=operator_header)
        assert r3.status_code == 403, f"{metot.upper()} {yol}: operatör 403 almalıydı, {r3.status_code} aldı"


# ------------------------------------------------------------------
# Site / Erişim Noktası (kamera ↔ bariyer bağlantısı)
# ------------------------------------------------------------------
# Nokta, kameralar (cameras.json'da tutulan) ile bariyerler (SQL tablosu)
# arasındaki TEK bağlantıdır — bu yüzden her iki foreign key de (var
# olduklarında) gerçekten var olan kayıtları göstermeli, aksi halde panel
# sessizce geçersiz bir bağlantı gösterir (bkz. main.py::nokta_ekle).

def test_site_olusturulur_ve_listelenir(client, yetkili_header):
    r = client.post("/siteler", json={"ad": "Test Yerleşkesi", "aciklama": "RBAC testi için"}, headers=yetkili_header)
    assert r.status_code == 200, r.text
    site_id = r.json()["id"]

    r2 = client.get("/siteler", headers=yetkili_header)
    assert r2.status_code == 200
    assert any(s["id"] == site_id for s in r2.json())


def test_nokta_olmayan_bariyer_id_ile_reddedilir(client, yetkili_header):
    site = client.post("/siteler", json={"ad": "Nokta Test Sitesi 1"}, headers=yetkili_header).json()
    r = client.post("/noktalar", json={
        "site_id": site["id"], "ad": "Nizamiye", "bariyer_id": 999999,
    }, headers=yetkili_header)
    assert r.status_code == 404


def test_nokta_olmayan_kamera_id_ile_reddedilir(client, yetkili_header):
    site = client.post("/siteler", json={"ad": "Nokta Test Sitesi 2"}, headers=yetkili_header).json()
    r = client.post("/noktalar", json={
        "site_id": site["id"], "ad": "Nizamiye", "kamera_id": "olmayan-kamera-uuid",
    }, headers=yetkili_header)
    assert r.status_code == 404


def test_nokta_gecerli_bariyerle_olusturulur_ve_bariyer_acmaya_baglanir(client, yetkili_header):
    site = client.post("/siteler", json={"ad": "Nokta Test Sitesi 3"}, headers=yetkili_header).json()
    bariyer = client.post("/bariyer/ayarlar", json={"ad": "Nizamiye Bariyeri", "mod": "simulate"}, headers=yetkili_header).json()

    r = client.post("/noktalar", json={
        "site_id": site["id"], "ad": "Nizamiye Giriş", "yon": "giris", "bariyer_id": bariyer["id"],
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text
    nokta = r.json()
    assert nokta["bariyer_id"] == bariyer["id"]

    # Panelin "Bariyer Aç" akışının dayandığı gerçek zincir: nokta -> bariyer_id -> /bariyer/{id}/ac
    r2 = client.post(f"/bariyer/{nokta['bariyer_id']}/ac", headers=yetkili_header)
    assert r2.status_code == 200, r2.text


def test_nokta_kamera_adi_id_degil_gercek_ad_degerini_dondurur(client, yetkili_header):
    """schemas.NoktaCevap.kamera_adi (2026-09-21, "id vs ad" hata sınıfı --
    bkz. o alanın docstring'i): GET /noktalar, bağlı kameranın "id"sini
    DEĞİL, gerçek geçiş kayıtlarında (Kayit.kamera_id) kullanılan "ad"
    değerini de döndürmeli -- aksi halde frontend (app.js::olayDetayAc) bir
    olayın hangi Nokta'ya ait olduğunu HİÇBİR ZAMAN bulamaz (id != ad)."""
    anahtar = lisans.uret("Nokta Kamera Adi Test", kamera_limiti=10, gun=30)
    r = client.post("/lisans/aktive-et", json={"anahtar": anahtar}, headers=yetkili_header)
    assert r.status_code == 200, r.text
    kamera_ad = "Nokta Testi Kamerası"
    kamera = client.post("/kameralar", json={
        "ad": kamera_ad, "rtsp_url": "rtsp://127.0.0.1/nokta-testi", "yon": "giris",
    }, headers=yetkili_header).json()
    assert kamera["id"] != kamera_ad

    site = client.post("/siteler", json={"ad": "Nokta Kamera Adi Test Sitesi"}, headers=yetkili_header).json()
    nokta = client.post("/noktalar", json={
        "site_id": site["id"], "ad": "Nokta Kamera Adi Test Noktası", "kamera_id": kamera["id"],
    }, headers=yetkili_header).json()
    assert nokta["kamera_id"] == kamera["id"]

    r = client.get("/noktalar", headers=yetkili_header)
    assert r.status_code == 200, r.text
    donen_nokta = next(n for n in r.json() if n["id"] == nokta["id"])
    assert donen_nokta["kamera_adi"] == kamera_ad, (
        "kamera_adi, kameranın 'ad' alanına eşit olmalı (id'ye DEĞİL) -- "
        "frontend olay detayında Nokta'yı Kayit.kamera_id (ad bazlı) ile eşleştirir"
    )


def test_site_silinince_bagli_nokta_da_silinir(client, yetkili_header):
    site = client.post("/siteler", json={"ad": "Silinecek Site"}, headers=yetkili_header).json()
    nokta = client.post("/noktalar", json={"site_id": site["id"], "ad": "Silinecek Nokta"}, headers=yetkili_header).json()

    r = client.delete(f"/siteler/{site['id']}", headers=yetkili_header)
    assert r.status_code == 200, r.text

    kalanlar = client.get("/noktalar", headers=yetkili_header).json()
    assert not any(n["id"] == nokta["id"] for n in kalanlar), "Site silindiğinde bağlı nokta da silinmeliydi (cascade)"


# ------------------------------------------------------------------
# /sistem/saglik — kimlik doğrulama zorunluluğu ve "güvenlik uyarıları" (2026-09-20)
# ------------------------------------------------------------------
# Geniş kapsamlı bir kod denetiminde tespit edildi: bu uç nokta hiçbir kimlik
# doğrulaması istemiyordu -- ağa erişimi olan HERKES (oturum açmadan) aktif
# pipeline/SSE istemci sayısını, yedek durumunu vb. görebiliyordu. Ayrıca,
# PTS_LICENSE_SECRET/PTS_KAMERA_ANAHTARI/PTS_CORS_ORIGINS gibi "varsayılana
# sessizce düşülürse güvensiz" ayarlar yalnızca başlangıçta BİR KEZ log
# dosyasına yazılıyordu -- kimse günlük olarak log dosyasını açıp okumadığı
# için bu, fark edilmeyen bir güvenlik borcuydu. Artık hem kimlik doğrulaması
# zorunlu hem de bu uyarılar panelin de kullandığı yanıtın bir parçası.

def test_sistem_sagligi_girissiz_401_doner(client):
    r = client.get("/sistem/saglik")
    assert r.status_code == 401, r.text


def test_sistem_sagligi_guvenlik_uyarilari_alani_var(client, izleyici_header):
    r = client.get("/sistem/saglik", headers=izleyici_header)
    assert r.status_code == 200, r.text
    uyarilar = r.json()["guvenlik_uyarilari"]
    # Test paketi PTS_LICENSE_SECRET'ı conftest.py'de ZATEN ayarlıyor (bkz.
    # test_suiti_icin_sabit_lisans_secret) -- yani bu bayrak burada True olmalı;
    # PTS_KAMERA_ANAHTARI ve PTS_CORS_ORIGINS ise test ortamında hiç ayarlı değil.
    assert uyarilar == {
        "lisans_secret_ayarli_mi": True,
        "kamera_anahtari_ayarli_mi": False,
        "cors_tum_originlere_acik": False,
    }


def test_sistem_sagligi_guvenlik_uyarilari_ortam_degiskenlerine_gore_degisir(client, izleyici_header, monkeypatch):
    monkeypatch.delenv("PTS_LICENSE_SECRET", raising=False)
    monkeypatch.setenv("PTS_KAMERA_ANAHTARI", "test-kamera-anahtari")
    monkeypatch.setenv("PTS_CORS_ORIGINS", "*")
    r = client.get("/sistem/saglik", headers=izleyici_header)
    assert r.status_code == 200, r.text
    assert r.json()["guvenlik_uyarilari"] == {
        "lisans_secret_ayarli_mi": False,
        "kamera_anahtari_ayarli_mi": True,
        "cors_tum_originlere_acik": True,
    }


# ------------------------------------------------------------------
# /sistem/saglik — SQL Server yedeğinin GERÇEKTEN alınıp alınmadığının izlenmesi
# ------------------------------------------------------------------
# Önceden sistem, SQL Server Agent bakım planının çalışıp çalışmadığını hiçbir
# şekilde izlemiyordu — plan hiç kurulmasa bile sessizce fark edilmezdi.

def test_yedek_izleme_ayarlanmamissa_hicbir_davranis_degismez(client, izleyici_header, monkeypatch):
    monkeypatch.delenv("PTS_SQL_YEDEK_KLASORU", raising=False)
    r = client.get("/sistem/saglik", headers=izleyici_header)
    assert r.status_code == 200, r.text
    assert r.json()["yedek"] == {"izleniyor": False, "son_yedek_zamani": None, "yedek_gecikmis": None}


def test_yedek_izleme_ayarliysa_en_yeni_dosyanin_yasini_raporlar(client, izleyici_header, monkeypatch, tmp_path):
    eski = tmp_path / "eski_yedek.bak"
    eski.write_bytes(b"eski")
    yeni = tmp_path / "yeni_yedek.bak"
    yeni.write_bytes(b"yeni")
    gecmis_zaman = (datetime.now() - timedelta(days=10)).timestamp()
    os.utime(eski, (gecmis_zaman, gecmis_zaman))
    # yeni dosya için mtime'ı elle ayarlamıyoruz — az önce yazıldığı için zaten güncel.

    monkeypatch.setenv("PTS_SQL_YEDEK_KLASORU", str(tmp_path))
    r = client.get("/sistem/saglik", headers=izleyici_header)
    assert r.status_code == 200, r.text
    yedek = r.json()["yedek"]
    assert yedek["izleniyor"] is True
    assert yedek["yedek_gecikmis"] is False  # en yeni dosya (yeni_yedek.bak) az önce yazıldı
    assert yedek["son_yedek_zamani"] is not None


def test_yedek_izleme_klasor_bossa_gecikmis_sayilir(client, izleyici_header, monkeypatch, tmp_path):
    monkeypatch.setenv("PTS_SQL_YEDEK_KLASORU", str(tmp_path))  # var ama boş bir klasör
    r = client.get("/sistem/saglik", headers=izleyici_header)
    yedek = r.json()["yedek"]
    assert yedek["izleniyor"] is True
    assert yedek["yedek_gecikmis"] is True
    assert yedek["son_yedek_zamani"] is None


def test_gecersiz_istek_govdesi_tutarli_422_doner(client, yetkili_header):
    """RequestValidationError (422) için küresel hata yakalayıcı, diğer tüm
    hatalarla aynı `{"detail": ...}` gövde biçimini kullanmalı ve ek olarak
    ham pydantic hata listesini de (`hatalar`) içermeli."""
    r = client.post(
        "/kara-listesi",
        json={"plaka_no": "';--", "sebep": "test"},  # yalnızca geçersiz karakter -> plaka_normalize ValueError fırlatır
        headers=yetkili_header,
    )
    assert r.status_code == 422, r.text
    govde = r.json()
    assert "detail" in govde
    assert "hatalar" in govde
    assert isinstance(govde["hatalar"], list)


# ------------------------------------------------------------------
# /kayitlar, /olaylar, /alarmlar, /sakin/gecmisim — limit/offset sınırları
# ve geçersiz tarih filtresi doğrulaması (2026-09-20)
# ------------------------------------------------------------------
# KÖK NEDEN (geniş kapsamlı denetimde tespit edildi): `limit` parametreleri
# `.limit(min(limit, 500))` deseniyle "üst sınıra kırpılıyordu" -- ama SQLite
# (ve bazı motorlar) NEGATİF bir LIMIT'i "sınırsız" sayıyor: `min(-1, 500)`
# hâlâ `-1`. Yani `?limit=-1` göndererek 500 satırlık üst sınır TAMAMEN
# atlatılabiliyordu (en düşük yetkili "izleyici" rolü dahil) -- KVKK
# kapsamındaki plaka/kişi verilerinin toplu dökümü anlamına gelirdi. Artık
# `Query(ge=1, le=500)` bu değeri isteğin FastAPI'ye ulaştığı anda (422 ile)
# reddediyor. Ayrıca `baslangic`/`bitis` için geçersiz bir tarih artık düz bir
# 500 yerine açık bir 400 döndürüyor.

@pytest.mark.parametrize("uc,parametre", [
    ("/kayitlar", "limit"), ("/kayitlar", "offset"),
    ("/kayitlar/sayfa-bilgisi", "limit"),
    ("/olaylar", "limit"), ("/olaylar", "since_id"),
    ("/alarmlar", "limit"),
])
def test_negatif_limit_offset_artik_422_ile_reddedilir(client, izleyici_header, uc, parametre):
    r = client.get(uc, params={parametre: -1}, headers=izleyici_header)
    assert r.status_code == 422, r.text


def test_asiri_buyuk_limit_500e_kirpilmiyor_422_doner(client, izleyici_header):
    """Önceden `min(limit, 500)` ile sessizce 500'e kırpılıyordu; artık
    açıkça reddediliyor -- istemci gerçek üst sınırı görüp buna göre
    sayfalama yapabilsin diye (sessizce farklı bir sonuç dönmek yerine)."""
    r = client.get("/kayitlar", params={"limit": 100000}, headers=izleyici_header)
    assert r.status_code == 422, r.text


def test_sakin_gecmisim_negatif_limit_422_doner(client, sakin_header):
    r = client.get("/sakin/gecmisim", params={"limit": -5}, headers=sakin_header)
    assert r.status_code == 422, r.text


@pytest.mark.parametrize("uc", ["/kayitlar", "/kayitlar/sayfa-bilgisi"])
@pytest.mark.parametrize("alan", ["baslangic", "bitis"])
def test_gecersiz_tarih_filtresi_500_yerine_400_doner(client, izleyici_header, uc, alan):
    r = client.get(uc, params={alan: "bu-bir-tarih-degil"}, headers=izleyici_header)
    assert r.status_code == 400, r.text
    assert alan in r.json()["detail"]


def test_gecerli_tarih_filtresi_normal_calisir(client, izleyici_header):
    """Yukarıdaki 400 doğrulamasının, GEÇERLİ bir ISO tarihini yanlışlıkla
    reddetmediğinden emin olmak için (regresyona karşı)."""
    r = client.get("/kayitlar", params={"baslangic": "2026-01-01", "bitis": "2026-12-31"}, headers=izleyici_header)
    assert r.status_code == 200, r.text


# ------------------------------------------------------------------
# Kayıt filtrelemede SAAT aralığı (2026-09-21 kullanıcı talebi: "kayıt
# filtreleme kısmına saat seçme özelliği de ekler misin, sadece tarih var,
# belirli saat aralıklarıyla da kayıt almam gerekiyor") -- baslangic/bitis
# artık yalnızca tarih değil, saat de içerebiliyor (örn.
# "2026-01-15T14:00:00", frontend'deki yeni saat seçiciyle üretilir, bkz.
# app.js::_tarihSaatDegeriOlustur). Bkz. main.py::_bitis_tarih_filtresi_sinirini_hesapla.
# ------------------------------------------------------------------

def test_kayitlar_saat_araligi_ile_filtrelenebilir(client, yetkili_header):
    """AYNI güne düşen ama farklı SAATLERDE olan üç kaydın, dar bir saat
    aralığı filtresiyle doğru ayrıştırıldığını doğrular."""
    from backend.database import SessionLocal
    from backend import models

    on_ek = "34 SAAT01"  # bu teste özgü, başka hiçbir testle çakışmayacak benzersiz plaka öneki
    db = SessionLocal()
    try:
        db.add_all([
            models.Kayit(plaka_no=f"{on_ek} SABAH", kamera_id="TEST", tarih_saat=datetime(2026, 1, 15, 8, 0, 0)),
            models.Kayit(plaka_no=f"{on_ek} OGLE", kamera_id="TEST", tarih_saat=datetime(2026, 1, 15, 14, 0, 0)),
            models.Kayit(plaka_no=f"{on_ek} AKSAM", kamera_id="TEST", tarih_saat=datetime(2026, 1, 15, 20, 0, 0)),
        ])
        db.commit()
    finally:
        db.close()

    r = client.get("/kayitlar", params={
        "plaka": on_ek,
        "baslangic": "2026-01-15T12:00:00",
        "bitis": "2026-01-15T18:00:00",
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text
    plakalar = {k["plaka_no"] for k in r.json()}
    assert plakalar == {f"{on_ek} OGLE"}, (
        f"12:00-18:00 saat aralığı filtresi yalnızca öğle kaydını döndürmeliydi, dönen: {plakalar}"
    )


def test_kayitlar_bitis_saatsiz_gunun_sonuna_kadar_kapsar(client, yetkili_header):
    """Regresyon: `bitis` saat İÇERMEDİĞİNDE (salt tarih, örn. '2026-01-16')
    eski davranış korunmalı -- o günün TAMAMI (23:59:59'a kadar) dahil
    edilmeli. Saat seçme özelliği eklenirken bu davranış YANLIŞLIKLA
    bozulmamalı."""
    from backend.database import SessionLocal
    from backend import models

    on_ek = "34 SAAT02"
    db = SessionLocal()
    try:
        db.add(models.Kayit(plaka_no=on_ek, kamera_id="TEST", tarih_saat=datetime(2026, 1, 16, 23, 59, 59)))
        db.commit()
    finally:
        db.close()

    r = client.get(
        "/kayitlar", params={"plaka": on_ek, "baslangic": "2026-01-16", "bitis": "2026-01-16"}, headers=yetkili_header,
    )
    assert r.status_code == 200, r.text
    plakalar = {k["plaka_no"] for k in r.json()}
    assert plakalar == {on_ek}, (
        f"Saat içermeyen bitis='2026-01-16' filtresi, o günün 23:59:59'undaki kaydı da KAPSAMALIYDI, dönen: {plakalar}"
    )


def test_kayitlar_sayfa_bilgisi_saat_araligini_kayitlar_ile_tutarli_sayar(client, yetkili_header):
    """/kayitlar ve /kayitlar/sayfa-bilgisi AYNI saat aralığı filtresiyle
    çağrıldığında TUTARLI bir 'toplam' sayısı döndürmeli -- ikisi ayrı ayrı
    aynı üst sınır mantığını (bkz. main.py::_bitis_tarih_filtresi_sinirini_hesapla)
    kullanıyor, bu regresyona karşı doğrulanıyor."""
    from backend.database import SessionLocal
    from backend import models

    on_ek = "34 SAAT03"
    db = SessionLocal()
    try:
        db.add_all([
            models.Kayit(plaka_no=f"{on_ek} A", kamera_id="TEST", tarih_saat=datetime(2026, 1, 17, 9, 0, 0)),
            models.Kayit(plaka_no=f"{on_ek} B", kamera_id="TEST", tarih_saat=datetime(2026, 1, 17, 11, 0, 0)),
        ])
        db.commit()
    finally:
        db.close()

    ortak_params = {"plaka": on_ek, "baslangic": "2026-01-17T08:00:00", "bitis": "2026-01-17T12:00:00"}
    r1 = client.get("/kayitlar", params=ortak_params, headers=yetkili_header)
    r2 = client.get("/kayitlar/sayfa-bilgisi", params=ortak_params, headers=yetkili_header)
    assert r1.status_code == 200, r1.text
    assert r2.status_code == 200, r2.text
    assert len(r1.json()) == 2
    assert r2.json()["toplam"] == 2


def test_guvenlik_baslikları_her_yanitta_var(client):
    """2026-09-20: her yanıta clickjacking/MIME-sniffing'e karşı ek bir
    savunma katmanı ekleyen standart güvenlik başlıkları eklendi (bkz.
    main.py::_guvenlik_basliklarini_ekle). Kimlik doğrulaması gerektirmeyen
    bir uç noktada bile (health-check benzeri) bu başlıkların var olduğunu
    doğruluyoruz -- middleware TÜM yanıtları kapsamalı."""
    r = client.get("/")
    assert r.headers["X-Content-Type-Options"] == "nosniff"
    assert r.headers["X-Frame-Options"] == "SAMEORIGIN"
    assert r.headers["Referrer-Policy"] == "strict-origin-when-cross-origin"
    assert "max-age=31536000" in r.headers["Strict-Transport-Security"]


def test_api_dokumantasyonu_varsayilan_olarak_kapali(client):
    """2026-09-20 ("daha profesyonel neler yapabilirsin" denetimi): Swagger UI
    (/docs), ReDoc (/redoc) ve ham OpenAPI şeması (/openapi.json) FastAPI'de
    varsayılan olarak açık ve kimliksizdi -- kimlik doğrulaması olmayan bir
    saldırı yüzeyi haritası sunuyordu. Test ortamı PTS_API_DOKUMANTASYONU_AC'ı
    AYARLAMAZ, bu yüzden bu üç uç nokta burada 404 dönmeli (bkz.
    main.py::_api_dokumantasyonu_acik_mi)."""
    assert client.get("/docs").status_code == 404
    assert client.get("/redoc").status_code == 404
    assert client.get("/openapi.json").status_code == 404


def test_sistem_sagligi_surum_app_version_ile_tek_kaynaktan_geliyor(client, izleyici_header):
    """2026-09-20: /sistem/saglik'teki "surum" alanı önceden "2.0" olarak
    AYRICA sabitlenmişti -- `app = FastAPI(..., version="2.0")`'daki AYNI
    değerin bağımsız bir kopyası. Artık tek kaynak `app.version` (bkz.
    main.py::sistem_sagligi)."""
    r = client.get("/sistem/saglik", headers=izleyici_header)
    assert r.status_code == 200, r.text
    assert r.json()["surum"] == pts_main.app.version


# ------------------------------------------------------------------
# GET /olaylar, GET /alarmlar — temel işlevsellik (2026-09-20)
# ------------------------------------------------------------------
# Denetimde bu iki uç noktanın HİÇ test kapsamında olmadığı tespit edildi
# (yalnızca yukarıdaki negatif-limit doğrulaması dolaylı olarak dokunuyordu).

def test_olaylar_girissiz_401_doner(client):
    r = client.get("/olaylar")
    assert r.status_code == 401, r.text


def test_olaylar_since_id_ile_sonraki_kayitlari_getirir(client, yetkili_header):
    r1 = client.post("/kayitlar", json={"plaka_no": "34 OLY 01", "kamera_id": "TEST", "yon": "giris"}, headers=yetkili_header)
    ilk_id = r1.json()["id"]
    r2 = client.post("/kayitlar", json={"plaka_no": "34 OLY 02", "kamera_id": "TEST", "yon": "giris"}, headers=yetkili_header)
    ikinci_id = r2.json()["id"]

    r = client.get("/olaylar", params={"since_id": ilk_id}, headers=yetkili_header)
    assert r.status_code == 200, r.text
    idler = {k["id"] for k in r.json()}
    assert ikinci_id in idler
    assert ilk_id not in idler


def test_alarmlar_girissiz_401_doner(client):
    r = client.get("/alarmlar")
    assert r.status_code == 401, r.text


def test_alarmlar_kara_liste_gecisinde_olusur_ve_sadece_acik_filtresi_calisir(client, yetkili_header):
    plaka = "34 ALM 01"
    r0 = client.post("/kara-listesi", json={"plaka_no": plaka, "sebep": "test"}, headers=yetkili_header)
    assert r0.status_code == 200, r0.text
    r1 = client.post("/kayitlar", json={"plaka_no": plaka, "kamera_id": "TEST", "yon": "giris"}, headers=yetkili_header)
    assert r1.status_code == 200, r1.text
    assert r1.json()["yetki_durumu"] == "kara_liste"

    r2 = client.get("/alarmlar", headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    alarmlar = [a for a in r2.json() if a["plaka_no"] == plaka]
    assert len(alarmlar) >= 1
    assert alarmlar[0]["alarm_tipi"] == "kara_liste"
    assert alarmlar[0]["okundu"] is False

    r3 = client.patch(f"/alarmlar/{alarmlar[0]['id']}/okundu", headers=yetkili_header)
    assert r3.status_code == 200, r3.text

    r4 = client.get("/alarmlar", params={"sadece_acik": True}, headers=yetkili_header)
    kalan_idler = {a["id"] for a in r4.json()}
    assert alarmlar[0]["id"] not in kalan_idler, "okundu=True işaretlenen alarm 'sadece_acik' filtresinde hâlâ görünüyor"


def test_beklenmeyen_hata_loglanir_ve_tutarli_500_doner(client, caplog, izleyici_header):
    """Bilerek fırlatılmamış (HTTPException olmayan) bir hata; istemciye
    traceback sızdırmadan tutarlı bir JSON gövdesiyle dönmeli VE sunucu
    tarafında tam iz düşümüyle loglanmalı (bkz. main.py::_beklenmeyen_hata_yakalayici) —
    aksi halde bu sınıftaki hatalar `/sistem/loglar` üzerinden hiç görülemez.

    NOT (2026-09-20): `/sistem/saglik` artık personel girişi istiyor (bkz.
    test_sistem_sagligi_girissiz_401_doner), bu yüzden burada da
    `izleyici_header` gönderiliyor -- aksi halde bu test kendi patch'lediği
    fonksiyona hiç ulaşmadan 401'e düşerdi."""

    def _patlayan_fonksiyon():
        raise RuntimeError("kaçınılmaz test hatası")

    with caplog.at_level("ERROR", logger="pts"):
        with pytest.MonkeyPatch.context() as mp:
            mp.setattr(pts_main, "_son_yedek_bilgisini_al", _patlayan_fonksiyon)
            r = client.get("/sistem/saglik", headers=izleyici_header)

    assert r.status_code == 500, r.text
    govde = r.json()
    assert govde == {
        "detail": "Sunucuda beklenmeyen bir hata oluştu. Lütfen tekrar deneyin veya sistem yöneticisine bildirin."
    }
    assert "kaçınılmaz test hatası" not in r.text  # istemciye traceback sızmamalı
    assert "Beklenmeyen hata" in caplog.text  # ama sunucu logunda İZİ olmalı
    assert "kaçınılmaz test hatası" in caplog.text  # traceback sunucu logunda GÖRÜNÜR olmalı


def test_gorsel_izleme_ayarlanmamissa_pasif_raporlanir(client, izleyici_header):
    """PTS_GORSEL_IZLEME_DIZINI test ortamında ayarlı değil, dolayısıyla
    _klasor_izleyici hiç başlatılmamış olmalı (bkz. main.py::
    _klasor_izlemeyi_baslat_gerekirse) — varsayılan/eski davranış korunur."""
    r = client.get("/sistem/saglik", headers=izleyici_header)
    assert r.status_code == 200, r.text
    assert r.json()["gorsel_izleme"] == {"aktif": False, "klasor": None}


def test_gorsel_izleme_aktifken_klasoruyle_birlikte_raporlanir(client, izleyici_header, monkeypatch):
    """Klasör izleyici çalışıyorken /sistem/saglik bunu ve hangi klasörü
    izlediğini bildirmeli (panelin Sistem sekmesindeki satır buna dayanır)."""

    class _SahteKlasorIzleyici:
        calisiyor = True
        kok_klasor = "/tmp/ornek-izleme-klasoru"

    monkeypatch.setattr(pts_main, "_klasor_izleyici", _SahteKlasorIzleyici())
    r = client.get("/sistem/saglik", headers=izleyici_header)
    assert r.status_code == 200, r.text
    assert r.json()["gorsel_izleme"] == {"aktif": True, "klasor": "/tmp/ornek-izleme-klasoru"}


def test_anpr_dedektor_esigi_kutuphane_yoksa_none_doner(client, izleyici_header, monkeypatch):
    """Kamera kütüphaneleri (cv2/fast-alpr bağımlılıkları) kurulu değilse
    dedektör eşiği bilgisi hiç sorulmaz, sade None raporlanır."""
    monkeypatch.setattr(pts_main, "_CAM_LIBS", False)
    r = client.get("/sistem/saglik", headers=izleyici_header)
    assert r.status_code == 200, r.text
    assert r.json()["anpr_dedektor_esigi"] is None


def test_anpr_dedektor_esigi_motor_olusunca_fiili_deger_raporlanir(client, izleyici_header, monkeypatch):
    """PTS_ANPR_DETECTOR_ESIGI ortam değişkeninin panelden ayarlanan
    'Min. plaka tanıma güveni' ile karıştırılması geçmişte kullanıcı
    karışıklığına yol açmıştı: bu test, motor bir kez oluşturulduktan sonra
    /sistem/saglik'in FİİLEN uygulanan eşiği ve kaynağını (ortam değişkeni mi,
    kütüphane varsayılanı mı) doğru yansıttığını doğrular (bkz.
    camera_reader.py::dedektor_esigi_bilgisi, anpr_engine.py)."""
    from backend import camera_reader as cr

    monkeypatch.setattr(pts_main, "_CAM_LIBS", True)
    monkeypatch.setattr(
        cr, "dedektor_esigi_bilgisi",
        lambda: {
            "esik": 0.25, "kaynak": "PTS_ANPR_DETECTOR_ESIGI ortam değişkeni ('0.25')",
            "model": "yolo-v9-s-608-license-plate-end2end",
            "model_kaynagi": "PTS_ANPR_DETECTOR_MODEL ortam değişkeni",
        },
    )
    r = client.get("/sistem/saglik", headers=izleyici_header)
    assert r.status_code == 200, r.text
    assert r.json()["anpr_dedektor_esigi"] == {
        "esik": 0.25, "kaynak": "PTS_ANPR_DETECTOR_ESIGI ortam değişkeni ('0.25')",
        "model": "yolo-v9-s-608-license-plate-end2end",
        "model_kaynagi": "PTS_ANPR_DETECTOR_MODEL ortam değişkeni",
    }


def test_dogruluk_testi_izleyici_erisemez(client, izleyici_header):
    """Sunucudaki dosya sistemini okuyan bir işlem olduğu için izleyici
    rolüne açık değil -- bkz. main.py::dogruluk_testi_calistir."""
    r = client.post("/sistem/dogruluk-testi", json={"klasor": "/tmp/herhangi"}, headers=izleyici_header)
    assert r.status_code == 403


def test_dogruluk_testi_kutuphane_yoksa_400_doner(client, operator_header, monkeypatch):
    monkeypatch.setattr(pts_main, "_CAM_LIBS", False)
    r = client.post("/sistem/dogruluk-testi", json={"klasor": "/tmp/herhangi"}, headers=operator_header)
    assert r.status_code == 400


def test_dogruluk_testi_operator_calistirabilir_ve_sonucu_doner(client, operator_header, monkeypatch):
    """toplu_dogruluk_testi() gerçekten çağrılıyor mu, verilen parametreler
    (klasor/min_guven_skoru/kontrast_iyilestir) doğru iletiliyor mu ve
    dönen sonuç istemciye olduğu gibi ulaşıyor mu -- bunu doğrular."""
    monkeypatch.setattr(pts_main, "_CAM_LIBS", True)

    cagrilar = []

    def _sahte_test(klasor, min_guven_skoru, kontrast_iyilestir):
        cagrilar.append((klasor, min_guven_skoru, kontrast_iyilestir))
        return {
            "toplam": 2, "dogru": 1, "yanlis": 0, "esik_altinda": 0,
            "tespit_edilemedi": 1, "etiketlenemedi": 0, "dogruluk_orani": 0.5,
            "detaylar": [],
        }

    from backend import camera_reader as cr
    monkeypatch.setattr(cr, "toplu_dogruluk_testi", _sahte_test)

    r = client.post(
        "/sistem/dogruluk-testi",
        json={"klasor": "/tmp/etiketli-fotograflar", "min_guven_skoru": 0.25, "kontrast_iyilestir": True},
        headers=operator_header,
    )
    assert r.status_code == 200, r.text
    assert r.json()["dogruluk_orani"] == 0.5
    assert cagrilar == [("/tmp/etiketli-fotograflar", 0.25, True)]


def test_dogruluk_testi_klasor_bulunamazsa_400_doner(client, operator_header, monkeypatch):
    monkeypatch.setattr(pts_main, "_CAM_LIBS", True)

    def _hata_firlat(klasor, min_guven_skoru, kontrast_iyilestir):
        raise ValueError(f"Klasör bulunamadı: {klasor}")

    from backend import camera_reader as cr
    monkeypatch.setattr(cr, "toplu_dogruluk_testi", _hata_firlat)

    r = client.post("/sistem/dogruluk-testi", json={"klasor": "/olmayan"}, headers=operator_header)
    assert r.status_code == 400


# ------------------------------------------------------------------
# Geçiş kaydı tam düzenleme + manuel kayıt notu (bkz. main.py::kayit_duzenle,
# kayit_ekle_manuel; "Plaka Analizi" ekranındaki manuel kayıt/düzenleme
# özellikleri bu uç noktaları kullanır)
# ------------------------------------------------------------------

def test_kayit_manuel_not_metni_ve_manuel_giris_bayragi_kaydedilir(client, operator_header):
    """'Plaka Analizi' ekranındaki 'Manuel Kayıt Ekle' formu bu uca not_metni
    ile POST atıyor -- kayıt hem manuel_giris=True hem de girilen notla
    dönmeli (görsel/güven skoru olmadan)."""
    r = client.post("/kayitlar", json={
        "plaka_no": "34 MNL 01", "kamera_id": "PANEL-MANUEL", "yon": "giris",
        "not_metni": "teslimat aracı, güvenlik onayıyla alındı",
    }, headers=operator_header)
    assert r.status_code == 200, r.text
    veri = r.json()
    assert veri["manuel_giris"] is True
    assert veri["not_metni"] == "teslimat aracı, güvenlik onayıyla alındı"
    assert veri["duzenleyen"] is None  # manuel eklemek "düzenleme" sayılmaz


def test_kayit_duzenle_plaka_durum_ve_not_guncellenir_ve_denetim_izi_tutulur(client, operator_header):
    r = client.post("/kayitlar", json={"plaka_no": "34 DZL 02", "kamera_id": "TEST", "yon": "giris"},
                     headers=operator_header)
    assert r.status_code == 200, r.text
    kayit_id = r.json()["id"]
    assert r.json()["duzenleyen"] is None

    r2 = client.patch(f"/kayitlar/{kayit_id}", json={
        "plaka_no": "34 dzl 02",  # küçük harf + boşluksuz -- normalize edilmeli
        "yetki_durumu": "yetkili",
        "not_metni": "operatör tarafından manuel olarak yetkilendirildi",
    }, headers=operator_header)
    assert r2.status_code == 200, r2.text
    veri = r2.json()
    assert veri["plaka_no"] == "34 DZL 02"
    assert veri["yetki_durumu"] == "yetkili"
    assert veri["not_metni"] == "operatör tarafından manuel olarak yetkilendirildi"
    assert veri["duzenleyen"] == "rbac-operator"
    assert veri["duzenleme_tarihi"] is not None


def test_kayit_duzenle_gecersiz_yetki_durumu_400_doner(client, operator_header):
    r = client.post("/kayitlar", json={"plaka_no": "34 DZL 03", "kamera_id": "TEST", "yon": "giris"},
                     headers=operator_header)
    kayit_id = r.json()["id"]
    r2 = client.patch(f"/kayitlar/{kayit_id}", json={"yetki_durumu": "gecersiz-deger"}, headers=operator_header)
    assert r2.status_code == 400


def test_kayit_duzenle_kisi_eslestirme_atanir_ve_temizlenir(client, operator_header):
    rk = client.post("/kisiler", json={
        "ad_soyad": "Düzenleme Testi Kişi", "plaka_no": "34 DZL 04", "tip": "personel",
    }, headers=operator_header)
    assert rk.status_code == 200, rk.text
    kisi_id = rk.json()["id"]

    r = client.post("/kayitlar", json={"plaka_no": "06 BSK 05", "kamera_id": "TEST", "yon": "giris"},
                     headers=operator_header)
    kayit_id = r.json()["id"]

    r2 = client.patch(f"/kayitlar/{kayit_id}", json={"kisi_id": kisi_id}, headers=operator_header)
    assert r2.status_code == 200, r2.text
    assert r2.json()["kisi_id"] == kisi_id
    assert r2.json()["kisi_tip_anlik"] == "personel"

    r3 = client.patch(f"/kayitlar/{kayit_id}", json={"kisi_id_temizle": True}, headers=operator_header)
    assert r3.status_code == 200, r3.text
    assert r3.json()["kisi_id"] is None


def test_kayit_duzenle_izleyici_yetkisiz_403_doner(client, operator_header, izleyici_header):
    r = client.post("/kayitlar", json={"plaka_no": "34 DZL 06", "kamera_id": "TEST", "yon": "giris"},
                     headers=operator_header)
    kayit_id = r.json()["id"]
    r2 = client.patch(f"/kayitlar/{kayit_id}", json={"not_metni": "izleyici bunu yapamamalı"},
                       headers=izleyici_header)
    assert r2.status_code == 403


def test_kayit_sil_yalnizca_yonetici_operator_403_doner(client, operator_header):
    r = client.post("/kayitlar", json={"plaka_no": "34 DZL 07", "kamera_id": "TEST", "yon": "giris"},
                     headers=operator_header)
    kayit_id = r.json()["id"]
    r2 = client.delete(f"/kayitlar/{kayit_id}", headers=operator_header)
    assert r2.status_code == 403


def test_kayit_duzenle_bulunamayan_kayit_404_doner(client, operator_header):
    r = client.patch("/kayitlar/999999", json={"not_metni": "yok"}, headers=operator_header)
    assert r.status_code == 404


def test_manuel_giris_null_olan_eski_kayit_500_patlamiyor(client, operator_header):
    """REGRESYON (2026-09-17, gerçek kullanıcı ortamında bulundu): SQL
    Server'da ALTER TABLE ... ADD manuel_giris BIT DEFAULT 0, "WITH VALUES"
    açıkça verilmedikçe TABLODA HALİHAZIRDA VAR OLAN satırları NULL bırakıyor
    (SQLite'ın aksine). `manuel_giris` alanı `bool` (Optional değil) olduğu
    için bu, GET /kayitlar gibi uçlarda ResponseValidationError ile 500'e yol
    açtı. Burada bunu, ORM'i atlayıp doğrudan SQL ile manuel_giris'i NULL
    bırakan bir kayıt ekleyerek simüle ediyoruz -- normal POST /kayitlar akışı
    Python tarafında zaten default=False uyguladığı için bunu üretemez, gerçek
    hata yalnızca ham SQL/ALTER TABLE seviyesinde ortaya çıkıyor."""
    from backend.database import engine
    from sqlalchemy import text as sqltext
    with engine.connect() as conn:
        conn.execute(sqltext(
            "INSERT INTO plaka_kayitlari (plaka_no, kamera_id, yon, yetki_durumu, manuel_giris) "
            "VALUES ('34 NULLTEST 09', 'TEST', 'giris', 'bilinmiyor', NULL)"
        ))
        conn.commit()

    r = client.get("/kayitlar", params={"plaka": "NULLTEST"}, headers=operator_header)
    assert r.status_code == 200, r.text
    kayitlar = r.json()
    assert len(kayitlar) == 1
    assert kayitlar[0]["manuel_giris"] is False


def test_plaka_analiz_yeni_alanlari_dondurur(client, operator_header):
    """Plaka Analizi ekranının görsel/doğrulama/not/manuel-giriş/düzenleme
    denetim bilgilerini gösterebilmesi için bu alanların analiz uç
    noktasından da gelmesi gerekiyor (bkz. main.py::plaka_analiz)."""
    r = client.post("/kayitlar", json={
        "plaka_no": "34 ANLZ 08", "kamera_id": "PANEL-MANUEL", "yon": "giris",
        "not_metni": "analiz ekranı testi",
    }, headers=operator_header)
    assert r.status_code == 200, r.text

    r2 = client.get("/kayitlar/analiz/34 ANLZ 08", headers=operator_header)
    assert r2.status_code == 200, r2.text
    kayit = r2.json()["son_kayitlar"][0]
    for alan in ("id", "plaka_no", "goruntu_yolu", "guven_skoru", "dogrulama_kare_sayisi",
                 "farkli_okuma_sayisi",
                 "not_metni", "manuel_giris", "kisi_id", "duzenleyen", "duzenleme_tarihi"):
        assert alan in kayit, f"'{alan}' alanı /kayitlar/analiz yanıtında eksik"
    assert kayit["not_metni"] == "analiz ekranı testi"
    assert kayit["manuel_giris"] is True


def test_plaka_analiz_toplam_gecis_dogru_sayilir(client, operator_header):
    """GÜVENLİK/DOĞRULUK KÖK NEDEN REGRESYON TESTİ (2026-09-17): kullanıcı,
    panelde bir plakaya tıklayınca açılan "Plaka Analizi" penceresinin,
    plakanın GERÇEKTEN var olan geçişlerine rağmen her zaman "Toplam Geçiş:
    0 / Kayıt yok" gösterdiğini bildirdi. Kök neden: main.py::plaka_analiz
    boşluksuz normalize edilmiş bir değerle, boşluklu SAKLANAN
    `Kayit.plaka_no` arasında DOĞRUDAN SQL eşitliği kullanıyordu -- iki taraf
    da farklı biçimde olduğu için sorgu HİÇBİR ZAMAN eşleşmiyordu."""
    for _ in range(3):
        r = client.post("/kayitlar", json={"plaka_no": "34 TGS 09", "kamera_id": "TEST", "yon": "giris"},
                         headers=operator_header)
        assert r.status_code == 200, r.text

    r2 = client.get("/kayitlar/analiz/34 TGS 09", headers=operator_header)
    assert r2.status_code == 200, r2.text
    veri = r2.json()
    assert veri["toplam_gecis"] == 3, "Plaka Analizi gerçek geçiş sayısını göstermeliydi (0 değil)"
    assert veri["son_gecis"] is not None
    assert len(veri["son_kayitlar"]) == 3
    assert veri["plaka_no"] == "34 TGS 09", (
        "Görüntülenen plaka, kayıtlardaki OKUNABİLİR (boşluklu) biçimde dönmeli, "
        "normalize edilmiş/boşluksuz hali değil"
    )


def test_plaka_analiz_kara_liste_durumunu_dogru_gosterir(client, yetkili_header):
    """Aynı kök nedenin bir başka belirtisi: Plaka Analizi'ndeki "Kara Liste"
    rozeti, plaka gerçekten kara listede olsa bile her zaman "Temiz"
    gösteriyordu (bkz. yukarıdaki not)."""
    r = client.post("/kara-listesi", json={"plaka_no": "34 PAK 10", "sebep": "test"}, headers=yetkili_header)
    assert r.status_code == 200, r.text
    r2 = client.get("/kayitlar/analiz/34 PAK 10", headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    assert r2.json()["kara_listesinde"] is True


# ------------------------------------------------------------------
# "Son kullanılan not" önerisi -- her gece 23:59'da sıfırlanır (2026-09-17, devam)
# ------------------------------------------------------------------
# Kullanıcı senaryosu: yetkisiz bir araca (örn. bir kargo aracına, "PTT
# Kargo" örneğindeki gibi) panelden elle not eklendiğinde, aynı plaka aynı
# gün tekrar geldiğinde görevli notu yeniden yazmasın diye bir sonraki not
# kutusuna ÖNERİ olarak sunulur (bkz. main.py::_SON_NOT_ONBELLEGI). Kullanıcı
# açıkça bu önerinin bir sonraki güne HİÇ taşınmamasını istedi -- gerçek
# denetim kaydı olan Kayit.not_metni'nden TAMAMEN AYRI, geçici bir önbellek.

def test_son_not_onerisi_kayitsiz_plaka_icin_bos_doner(client, yetkili_header):
    r = client.get("/kayitlar/son-not", params={"plaka": "34 YOK 99"}, headers=yetkili_header)
    assert r.status_code == 200, r.text
    assert r.json()["not_metni"] is None


def test_manuel_kayit_notu_son_not_onerisine_yansir(client, operator_header, yetkili_header):
    """Kullanıcının verdiği örneğin birebir tekrarı: yetkisiz bir araca
    (kargo aracı) "PTT Kargo" notuyla manuel bir giriş kaydı eklenir; aynı
    plaka için son-not önerisi bu notu döndürmeli."""
    r = client.post("/kayitlar", json={
        "plaka_no": "34 PTT 01", "kamera_id": "PANEL-MANUEL", "yon": "giris",
        "guven_skoru": None, "not_metni": "PTT Kargo",
    }, headers=operator_header)
    assert r.status_code == 200, r.text
    assert r.json()["yetki_durumu"] in ("yetkisiz", "bilinmiyor")

    r2 = client.get("/kayitlar/son-not", params={"plaka": "34 PTT 01"}, headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    assert r2.json()["not_metni"] == "PTT Kargo"

    # Aynı gün içinde aynı aracın çıkış kaydı ayrıca eklense bile (örn. saat
    # 13:00), öneri hâlâ aynı notu göstermeli -- kalıcı kayda hiç dokunmadan.
    r3 = client.post("/kayitlar", json={
        "plaka_no": "34 PTT 01", "kamera_id": "PANEL-MANUEL", "yon": "cikis",
    }, headers=operator_header)
    assert r3.status_code == 200, r3.text
    r4 = client.get("/kayitlar/son-not", params={"plaka": "34 PTT 01"}, headers=yetkili_header)
    assert r4.json()["not_metni"] == "PTT Kargo"


def test_kayit_duzenle_notu_da_son_not_onerisine_yansir(client, operator_header, yetkili_header):
    """Not yalnızca yeni kayıt eklerken değil, mevcut bir kaydı (Kayıt
    Düzenle / Ziyaretçi Girişi -- ikisi de PATCH /kayitlar/{id} kullanır)
    düzenlerken eklenirse de önbelleğe yansımalı."""
    r = client.post("/kayitlar", json={"plaka_no": "34 DZL 02", "kamera_id": "TEST", "yon": "giris"},
                     headers=operator_header)
    assert r.status_code == 200, r.text
    kayit_id = r.json()["id"]

    r2 = client.patch(f"/kayitlar/{kayit_id}", json={"not_metni": "Misafir - B Blok"}, headers=operator_header)
    assert r2.status_code == 200, r2.text

    r3 = client.get("/kayitlar/son-not", params={"plaka": "34 DZL 02"}, headers=yetkili_header)
    assert r3.status_code == 200, r3.text
    assert r3.json()["not_metni"] == "Misafir - B Blok"


def test_plaka_analiz_son_not_onerisini_dondurur(client, operator_header):
    """Plaka Analizi ekranındaki "Manuel Kayıt Ekle" not kutusunun
    doldurulabilmesi için, analiz uç noktası da öneriyi kendi yanıtında
    döndürmeli (bkz. main.py::plaka_analiz'deki son_not_onerisi alanı)."""
    r = client.post("/kayitlar", json={
        "plaka_no": "34 SNO 03", "kamera_id": "PANEL-MANUEL", "yon": "giris", "not_metni": "PTT Kargo",
    }, headers=operator_header)
    assert r.status_code == 200, r.text

    r2 = client.get("/kayitlar/analiz/34 SNO 03", headers=operator_header)
    assert r2.status_code == 200, r2.text
    assert r2.json()["son_not_onerisi"] == "PTT Kargo"


def test_son_not_onerisi_gun_degisince_sifirlanir(client, yetkili_header, operator_header):
    """Kullanıcının açık talebi: bir önceki günün notu bir sonraki güne HİÇ
    taşınmamalı. Arka plan temizlik döngüsünün saat 23:59'u beklemesini
    test içinde simüle etmek yerine (bu, gerçek zamanı beklemeyi gerektirir),
    önbellekteki tarihi doğrudan DÜN olarak ayarlayıp pasif sona erme
    mantığının (_son_not_oku) çalıştığını doğruluyoruz -- aktif 23:59
    döngüsü yalnızca belleği erkenden boşaltmak içindir, doğruluk buna değil
    bu pasif kontrole dayanır (bkz. main.py::_son_not_oku'nun docstring'i)."""
    r = client.post("/kayitlar", json={
        "plaka_no": "34 DUN 04", "kamera_id": "PANEL-MANUEL", "yon": "giris", "not_metni": "PTT Kargo",
    }, headers=operator_header)
    assert r.status_code == 200, r.text

    r2 = client.get("/kayitlar/son-not", params={"plaka": "34 DUN 04"}, headers=yetkili_header)
    assert r2.json()["not_metni"] == "PTT Kargo", "Önce normal şekilde önbelleğe düşmeli"

    # Önbellekteki tarihi doğrudan "dün"e çekerek gün değişimini simüle et.
    hedef = pts_main._plaka_normalize("34 DUN 04")
    assert hedef in pts_main._SON_NOT_ONBELLEGI
    pts_main._SON_NOT_ONBELLEGI[hedef]["tarih"] = (datetime.now() - timedelta(days=1)).date()

    r3 = client.get("/kayitlar/son-not", params={"plaka": "34 DUN 04"}, headers=yetkili_header)
    assert r3.status_code == 200, r3.text
    assert r3.json()["not_metni"] is None, "Dünden kalan not önerisi bugüne taşınmamalıydı"
    # Pasif sıfırlama, önbellekten de fiilen silmeli (bellek şişmesin diye).
    assert hedef not in pts_main._SON_NOT_ONBELLEGI


# ------------------------------------------------------------------
# Kamera yön değiştirme + tespit alanı (ROI) sınırlama
# ------------------------------------------------------------------
# Kök neden (2026-09-17, gerçek kullanıcı ortamında bulundu): giriş ve çıkış
# kameralarının açıları birbirinin şeridini de görecek şekilde örtüşünce,
# giriş yapan bir araç çıkış kamerasına da yansıyıp aynı anda hem "giriş" hem
# "çıkış" olarak iki ayrı kayıt oluşturuyordu. Bunu çözmek için: (1) bir
# kameranın yönünü RTSP adresini/parolasını yeniden girmeye gerek kalmadan
# yerinde değiştirebilen PATCH /kameralar/{id}/yon ve (2) her kameraya
# yüzde tabanlı bir tespit alanı (ROI) tanımlayıp bu alanın dışındaki
# tespitleri oy birikimine hiç sokmayan PATCH /kameralar/{id}/roi eklendi.

@pytest.fixture(scope="module")
def roi_test_kamera_id(client, yetkili_header):
    """Bu modüldeki lisans limiti daha önceki testlerde (bkz.
    test_lisans_aktivasyonu_ve_kamera_limiti) 1 olarak aktifleştirilip o tek
    hak da kullanılmış durumda -- bu yüzden burada limiti yükselten YENİ bir
    lisans aktive edip kendi test kameramızı ekliyoruz."""
    anahtar = lisans.uret("ROI Test Site", kamera_limiti=10, gun=30)
    r = client.post("/lisans/aktive-et", json={"anahtar": anahtar}, headers=yetkili_header)
    assert r.status_code == 200, r.text

    r2 = client.post("/kameralar", json={
        "ad": "ROI Test Kamerası", "rtsp_url": "rtsp://127.0.0.1/roitest", "yon": "giris",
    }, headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    return r2.json()["id"]


def test_kamera_yon_degistir_operator_calistirabilir_ve_yon_gunceller(client, operator_header, roi_test_kamera_id):
    r = client.patch(f"/kameralar/{roi_test_kamera_id}/yon", json={"yon": "cikis"}, headers=operator_header)
    assert r.status_code == 200, r.text
    assert r.json()["yon"] == "cikis"

    # Panel/başka testler karışmasın diye eski haline geri alıyoruz.
    r2 = client.patch(f"/kameralar/{roi_test_kamera_id}/yon", json={"yon": "giris"}, headers=operator_header)
    assert r2.status_code == 200
    assert r2.json()["yon"] == "giris"


def test_kamera_yon_degistir_gecersiz_deger_400_doner(client, operator_header, roi_test_kamera_id):
    r = client.patch(f"/kameralar/{roi_test_kamera_id}/yon", json={"yon": "yukari"}, headers=operator_header)
    assert r.status_code == 400


def test_kamera_yon_degistir_bulunamayan_kamera_404_doner(client, operator_header):
    r = client.patch("/kameralar/YOK-BOYLE-BIR-KAMERA/yon", json={"yon": "giris"}, headers=operator_header)
    assert r.status_code == 404


def test_kamera_yon_degistir_izleyici_yetkisiz_403_doner(client, izleyici_header, roi_test_kamera_id):
    r = client.patch(f"/kameralar/{roi_test_kamera_id}/yon", json={"yon": "cikis"}, headers=izleyici_header)
    assert r.status_code == 403


def test_kamera_roi_guncelle_operator_calistirabilir_ve_alani_kaydeder(client, operator_header, roi_test_kamera_id):
    r = client.patch(f"/kameralar/{roi_test_kamera_id}/roi",
                      json={"x1": 10, "y1": 5, "x2": 90, "y2": 95}, headers=operator_header)
    assert r.status_code == 200, r.text
    kamera = r.json()
    assert kamera["roi"] == {"x1": 10.0, "y1": 5.0, "x2": 90.0, "y2": 95.0}

    # Kalıcı mı diye /kameralar listesinden de doğrula.
    r2 = client.get("/kameralar", headers=operator_header)
    kaydedilen = next(k for k in r2.json() if k["id"] == roi_test_kamera_id)
    assert kaydedilen["roi"] == {"x1": 10.0, "y1": 5.0, "x2": 90.0, "y2": 95.0}


def test_kamera_roi_guncelle_temizle_alani_kaldirir(client, operator_header, roi_test_kamera_id):
    # Önce bir alan tanımlı olduğundan emin ol.
    r0 = client.patch(f"/kameralar/{roi_test_kamera_id}/roi",
                       json={"x1": 0, "y1": 0, "x2": 50, "y2": 50}, headers=operator_header)
    assert r0.status_code == 200 and r0.json()["roi"] is not None

    r = client.patch(f"/kameralar/{roi_test_kamera_id}/roi", json={"temizle": True}, headers=operator_header)
    assert r.status_code == 200, r.text
    # NOT: temizlenince "roi" anahtarı sözlükten tamamen kaldırılıyor
    # (kamera.pop("roi", None)) -- yanıt JSON'unda hiç bulunmaz, None değil.
    assert r.json().get("roi") is None


@pytest.mark.parametrize("govde", [
    {"x1": 50, "y1": 0, "x2": 40, "y2": 100},   # x1 >= x2
    {"x1": 0, "y1": 80, "x2": 100, "y2": 20},   # y1 >= y2
    {"x1": -5, "y1": 0, "x2": 100, "y2": 100},  # 0-100 dışı
    {"x1": 0, "y1": 0, "x2": 100, "y2": 150},   # 0-100 dışı
    {"x1": 0, "y1": 0, "x2": 100},              # eksik alan (y2 yok), temizle=false
])
def test_kamera_roi_guncelle_gecersiz_degerler_400_doner(client, operator_header, roi_test_kamera_id, govde):
    r = client.patch(f"/kameralar/{roi_test_kamera_id}/roi", json=govde, headers=operator_header)
    assert r.status_code == 400, r.text


def test_kamera_roi_guncelle_bulunamayan_kamera_404_doner(client, operator_header):
    r = client.patch("/kameralar/YOK-BOYLE-BIR-KAMERA/roi",
                      json={"x1": 0, "y1": 0, "x2": 100, "y2": 100}, headers=operator_header)
    assert r.status_code == 404


def test_kamera_roi_guncelle_izleyici_yetkisiz_403_doner(client, izleyici_header, roi_test_kamera_id):
    r = client.patch(f"/kameralar/{roi_test_kamera_id}/roi",
                      json={"x1": 0, "y1": 0, "x2": 100, "y2": 100}, headers=izleyici_header)
    assert r.status_code == 403


# ------------------------------------------------------------------
# Tespit alanı (ROI) — SERBEST ÇİZİM (polygon), 2026-09-20
#
# Kullanıcı geri bildirimi (birebir): "bir de alan sınırı eklemiştik bu alan
# sınırına serbest çizim ekleme şansımız var mı kare seçimde bazen farklı
# yönden geçen araçları da tespit ediyor bunu istemiyorum" -- dikdörtgenin
# yanına, keyfi köşe sayılı bir çokgen tanımlanabilen `polygon` alanı eklendi
# (bkz. schemas.KameraRoiGuncelle, main.py::_roi_polygon_gecerlilestir).
# ------------------------------------------------------------------

def test_kamera_roi_guncelle_polygon_ile_kaydeder(client, operator_header, roi_test_kamera_id):
    noktalar = [{"x": 0, "y": 0}, {"x": 30, "y": 0}, {"x": 70, "y": 100}, {"x": 0, "y": 100}]
    r = client.patch(f"/kameralar/{roi_test_kamera_id}/roi",
                      json={"polygon": noktalar}, headers=operator_header)
    assert r.status_code == 200, r.text
    kamera = r.json()
    assert kamera["roi"] == {"tip": "polygon", "noktalar": noktalar}

    # Kalıcı mı diye /kameralar listesinden de doğrula.
    r2 = client.get("/kameralar", headers=operator_header)
    kaydedilen = next(k for k in r2.json() if k["id"] == roi_test_kamera_id)
    assert kaydedilen["roi"] == {"tip": "polygon", "noktalar": noktalar}


def test_kamera_roi_guncelle_polygon_temizlenebilir(client, operator_header, roi_test_kamera_id):
    r0 = client.patch(f"/kameralar/{roi_test_kamera_id}/roi",
                       json={"polygon": [{"x": 0, "y": 0}, {"x": 50, "y": 0}, {"x": 50, "y": 50}]},
                       headers=operator_header)
    assert r0.status_code == 200 and r0.json()["roi"] is not None

    r = client.patch(f"/kameralar/{roi_test_kamera_id}/roi", json={"temizle": True}, headers=operator_header)
    assert r.status_code == 200, r.text
    assert r.json().get("roi") is None


def test_kamera_roi_guncelle_polygon_2_nokta_ile_400_doner(client, operator_header, roi_test_kamera_id):
    """Bir çokgen için en az 3 köşe gerekir -- 2 nokta (bir doğru parçası)
    geçerli bir alan tanımlamaz."""
    r = client.patch(f"/kameralar/{roi_test_kamera_id}/roi",
                      json={"polygon": [{"x": 0, "y": 0}, {"x": 50, "y": 50}]}, headers=operator_header)
    assert r.status_code == 400, r.text


def test_kamera_roi_guncelle_polygon_asiri_nokta_ile_400_doner(client, operator_header, roi_test_kamera_id):
    """Aşırı sayıda nokta (örn. bir sürükleme olayının yanlışlıkla yüzlerce
    tıklama olarak işlenmesi) reddedilmeli -- bkz. main.py::
    _ROI_POLIGON_MAX_NOKTA."""
    noktalar = [{"x": i % 100, "y": (i * 3) % 100} for i in range(30)]
    r = client.patch(f"/kameralar/{roi_test_kamera_id}/roi",
                      json={"polygon": noktalar}, headers=operator_header)
    assert r.status_code == 400, r.text


@pytest.mark.parametrize("noktalar", [
    [{"x": -5, "y": 0}, {"x": 50, "y": 0}, {"x": 50, "y": 50}],   # 0-100 dışı (x)
    [{"x": 0, "y": 0}, {"x": 50, "y": 150}, {"x": 50, "y": 50}],  # 0-100 dışı (y)
])
def test_kamera_roi_guncelle_polygon_gecersiz_koordinat_400_doner(client, operator_header, roi_test_kamera_id, noktalar):
    r = client.patch(f"/kameralar/{roi_test_kamera_id}/roi",
                      json={"polygon": noktalar}, headers=operator_header)
    assert r.status_code == 400, r.text


def test_kamera_roi_guncelle_polygon_izleyici_yetkisiz_403_doner(client, izleyici_header, roi_test_kamera_id):
    r = client.patch(f"/kameralar/{roi_test_kamera_id}/roi",
                      json={"polygon": [{"x": 0, "y": 0}, {"x": 50, "y": 0}, {"x": 50, "y": 50}]},
                      headers=izleyici_header)
    assert r.status_code == 403


# ------------------------------------------------------------------
# "sakin" (site sakini öz-hizmet portalı) rolü — 2026-09-17
# ------------------------------------------------------------------
# Kök neden: kullanıcı, ekran görüntüleriyle bir referans ürünün "Abone
# Düzenle" ekranını göstererek site sakinlerinin kendi kullanıcı adı/
# şifresiyle giriş yapıp kendi araç/geçmiş bilgisini yönetebildiği bir
# öz-hizmet hesabı istedi. Bu, panelin önceki üç rolünden (yönetici/
# operatör/izleyici — hepsi "güvenilir iç personel") temelde farklı: bir
# "sakin" DIŞARIDAN bir hesap olduğu için, yalnızca kendi verisine
# erişebilmesi ve tüm dahili/genel amaçlı uçlara (kişi listesi, tüm
# kayıtlar, kameralar, sistem logları vb.) KESİNLİKLE erişememesi gerekir
# (bkz. main.py::_personel_girisi_gerekli ve _sakin_girisi_gerekli).

def test_sakin_olusturma_kisi_id_zorunlu_400_doner(client, yetkili_header):
    r = client.post("/kullanicilar", json={
        "kullanici_adi": "sakin-kisisiz", "parola": "GucluParola123!", "rol": "sakin",
    }, headers=yetkili_header)
    assert r.status_code == 400, r.text


def test_sakin_olusturma_olmayan_kisi_404_doner(client, yetkili_header):
    r = client.post("/kullanicilar", json={
        "kullanici_adi": "sakin-yokkisi", "parola": "GucluParola123!", "rol": "sakin", "kisi_id": 999999,
    }, headers=yetkili_header)
    assert r.status_code == 404, r.text


def test_sakin_olusturma_gecerli_kisiyle_basarili(client, yetkili_header, sakin_kisi_id):
    r = client.post("/kullanicilar", json={
        "kullanici_adi": "sakin-basarili", "parola": "GucluParola123!", "rol": "sakin", "kisi_id": sakin_kisi_id,
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text
    assert r.json()["kisi_id"] == sakin_kisi_id
    assert r.json()["rol"] == "sakin"


def test_sakin_olmayan_rolde_kisi_id_yok_sayilir(client, yetkili_header, sakin_kisi_id):
    """rol != 'sakin' iken kisi_id gönderilse bile anlamsızdır ve kayda
    geçmemeli (bkz. main.py::_sakin_kisi_id_dogrula)."""
    r = client.post("/kullanicilar", json={
        "kullanici_adi": "izleyici-kisili", "parola": "GucluParola123!", "rol": "izleyici", "kisi_id": sakin_kisi_id,
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text
    assert r.json()["kisi_id"] is None


def test_kullanici_guncelle_sakine_terfi_kisi_id_gerektirir(client, yetkili_header):
    r = client.post("/kullanicilar", json={
        "kullanici_adi": "terfi-oncesi", "parola": "GucluParola123!", "rol": "izleyici",
    }, headers=yetkili_header)
    kid = r.json()["id"]
    r2 = client.put(f"/kullanicilar/{kid}", json={"rol": "sakin"}, headers=yetkili_header)
    assert r2.status_code == 400, r2.text


def test_kullanici_guncelle_sakinden_baska_role_kisi_id_temizlenir(client, yetkili_header, sakin_kisi_id):
    r = client.post("/kullanicilar", json={
        "kullanici_adi": "sakin-sonra-operator", "parola": "GucluParola123!", "rol": "sakin", "kisi_id": sakin_kisi_id,
    }, headers=yetkili_header)
    kid = r.json()["id"]
    r2 = client.put(f"/kullanicilar/{kid}", json={"rol": "operatör"}, headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    assert r2.json()["kisi_id"] is None


def test_kullanici_guncelle_rol_degisikligi_loglanir(client, caplog, yetkili_header):
    """2026-09-20: rol değişikliği, hesap aktif/pasif yapma ve parola
    sıfırlama gibi hassas kullanıcı-yönetimi işlemleri önceden HİÇBİR YERE
    loglanmıyordu -- "kim, kimin rolünü ne zaman değiştirdi" sorusu
    cevapsızdı (bkz. main.py::kullanici_guncelle). Parolanın KENDİSİNİN
    asla loglanmadığını da doğruluyoruz."""
    r = client.post("/kullanicilar", json={
        "kullanici_adi": "denetim-log-testi", "parola": "GucluParola123!", "rol": "izleyici",
    }, headers=yetkili_header)
    kid = r.json()["id"]
    with caplog.at_level("INFO", logger="pts"):
        r2 = client.put(
            f"/kullanicilar/{kid}",
            json={"rol": "operatör", "aktif": False, "parola": "YeniGucluParola456!"},
            headers=yetkili_header,
        )
    assert r2.status_code == 200, r2.text
    assert "denetim-log-testi" in caplog.text
    assert "rol: izleyici -> operatör" in caplog.text
    assert "aktif: True -> False" in caplog.text
    assert "parola sıfırlandı" in caplog.text
    assert "YeniGucluParola456!" not in caplog.text  # parolanın kendisi ASLA loglanmamalı

    r3 = client.get("/denetim-kayitlari", params={"kullanici_adi": "denetim-log-testi"}, headers=yetkili_header)
    assert r3.status_code == 200, r3.text
    kayit = next(k for k in r3.json() if k["eylem"] == "kullanici_guncelle")
    assert "rol: izleyici -> operatör" in kayit["aciklama"]
    assert "YeniGucluParola456!" not in kayit["aciklama"]


def test_kullanici_guncelle_degisiklik_yoksa_loglanmiyor(client, caplog, yetkili_header):
    """Boş/etkisiz bir PUT (mevcut değerlerin aynısı gönderilmesi) log
    spam'ine yol açmamalı (bkz. main.py::kullanici_guncelle)."""
    r = client.post("/kullanicilar", json={
        "kullanici_adi": "degisiklik-yok-testi", "parola": "GucluParola123!", "rol": "izleyici",
    }, headers=yetkili_header)
    kid = r.json()["id"]
    with caplog.at_level("INFO", logger="pts"):
        r2 = client.put(f"/kullanicilar/{kid}", json={}, headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    assert "degisiklik-yok-testi" not in caplog.text


def test_kullanici_silme_loglanir(client, caplog, yetkili_header):
    """2026-09-20: bir kullanıcı hesabının silinmesi önceden hiçbir yere
    loglanmıyordu (bkz. main.py::kullanici_sil)."""
    r = client.post("/kullanicilar", json={
        "kullanici_adi": "silinecek-denetim-testi", "parola": "GucluParola123!", "rol": "izleyici",
    }, headers=yetkili_header)
    kid = r.json()["id"]
    with caplog.at_level("INFO", logger="pts"):
        r2 = client.delete(f"/kullanicilar/{kid}", headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    assert "silinecek-denetim-testi" in caplog.text

    r3 = client.get("/denetim-kayitlari", params={"eylem": "kullanici_sil"}, headers=yetkili_header)
    assert r3.status_code == 200, r3.text
    assert any("silinecek-denetim-testi" in k["aciklama"] for k in r3.json())


# ------------------------------------------------------------------
# "Nizamiye Bazlı Kamera Erişimi" (2026-09-21) -- bir kullanıcı hesabını
# belirli kameralarla sınırlama (bkz. models.Kullanici.kamera_erisim_listesi,
# main.py::_kullanicinin_izinli_kameralari). Kullanıcı talebi: aynı ağdaki
# farklı bilgisayarlardan/noktalardan çalışan güvenlik personelinin yalnızca
# KENDİ noktasının kameralarını görebilmesi ("Bülent" örneği -- Lojman
# Nizamiye'de yalnızca 2 kamera görmesi gerekiyor, Ana Nizamiye'nin 4
# kamerasını GÖRMEMELİ).
# ------------------------------------------------------------------

@pytest.fixture(scope="module")
def kamera_erisim_test_kameralari(client, yetkili_header):
    """Bu testler için İKİ ayrı GERÇEK kamera (bkz. roi_test_kamera_id'deki
    AYNI lisans-yükseltme deseni -- bu modüldeki lisans limiti önceki
    testlerde tüketilmiş olabilir).

    Hem "id" (cameras.json'daki KİMLİK -- her zaman rastgele bir UUID, bkz.
    main.py::kamera_ekle) hem de "ad" (gerçek geçiş kayıtlarının
    `Kayit.kamera_id` alanına damgalanan değer, bkz. main.py::_pipeline_
    baslat) döndürülür -- ikisi KASITLI olarak FARKLI (id ASLA "ad" ile aynı
    değildir, gerçek bir kamera eklerken de böyle), ki testler "id ile ad
    tesadüfen aynı" durumunu maskeleyip 2026-09-21'de gerçek üretim
    verisiyle bulunan id/ad karışıklığı hatasını YENİDEN gözden
    kaçırmasın."""
    anahtar = lisans.uret("Kamera Erisim Test Site", kamera_limiti=10, gun=30)
    r = client.post("/lisans/aktive-et", json={"anahtar": anahtar}, headers=yetkili_header)
    assert r.status_code == 200, r.text
    ana_ad = "Ana Nizamiye Kamerası"
    lojman_ad = "Lojman Nizamiye Kamerası"
    r1 = client.post("/kameralar", json={
        "ad": ana_ad, "rtsp_url": "rtsp://127.0.0.1/ana-nizamiye", "yon": "giris",
    }, headers=yetkili_header)
    assert r1.status_code == 200, r1.text
    r2 = client.post("/kameralar", json={
        "ad": lojman_ad, "rtsp_url": "rtsp://127.0.0.1/lojman-nizamiye", "yon": "giris",
    }, headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    assert r1.json()["id"] != ana_ad and r2.json()["id"] != lojman_ad, (
        "id, UUID olduğu için 'ad' ile ASLA aynı olmamalı -- aksi halde bu fixture "
        "id/ad karışıklığı hatasını maskeler"
    )
    return {
        "ana": r1.json()["id"], "ana_ad": ana_ad,
        "lojman": r2.json()["id"], "lojman_ad": lojman_ad,
    }


def _kullanici_id_bul(client, yetkili_header, kullanici_adi: str) -> int:
    r = client.get("/kullanicilar", headers=yetkili_header)
    return next(k["id"] for k in r.json() if k["kullanici_adi"] == kullanici_adi)


def test_kullanici_ekle_gecersiz_kamera_id_ile_reddedilir(client, yetkili_header):
    r = client.post("/kullanicilar", json={
        "kullanici_adi": "kamera-kisitli-gecersiz", "parola": "GucluParola123!", "rol": "güvenlik",
        "kamera_erisim_listesi": ["olmayan-kamera-id"],
    }, headers=yetkili_header)
    assert r.status_code == 400, r.text


def test_kullanici_ekle_kamera_kisitlamasiyla_olusturulur(client, yetkili_header, kamera_erisim_test_kameralari):
    r = client.post("/kullanicilar", json={
        "kullanici_adi": "bulent-lojman", "parola": "GucluParola123!", "rol": "güvenlik",
        "kamera_erisim_listesi": [kamera_erisim_test_kameralari["lojman"]],
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text
    assert r.json()["kamera_erisim_listesi"] == [kamera_erisim_test_kameralari["lojman"]]


def test_kullanici_ekle_kisitlamasiz_kamera_erisim_listesi_null_doner(client, yetkili_header):
    """Geriye dönük uyumluluk: `kamera_erisim_listesi` hiç gönderilmezse
    hesap kısıtlamasız (tüm kameraları görebilir) kalmalı -- mevcut tüm
    hesaplar bu haldeydi, bu özellik onları ETKİLEMEMELİ."""
    r = client.post("/kullanicilar", json={
        "kullanici_adi": "kisitlamasiz-hesap", "parola": "GucluParola123!", "rol": "izleyici",
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text
    assert r.json()["kamera_erisim_listesi"] is None


def test_kamera_kisitli_hesap_kameralar_listesinde_sadece_izinliyi_gorur(client, yetkili_header, kamera_erisim_test_kameralari):
    hedef = _rol_ile_kullanici_olustur_ve_giris_yap(client, yetkili_header, "lojman-izleyici-1", "izleyici")
    kid = _kullanici_id_bul(client, yetkili_header, "lojman-izleyici-1")
    r = client.put(f"/kullanicilar/{kid}", json={
        "kamera_erisim_listesi": [kamera_erisim_test_kameralari["lojman"]],
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text

    r2 = client.get("/kameralar", headers=hedef)
    assert r2.status_code == 200, r2.text
    assert [k["id"] for k in r2.json()] == [kamera_erisim_test_kameralari["lojman"]]


def test_kamera_kisitli_hesap_izinsiz_kameranin_goruntusunu_alamaz(client, yetkili_header, kamera_erisim_test_kameralari):
    hedef = _rol_ile_kullanici_olustur_ve_giris_yap(client, yetkili_header, "lojman-izleyici-2", "izleyici")
    kid = _kullanici_id_bul(client, yetkili_header, "lojman-izleyici-2")
    client.put(f"/kullanicilar/{kid}", json={
        "kamera_erisim_listesi": [kamera_erisim_test_kameralari["lojman"]],
    }, headers=yetkili_header)

    r = client.get(f"/kameralar/{kamera_erisim_test_kameralari['ana']}/goruntu", headers=hedef)
    assert r.status_code == 403, r.text
    r2 = client.get(f"/kameralar/{kamera_erisim_test_kameralari['ana']}/son-plaka", headers=hedef)
    assert r2.status_code == 403, r2.text
    # İZİNLİ kameraya erişim REDDEDİLMEMELİ -- gerçek bir RTSP bağlantısı
    # olmadığı için "son-plaka" boş liste döner, önemli olan 403 OLMAMASI.
    r3 = client.get(f"/kameralar/{kamera_erisim_test_kameralari['lojman']}/son-plaka", headers=hedef)
    assert r3.status_code == 200, r3.text


def test_kamera_kisitli_hesap_kayitlar_listesinde_sadece_izinli_kameradan_gelenleri_gorur(client, yetkili_header, kamera_erisim_test_kameralari):
    """KÖK NEDEN DÜZELTMESİ (2026-09-21, gerçek üretim verisiyle bulunan hata
    -- "Lojman A Vardiyası" hesabı hiçbir geçiş göremiyordu): gerçek geçiş
    kayıtları `Kayit.kamera_id` alanına kameranın "id"si DEĞİL "ad"ıyla
    damgalanır (bkz. main.py::_pipeline_baslat) -- bu yüzden burada da
    kayıtlar `..._ad` değerleriyle oluşturulur (gerçek pipeline'ı simüle
    eder), ama kısıtlama (kamera_erisim_listesi) YİNE "id" ile ayarlanır
    (panelin/API'nin beklediği biçim budur) -- bkz. main.py::
    _kullanicinin_izinli_kamera_adlari'nın bu ikisini nasıl eşleştirdiği."""
    ana_ad = kamera_erisim_test_kameralari["ana_ad"]
    lojman_ad = kamera_erisim_test_kameralari["lojman_ad"]
    lojman_id = kamera_erisim_test_kameralari["lojman"]
    r1 = client.post("/kayitlar", json={"plaka_no": "34 KAM 01", "kamera_id": ana_ad, "yon": "giris"}, headers=yetkili_header)
    assert r1.status_code == 200, r1.text
    r2 = client.post("/kayitlar", json={"plaka_no": "34 KAM 02", "kamera_id": lojman_ad, "yon": "giris"}, headers=yetkili_header)
    assert r2.status_code == 200, r2.text

    hedef = _rol_ile_kullanici_olustur_ve_giris_yap(client, yetkili_header, "lojman-izleyici-3", "izleyici")
    kid = _kullanici_id_bul(client, yetkili_header, "lojman-izleyici-3")
    client.put(f"/kullanicilar/{kid}", json={"kamera_erisim_listesi": [lojman_id]}, headers=yetkili_header)

    r3 = client.get("/kayitlar", params={"plaka": "34 KAM"}, headers=hedef)
    assert r3.status_code == 200, r3.text
    plakalar = {k["plaka_no"] for k in r3.json()}
    assert "34 KAM 02" in plakalar, (
        "kısıtlama id bazlı saklansa da, id'nin cameras.json üzerinden karşılık geldiği "
        "'ad' değeriyle damgalanmış kayıt GÖRÜNMELİ (bkz. _kullanicinin_izinli_kamera_adlari)"
    )
    assert "34 KAM 01" not in plakalar


def test_plaka_analizi_kamera_kisitlamasindan_muaftir(client, yetkili_header, kamera_erisim_test_kameralari):
    """Kullanıcı talebi (2026-09-21): bir vardiya/nokta, BAŞKA bir
    vardiyanın/noktanın kamerasından geçen belirli bir aracı "Plaka Analizi"
    ile TESPİT edebilmeli ("A Vardiyasının nöbet saatinde giriş yapan bir
    aracı, B vardiyası geldiğinde tespit edebilmesi gerekiyor") -- bkz.
    main.py::plaka_analiz'in güncellenen docstring'i. Bu test hem KAMERA hem
    de VARDİYA PENCERESİ kısıtlamasının BİRLİKTE muaf tutulduğunu doğrular:
    kayıt, hedef güvenlik hesabının vardiyası başlamadan ÖNCE ve izinsiz bir
    kameradan oluşturuluyor."""
    ana_id = kamera_erisim_test_kameralari["ana"]
    r1 = client.post("/kayitlar", json={"plaka_no": "34 MUAF 01", "kamera_id": ana_id, "yon": "giris"}, headers=yetkili_header)
    assert r1.status_code == 200, r1.text

    # Giriş (bkz. _rol_ile_kullanici_olustur_ve_giris_yap -> /auth/giris),
    # güvenlik rolü için ŞİMDİ başlayan bir vardiya oturumu açar -- yukarıdaki
    # kayıt bundan ÖNCE oluşturulduğu için normal şartlarda vardiya
    # penceresinin de dışında kalırdı.
    hedef = _rol_ile_kullanici_olustur_ve_giris_yap(client, yetkili_header, "lojman-guvenlik-4", "güvenlik")
    kid = _kullanici_id_bul(client, yetkili_header, "lojman-guvenlik-4")
    client.put(f"/kullanicilar/{kid}", json={
        "kamera_erisim_listesi": [kamera_erisim_test_kameralari["lojman"]],
    }, headers=yetkili_header)

    # Genel "Kayıtlar" listesinde göremez (hem kamera hem vardiya kısıtlaması) ...
    r2 = client.get("/kayitlar", params={"plaka": "34 MUAF"}, headers=hedef)
    assert r2.status_code == 200, r2.text
    assert not any(k["plaka_no"] == "34 MUAF 01" for k in r2.json())

    # ... ama "Plaka Analizi" ile TAM geçmişi görebilir.
    r3 = client.get("/kayitlar/analiz/34 MUAF 01", headers=hedef)
    assert r3.status_code == 200, r3.text
    assert r3.json()["toplam_gecis"] == 1


def test_kamera_erisimi_temizle_kisitlamayi_kaldirir(client, yetkili_header, kamera_erisim_test_kameralari):
    r = client.post("/kullanicilar", json={
        "kullanici_adi": "kisitlama-kaldirilacak", "parola": "GucluParola123!", "rol": "izleyici",
        "kamera_erisim_listesi": [kamera_erisim_test_kameralari["lojman"]],
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text
    kid = r.json()["id"]
    assert r.json()["kamera_erisim_listesi"] == [kamera_erisim_test_kameralari["lojman"]]

    r2 = client.put(f"/kullanicilar/{kid}", json={"kamera_erisimi_temizle": True}, headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    assert r2.json()["kamera_erisim_listesi"] is None


def test_kamera_erisimi_id_ad_karisikligindan_etkilenmez(client, yetkili_header, kamera_erisim_test_kameralari):
    """KÖK NEDEN REGRESYON TESTİ (2026-09-21, gerçek üretim verisiyle
    bulunan hata): `_kullanicinin_izinli_kamera_adlari`'nin, id bazlı bir
    kısıtlamayı doğru "ad" değerlerine çevirdiğini DOĞRUDAN doğrular --
    `pts_main._kullanicinin_izinli_kameralari` (eski, id bazlı) ile
    KARIŞTIRILMAMASI gereken yeni fonksiyon. Ayrıca tekil kayıt görünürlüğü
    uç noktasının (`/disa-aktar/pdf/kayit/{id}`, bkz. main.py::
    _guvenlik_kayit_gorunur_mu) da AYNI ad-bazlı eşleştirmeyi kullandığını
    doğrular -- bu, `_guvenlik_kayit_filtresi_uygula`'dan AYRI bir kod
    yoludur, düzeltme unutulursa bağımsız olarak bozuk kalabilirdi."""
    import json as _json
    from backend.database import SessionLocal
    from backend import models as _models

    lojman_id = kamera_erisim_test_kameralari["lojman"]
    lojman_ad = kamera_erisim_test_kameralari["lojman_ad"]
    ana_ad = kamera_erisim_test_kameralari["ana_ad"]

    db = SessionLocal()
    try:
        kisitli = _models.Kullanici(
            kullanici_adi="birim-testi-kamera-kisitli-x", parola_hash="x", rol="izleyici",
            kamera_erisim_listesi=_json.dumps([lojman_id]),
        )
        # Fonksiyon sadece kullanici.kamera_erisim_listesi'ne bakıyor, DB'ye
        # eklemeye gerek yok -- ama import döngüsünü basit tutmak için
        # gerçek bir ORM nesnesi kullanılıyor.
        izinli_adlar = pts_main._kullanicinin_izinli_kamera_adlari(kisitli)
    finally:
        db.close()
    assert izinli_adlar == {lojman_ad}, (
        f"id bazlı kısıtlama ({lojman_id}) karşılık gelen 'ad' değerine ({lojman_ad}) "
        f"çevrilmeli, ham id'ye DEĞİL -- döndü: {izinli_adlar}"
    )

    r1 = client.post("/kayitlar", json={"plaka_no": "34 KAM 03", "kamera_id": lojman_ad, "yon": "giris"}, headers=yetkili_header)
    assert r1.status_code == 200, r1.text
    kayit_id_izinli = r1.json()["id"]
    r2 = client.post("/kayitlar", json={"plaka_no": "34 KAM 04", "kamera_id": ana_ad, "yon": "giris"}, headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    kayit_id_izinsiz = r2.json()["id"]

    hedef = _rol_ile_kullanici_olustur_ve_giris_yap(client, yetkili_header, "lojman-izleyici-kamera-x", "izleyici")
    kid = _kullanici_id_bul(client, yetkili_header, "lojman-izleyici-kamera-x")
    client.put(f"/kullanicilar/{kid}", json={"kamera_erisim_listesi": [lojman_id]}, headers=yetkili_header)

    assert client.get(f"/disa-aktar/pdf/kayit/{kayit_id_izinli}", headers=hedef).status_code == 200
    assert client.get(f"/disa-aktar/pdf/kayit/{kayit_id_izinsiz}", headers=hedef).status_code == 403


def test_kisi_silinince_bagli_sakin_hesabinin_kisi_id_temizlenir(client, yetkili_header):
    rk = client.post("/kisiler", json={
        "ad_soyad": "Silinecek Sakin Kişisi", "plaka_no": "34 SLN 02", "tip": "abone",
    }, headers=yetkili_header)
    kisi_id = rk.json()["id"]
    ru = client.post("/kullanicilar", json={
        "kullanici_adi": "sakin-yetim-olacak", "parola": "GucluParola123!", "rol": "sakin", "kisi_id": kisi_id,
    }, headers=yetkili_header)
    kullanici_id = ru.json()["id"]

    rs = client.delete(f"/kisiler/{kisi_id}", headers=yetkili_header)
    assert rs.status_code == 200, rs.text

    rl = client.get("/kullanicilar", headers=yetkili_header)
    hedef = next(k for k in rl.json() if k["id"] == kullanici_id)
    assert hedef["kisi_id"] is None


@pytest.mark.parametrize("yol", ["/kisiler", "/kayitlar", "/kameralar", "/kara-listesi", "/bariyer/ayarlar", "/siteler"])
def test_sakin_genel_amacli_uclara_erisemez_403_doner(client, sakin_header, yol):
    """KÖK NEDEN testi: _personel_girisi_gerekli eklenmeden önce bu uçlar
    yalnızca _giris_gerekli ile korunuyordu -- yani giriş yapmış HERHANGİ
    bir hesap (sakin dahil) tüm kişileri/kayıtları/kameraları görebilirdi."""
    r = client.get(yol, headers=sakin_header)
    assert r.status_code == 403, f"GET {yol}: sakin 403 almalıydı, {r.status_code} aldı ({r.text})"


def test_sakin_auth_me_cagirabilir(client, sakin_header):
    """Bu tek istisna: frontend'in doğru arayüze (öz-hizmet paneli) karar
    verebilmesi için sakin de /auth/me'yi çağırabilmeli."""
    r = client.get("/auth/me", headers=sakin_header)
    assert r.status_code == 200, r.text
    assert r.json()["rol"] == "sakin"


def test_sakin_olmayan_roller_sakin_uclarina_erisemez_403_doner(client, izleyici_header, operator_header, yetkili_header):
    for header in (izleyici_header, operator_header, yetkili_header):
        r = client.get("/sakin/profilim", headers=header)
        assert r.status_code == 403, r.text


def test_sakin_profilim_kendi_kisi_kaydini_doner(client, sakin_header, sakin_kisi_id):
    r = client.get("/sakin/profilim", headers=sakin_header)
    assert r.status_code == 200, r.text
    assert r.json()["id"] == sakin_kisi_id
    assert r.json()["ad_soyad"] == "Sakin Test Kişi"
    assert r.json()["ek_plakalar"] == []


def test_sakin_arac_ekle_ve_sil(client, sakin_header):
    r = client.post("/sakin/arac-ekle", json={"plaka_no": "34 SKN 02", "aciklama": "ikinci aracım"}, headers=sakin_header)
    assert r.status_code == 200, r.text
    plaka_id = r.json()["id"]

    r2 = client.get("/sakin/profilim", headers=sakin_header)
    assert any(p["id"] == plaka_id for p in r2.json()["ek_plakalar"])

    r3 = client.delete(f"/sakin/arac/{plaka_id}", headers=sakin_header)
    assert r3.status_code == 200, r3.text

    r4 = client.get("/sakin/profilim", headers=sakin_header)
    assert not any(p["id"] == plaka_id for p in r4.json()["ek_plakalar"])


def test_sakin_baska_kisinin_plakasini_silemez_404_doner(client, sakin_header, yetkili_header):
    """IDOR regresyonu: bir sakin, kendi kisi_id'sine ait olmayan bir
    KisiPlaka id'sini tahmin ederek silemesin (bkz. main.py::sakin_arac_sil)."""
    rk = client.post("/kisiler", json={
        "ad_soyad": "Başka Sakin", "plaka_no": "34 BSK 03", "tip": "abone",
    }, headers=yetkili_header)
    baska_kisi_id = rk.json()["id"]
    rp = client.post(f"/kisiler/{baska_kisi_id}/plakalar", json={"plaka_no": "34 BSK 04"}, headers=yetkili_header)
    baska_plaka_id = rp.json()["id"]

    r = client.delete(f"/sakin/arac/{baska_plaka_id}", headers=sakin_header)
    assert r.status_code == 404, r.text


def test_sakin_gecmisim_yalnizca_kendi_kayitlarini_doner(client, sakin_header, operator_header, sakin_kisi_id):
    # Sakinin kendi kişisine bağlı bir kayıt oluştur.
    r1 = client.post("/kayitlar", json={"plaka_no": "34 SKN 01", "kamera_id": "TEST-SAKIN", "yon": "giris"}, headers=operator_header)
    assert r1.status_code == 200, r1.text
    assert r1.json()["kisi_id"] == sakin_kisi_id, "Test öncülü: plaka zaten sakin_kisi_id'ye eşleşmeliydi"

    # Başka bir plakaya ait, sakine bağlı OLMAYAN bir kayıt da oluştur.
    r2 = client.post("/kayitlar", json={"plaka_no": "34 BASKASI 05", "kamera_id": "TEST-SAKIN", "yon": "giris"}, headers=operator_header)
    assert r2.status_code == 200, r2.text

    r3 = client.get("/sakin/gecmisim", headers=sakin_header)
    assert r3.status_code == 200, r3.text
    plakalar = {k["plaka_no"] for k in r3.json()}
    assert "34 SKN 01" in plakalar
    assert "34 BASKASI 05" not in plakalar


def test_sakin_hesabi_yetim_kalinca_400_doner(client, yetkili_header):
    """_sakin_kisisini_al'ın 'hesap henüz bağlanmamış' savunmasını, normal
    API akışında bu duruma ULAŞILABİLECEK tek yoldan (bağlı Kişi silinince
    kisi_id NULL'lanır, bkz. test_kisi_silinince_bagli_sakin... ve
    main.py::kisi_sil) uçtan uca doğrular -- ayrı bir kişi/hesap kullanılır
    ki paylaşılan `sakin_header`/`sakin_kisi_id` fixture'ları başka testler
    için bozulmasın."""
    rk = client.post("/kisiler", json={
        "ad_soyad": "Yetim Kalacak Sakin", "plaka_no": "34 YTM 07", "tip": "abone",
    }, headers=yetkili_header)
    kisi_id = rk.json()["id"]
    ru = client.post("/kullanicilar", json={
        "kullanici_adi": "sakin-yetim-kalan", "parola": "GucluParola123!", "rol": "sakin", "kisi_id": kisi_id,
    }, headers=yetkili_header)
    assert ru.status_code == 200, ru.text
    r2 = client.post("/auth/giris", json={"kullanici_adi": "sakin-yetim-kalan", "parola": "GucluParola123!"})
    yetim_header = {"Authorization": f"Bearer {r2.json()['token']}"}

    rs = client.delete(f"/kisiler/{kisi_id}", headers=yetkili_header)
    assert rs.status_code == 200, rs.text

    r3 = client.get("/sakin/profilim", headers=yetim_header)
    assert r3.status_code == 400, r3.text


def test_ziyaretci_onayli_yetki_durumu_kayit_duzenlede_kabul_edilir(client, operator_header):
    """Yeni "Ziyaretçi Girişi" akışının (bkz. app.js::_ziyaretciGirisiKutusunuAyarla,
    ziyaretciBilgileriKaydet) dayandığı whitelist genişletmesi: kayit_duzenle
    artık yetki_durumu="ziyaretci_onayli" değerini kabul etmeli."""
    r = client.post("/kayitlar", json={"plaka_no": "34 ZYR 06", "kamera_id": "TEST-ZYR", "yon": "giris"}, headers=operator_header)
    kayit_id = r.json()["id"]
    r2 = client.patch(f"/kayitlar/{kayit_id}", json={"yetki_durumu": "ziyaretci_onayli"}, headers=operator_header)
    assert r2.status_code == 200, r2.text
    assert r2.json()["yetki_durumu"] == "ziyaretci_onayli"


# ------------------------------------------------------------------
# Otomatik kayıt güven eşiği filtresi (2026-09-17)
# ------------------------------------------------------------------
# Kök neden: kullanıcı, hatalı/yanlış okunan düşük güvenli OCR tespitlerinin
# panele/kayıtlara düşüp kayıtları şişirdiğini bildirdi ve yalnızca %97-%100
# güvenli tespitlerin kayda geçmesini, geri kalanının HİÇ kaydedilmemesini
# istedi (bkz. main.py::kayit_ekle_otomatik'teki gerçek-zamanlı filtre ve
# _VARSAYILAN_AYARLAR'daki "otomatik_kayit_min_guven_skoru" notu). Sabit bir
# değere güvenmek yerine testler her zaman GEÇERLİ ayarı (/sistem/ayarlar)
# okuyup ona göre eşik üstü/altı değerler üretir -- böylece modül içindeki
# başka bir testin ayarı değiştirmesi (şu an hiçbiri değiştirmiyor, ama
# ileride değiştirebilir) bu testleri kırılgan hale getirmez.

def _mevcut_otomatik_kayit_esigi(client, yetkili_header) -> float:
    r = client.get("/sistem/ayarlar", headers=yetkili_header)
    assert r.status_code == 200, r.text
    return float(r.json().get("otomatik_kayit_min_guven_skoru", 0.97))


def test_otomatik_kayit_esigi_varsayilan_097(client, yetkili_header):
    assert _mevcut_otomatik_kayit_esigi(client, yetkili_header) == pytest.approx(0.97)


def test_kayit_ekle_otomatik_dusuk_guven_kayda_dusurulmez(client, yetkili_header):
    esik = _mevcut_otomatik_kayit_esigi(client, yetkili_header)
    dusuk_guven = max(0.0, esik - 0.20)
    r = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 DGV 01", "kamera_id": "TEST-DUSUK-GUVEN", "yon": "giris",
        "guven_skoru": dusuk_guven,
    })
    assert r.status_code == 200, r.text
    veri = r.json()
    assert veri["atlandi"] is True
    assert veri["sebep"] == "dusuk_guven_skoru"

    r2 = client.get("/kayitlar", params={"plaka": "34 DGV 01"}, headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    assert r2.json() == [], "Düşük güvenli tespit hiç kayda düşmemeliydi"


def test_kayit_ekle_otomatik_esik_ve_uzeri_kayda_dusurulur(client, yetkili_header):
    esik = _mevcut_otomatik_kayit_esigi(client, yetkili_header)
    r = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 YGV 02", "kamera_id": "TEST-YUKSEK-GUVEN", "yon": "giris",
        "guven_skoru": esik,  # eşiğe TAM eşit -- sınır dahil (>=) olmalı
    })
    assert r.status_code == 200, r.text
    veri = r.json()
    assert "atlandi" not in veri
    assert veri["plaka_no"] == "34 YGV 02"

    r2 = client.get("/kayitlar", params={"plaka": "34 YGV 02"}, headers=yetkili_header)
    assert len(r2.json()) == 1


# ------------------------------------------------------------------
# /kayitlar/otomatik — PTS_KAMERA_ANAHTARI başlık kontrolü
# (2026-09-21, GERÇEK ÜRETİMDE BULUNAN HATA: bkz. main.py::
# _kamera_anahtari_degeri'nin kök neden notu -- ortam değişkenine sona
# karışan görünmez bir satır sonu (\n), `requests` kütüphanesinin isteği
# HİÇ GÖNDERMEMESİNE yol açıyordu; yani gerçek kameraların doğruladığı HER
# plaka sessizce kayboluyordu. Bu testler hem normal başlık kontrolünü hem
# de .strip() ile bu sınıf hataya karşı bağışıklığı doğrular.)
# ------------------------------------------------------------------

def test_kayit_ekle_otomatik_kamera_anahtari_ayarliyken_basliksiz_401_doner(client, monkeypatch):
    monkeypatch.setenv("PTS_KAMERA_ANAHTARI", "gizli-kamera-anahtari")
    r = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 KAT 01", "kamera_id": "TEST-ANAHTAR-YOK", "yon": "giris",
    })
    assert r.status_code == 401, r.text


def test_kayit_ekle_otomatik_kamera_anahtari_dogru_basliktan_kabul_edilir(client, monkeypatch):
    monkeypatch.setenv("PTS_KAMERA_ANAHTARI", "gizli-kamera-anahtari")
    r = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 KAT 02", "kamera_id": "TEST-ANAHTAR-DOGRU", "yon": "giris",
    }, headers={"X-PTS-Kamera-Anahtari": "gizli-kamera-anahtari"})
    assert r.status_code == 200, r.text


def test_kayit_ekle_otomatik_kamera_anahtarindaki_sondaki_satir_sonu_sessizce_temizlenir(client, monkeypatch):
    """Kök neden testi: ortam değişkeninin (sunucu tarafı) VEYA gönderilen
    başlığın (kamera/CCTV tarafı) sonunda bir \n/boşluk olması, gerçek bir
    kamera pipeline'ının doğruladığı plakanın sessizce reddedilmesine
    (ya da hiç gönderilememesine) yol açmamalı."""
    monkeypatch.setenv("PTS_KAMERA_ANAHTARI", "gizli-kamera-anahtari\n")
    r = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 KAT 03", "kamera_id": "TEST-ANAHTAR-NEWLINE", "yon": "giris",
    }, headers={"X-PTS-Kamera-Anahtari": "gizli-kamera-anahtari"})
    assert r.status_code == 200, r.text


def test_kayit_ekle_otomatik_farkli_okuma_sayisi_kaydedilir_ve_dondurulur(client, yetkili_header):
    """2026-09-18: kamera pipeline'ının bu oturumda kaç FARKLI OCR metin
    varyantı gördüğü (bkz. camera_reader.py::PlakaOyBirikimi.kazanan) artık
    kayıtla birlikte gönderiliyor ve DB'ye kalıcı olarak yazılıyor -- panelin
    "6 kare, ama içlerinde çelişki vardı" durumunu "6 kare, tam oydaşma"
    durumundan ayırt edebilmesi için. Kamera pipeline'ından gelmeyen
    kayıtlarda (bu testte olmayan durum) None kalmalı; burada değeri
    doğrudan kontrol ediyoruz."""
    esik = _mevcut_otomatik_kayit_esigi(client, yetkili_header)
    r = client.post("/kayitlar/otomatik", data={
        "plaka_no": "02 AFP 552", "kamera_id": "TEST-CELISKILI-OKUMA", "yon": "giris",
        "guven_skoru": esik, "dogrulama_kare_sayisi": 6, "farkli_okuma_sayisi": 2,
    })
    assert r.status_code == 200, r.text
    veri = r.json()
    assert veri["dogrulama_kare_sayisi"] == 6
    assert veri["farkli_okuma_sayisi"] == 2

    r2 = client.get("/kayitlar", params={"plaka": "02 AFP 552"}, headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    kayitlar = r2.json()
    assert len(kayitlar) == 1
    assert kayitlar[0]["farkli_okuma_sayisi"] == 2


def test_kayit_ekle_manuel_farkli_okuma_sayisi_none_kalir(client, operator_header):
    """Manuel/elle eklenen kayıtlar (bkz. kayit_ekle_manuel) kamera
    pipeline'ından gelmediği için bu alan HİÇ gönderilmez -- None kalmalı,
    0 veya başka bir varsayılana sessizce düşmemeli."""
    r = client.post("/kayitlar", json={
        "plaka_no": "34 MNL 07", "kamera_id": "PANEL-MANUEL", "yon": "giris",
    }, headers=operator_header)
    assert r.status_code == 200, r.text
    assert r.json()["farkli_okuma_sayisi"] is None


# ------------------------------------------------------------------
# Çapraz kamera kısa süreli tekrarı (2026-09-18)
# ------------------------------------------------------------------
# Kullanıcı bildirimi: giriş kamerası bir aracı kaydettikten sonra, araç
# geçişine devam ederken çıkış kamerasının da görüş açısına girebiliyor --
# bu TEK bir fiziksel geçiş olmasına rağmen ikinci kamera bunu AYRI (yanlış
# yönde) bir kayıt olarak düşürmemeli. Bkz. main.py::
# _capraz_kamera_kisa_sureli_tekrar_mi.

def test_capraz_kamera_kisa_surede_farkli_kameradan_ayni_plaka_atlanir(client, yetkili_header):
    r1 = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 CKT 01", "kamera_id": "GIRIS-KAM", "yon": "giris", "guven_skoru": 0.99,
    })
    assert r1.status_code == 200, r1.text
    assert "atlandi" not in r1.json()

    r2 = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 CKT 01", "kamera_id": "CIKIS-KAM", "yon": "cikis", "guven_skoru": 0.99,
    })
    assert r2.status_code == 200, r2.text
    veri2 = r2.json()
    assert veri2["atlandi"] is True
    assert veri2["sebep"] == "capraz_kamera_kisa_sureli_tekrar"
    assert veri2["onceki_kamera_id"] == "GIRIS-KAM"

    r3 = client.get("/kayitlar", params={"plaka": "34 CKT 01"}, headers=yetkili_header)
    assert r3.status_code == 200, r3.text
    assert len(r3.json()) == 1, "Çapraz kamera tekrarı yine de İKİNCİ bir kayıt oluşturdu"


def test_capraz_kamera_ayni_kameradan_tekrar_engellenmez(client, yetkili_header):
    """Bu kontrol yalnızca FARKLI bir kamera_id için geçerlidir -- AYNI
    kameranın kendi tekrarını bastırmak camera_reader.py'nin (in-process,
    tekrar_gecikme_sn) işidir, backend bunu engellemez/engellememeli (aksi
    halde tek bir kameradan gelen MEŞRU art arda geçişler -- örn. giriş
    kamerasının kendisi -- de yanlışlıkla atlanırdı)."""
    r1 = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 CKT 02", "kamera_id": "GIRIS-KAM", "yon": "giris", "guven_skoru": 0.99,
    })
    assert r1.status_code == 200 and "atlandi" not in r1.json()

    r2 = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 CKT 02", "kamera_id": "GIRIS-KAM", "yon": "giris", "guven_skoru": 0.99,
    })
    assert r2.status_code == 200, r2.text
    assert "atlandi" not in r2.json(), "Aynı kameradan gelen kayıt yanlışlıkla çapraz-kamera tekrarı sayıldı"

    r3 = client.get("/kayitlar", params={"plaka": "34 CKT 02"}, headers=yetkili_header)
    assert len(r3.json()) == 2


def test_capraz_kamera_pencere_disindaki_eski_kayit_engellemez(client, yetkili_header):
    """Önceki kayıt yapılandırılmış pencereden (varsayılan 180 sn) daha
    ESKİYSE, bu artık aynı fiziksel geçiş sayılamayacak kadar uzun bir süre
    demektir -- yeni kayıt normal şekilde oluşturulmalı."""
    from backend.database import engine
    from sqlalchemy import text as sqltext

    with engine.connect() as conn:
        conn.execute(sqltext(
            "INSERT INTO plaka_kayitlari (plaka_no, kamera_id, yon, yetki_durumu, tarih_saat, manuel_giris) "
            "VALUES ('34 CKT 03', 'GIRIS-KAM', 'giris', 'yetkisiz', datetime('now', '-10 minutes'), 0)"
        ))
        conn.commit()

    r = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 CKT 03", "kamera_id": "CIKIS-KAM", "yon": "cikis", "guven_skoru": 0.99,
    })
    assert r.status_code == 200, r.text
    assert "atlandi" not in r.json(), "10 dakika önceki kayıt yanlışlıkla hâlâ 'yakın zamanlı' sayıldı"

    r2 = client.get("/kayitlar", params={"plaka": "34 CKT 03"}, headers=yetkili_header)
    assert len(r2.json()) == 2


def test_capraz_kamera_penceresi_sifirlanirsa_ozellik_kapanir(client, yetkili_header):
    """capraz_kamera_tekrar_penceresi_sn 0'a çekilirse özellik tamamen
    devre dışı kalmalı -- kullanıcı bunu istemezse tamamen kapatabilmeli."""
    r0 = client.put("/sistem/ayarlar", json={"capraz_kamera_tekrar_penceresi_sn": 0}, headers=yetkili_header)
    assert r0.status_code == 200, r0.text
    try:
        r1 = client.post("/kayitlar/otomatik", data={
            "plaka_no": "34 CKT 04", "kamera_id": "GIRIS-KAM", "yon": "giris", "guven_skoru": 0.99,
        })
        assert r1.status_code == 200 and "atlandi" not in r1.json()

        r2 = client.post("/kayitlar/otomatik", data={
            "plaka_no": "34 CKT 04", "kamera_id": "CIKIS-KAM", "yon": "cikis", "guven_skoru": 0.99,
        })
        assert r2.status_code == 200, r2.text
        assert "atlandi" not in r2.json(), "Pencere 0 iken özellik yine de devrede kaldı"
    finally:
        client.put("/sistem/ayarlar", json={"capraz_kamera_tekrar_penceresi_sn": 180}, headers=yetkili_header)


def test_sistem_ayarlari_degisikligi_loglanir(client, caplog, yetkili_header):
    """2026-09-20: sistem ayarları (ör. şüpheli-alarm eşiği gibi güvenlik
    açısından anlamlı bir değer) değiştirildiğinde önceden hiçbir log satırı
    yazılmıyordu -- "bu ayar ne zaman, kim tarafından değiştirildi" sorusu
    cevapsızdı (bkz. main.py::sistem_ayarlarini_guncelle). Ayrıca aynı
    değerin tekrar gönderilmesinin (gerçek bir değişiklik olmadığı için)
    log spam'i yaratmadığını da doğruluyoruz."""
    onceki = client.get("/sistem/ayarlar", headers=yetkili_header).json()["supheli_esik"]
    try:
        with caplog.at_level("INFO", logger="pts"):
            r = client.put("/sistem/ayarlar", json={"supheli_esik": onceki + 5}, headers=yetkili_header)
        assert r.status_code == 200, r.text
        assert f"supheli_esik: {onceki!r} -> {onceki + 5!r}" in caplog.text

        # 2026-09-20 (devam, kullanıcı talebi): log dosyasına EK olarak kalıcı
        # `denetim_kayitlari` tablosuna da yazılmalı.
        r_denetim = client.get("/denetim-kayitlari", params={"eylem": "sistem_ayarlari_guncelle"}, headers=yetkili_header)
        assert r_denetim.status_code == 200, r_denetim.text
        assert any(f"supheli_esik: {onceki!r} -> {onceki + 5!r}" in k["aciklama"] for k in r_denetim.json())

        caplog.clear()
        with caplog.at_level("INFO", logger="pts"):
            r2 = client.put("/sistem/ayarlar", json={"supheli_esik": onceki + 5}, headers=yetkili_header)
        assert r2.status_code == 200, r2.text
        assert "sistem_ayarlari_guncelle" not in caplog.text  # gerçek değişiklik yoksa log/denetim kaydı YOK
    finally:
        client.put("/sistem/ayarlar", json={"supheli_esik": onceki}, headers=yetkili_header)


# ------------------------------------------------------------------
# GET /denetim-kayitlari, GET /denetim-kayitlari/eylem-listesi (2026-09-20)
# ------------------------------------------------------------------
# Kullanıcı talebi: "kalıcı bir veritabanı tablosu + panelde ayrı bir Denetim
# Kayıtları ekranı olan tam bir audit-trail sistemi kurabilirim -- bunu
# yapabilirsin". Yukarıdaki *_loglanir testleri her bir çağrı noktasının
# (kullanici_guncelle/sil, kamera_sil, lisans_aktive_et, sistem_ayarlarini_
# guncelle) doğru şekilde kayıt oluşturduğunu doğruluyor; buradaki testler
# uç noktanın kendisinin (RBAC, filtreleme) doğru çalıştığını doğruluyor.

def test_denetim_kayitlari_girissiz_401_doner(client):
    r = client.get("/denetim-kayitlari")
    assert r.status_code == 401, r.text


def test_denetim_kayitlari_yalniz_yonetici_gorebilir(client, izleyici_header, operator_header, yetkili_header):
    """RBAC: hesap yönetimiyle ilgili hassas bilgi taşıdığı için (bkz.
    endpoint docstring'i) operatöre bile KAPALI olmalı -- yalnızca yönetici."""
    assert client.get("/denetim-kayitlari", headers=izleyici_header).status_code == 403
    assert client.get("/denetim-kayitlari", headers=operator_header).status_code == 403
    assert client.get("/denetim-kayitlari", headers=yetkili_header).status_code == 200


def test_kullanici_olusturma_da_denetim_kaydi_birakir(client, yetkili_header):
    """kullanici_ekle önceden de logluyordu (bkz. patch #50) ama YENİ
    denetim tablosuna da yazdığını doğrula (bkz. main.py::kullanici_ekle)."""
    r = client.post("/kullanicilar", json={
        "kullanici_adi": "denetim-olusturma-testi", "parola": "GucluParola123!", "rol": "izleyici",
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text
    r2 = client.get("/denetim-kayitlari", params={"eylem": "kullanici_olustur"}, headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    assert any("denetim-olusturma-testi" in k["aciklama"] for k in r2.json())


def test_denetim_kayitlari_kullanici_adina_gore_filtreler(client, yetkili_header):
    client.post("/kullanicilar", json={
        "kullanici_adi": "filtre-testi-a", "parola": "GucluParola123!", "rol": "izleyici",
    }, headers=yetkili_header)
    client.post("/kullanicilar", json={
        "kullanici_adi": "filtre-testi-b", "parola": "GucluParola123!", "rol": "izleyici",
    }, headers=yetkili_header)
    r = client.get("/denetim-kayitlari", params={"kullanici_adi": "filtre-testi-a"}, headers=yetkili_header)
    assert r.status_code == 200, r.text
    aciklamalar = [k["aciklama"] for k in r.json()]
    assert any("filtre-testi-a" in a for a in aciklamalar)
    assert not any("filtre-testi-b" in a for a in aciklamalar)


def test_denetim_kayitlari_negatif_limit_422_doner(client, yetkili_header):
    r = client.get("/denetim-kayitlari", params={"limit": -1}, headers=yetkili_header)
    assert r.status_code == 422, r.text


def test_denetim_eylem_listesi_sonuc_doner(client, yetkili_header):
    client.post("/kullanicilar", json={
        "kullanici_adi": "eylem-listesi-testi", "parola": "GucluParola123!", "rol": "izleyici",
    }, headers=yetkili_header)
    r = client.get("/denetim-kayitlari/eylem-listesi", headers=yetkili_header)
    assert r.status_code == 200, r.text
    assert "kullanici_olustur" in r.json()["eylemler"]


def test_denetim_eylem_listesi_operatore_kapali(client, operator_header):
    r = client.get("/denetim-kayitlari/eylem-listesi", headers=operator_header)
    assert r.status_code == 403, r.text


def test_capraz_kamera_manuel_kayitlari_hic_etkilemez(client, operator_header):
    """Görevlinin bilinçli olarak elle girdiği bir kayıt (kayit_ekle_manuel),
    çok yakın zamanda farklı bir kameradan aynı plaka otomatik kaydedilmiş
    olsa bile ASLA sessizce atlanmamalı -- operatör iradesi her zaman
    geçerli olmalı."""
    r1 = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 CKT 05", "kamera_id": "GIRIS-KAM", "yon": "giris", "guven_skoru": 0.99,
    })
    assert r1.status_code == 200 and "atlandi" not in r1.json()

    r2 = client.post("/kayitlar", json={
        "plaka_no": "34 CKT 05", "kamera_id": "CIKIS-KAM", "yon": "cikis",
    }, headers=operator_header)
    assert r2.status_code == 200, r2.text
    assert r2.json()["plaka_no"] == "34 CKT 05"
    assert r2.json()["manuel_giris"] is True


def test_kayit_ekle_otomatik_guven_skoru_gonderilmezse_filtrelenmez(client, yetkili_header):
    """Harici/eski entegrasyonlar güven skoru göndermeyebilir -- bu durumda
    filtre hiç uygulanmaz (geriye dönük uyumluluk), kayıt normal şekilde
    oluşur."""
    r = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 GYK 03", "kamera_id": "TEST-GUVENSIZ", "yon": "giris",
    })
    assert r.status_code == 200, r.text
    assert "atlandi" not in r.json()
    r2 = client.get("/kayitlar", params={"plaka": "34 GYK 03"}, headers=yetkili_header)
    assert len(r2.json()) == 1


def test_kayit_ekle_manuel_dusuk_guven_filtreden_etkilenmez(client, operator_header, yetkili_header):
    """Elle girilen (manuel_giris=True) kayıtlar, personelin bilinçli girişi
    olduğu için güven eşiği filtresinden HİÇ etkilenmemeli."""
    r = client.post("/kayitlar", json={
        "plaka_no": "34 MNL 04", "kamera_id": "TEST-MANUEL", "yon": "giris", "guven_skoru": 0.1,
    }, headers=operator_header)
    assert r.status_code == 200, r.text
    assert r.json()["plaka_no"] == "34 MNL 04"
    r2 = client.get("/kayitlar", params={"plaka": "34 MNL 04"}, headers=yetkili_header)
    assert len(r2.json()) == 1


# ------------------------------------------------------------------
# Otomatik kayıt güven eşiği filtresi — "bilinen araç" istisnası (2026-09-17, devam)
# ------------------------------------------------------------------
# Kök neden (kullanıcı bildirimi): kullanıcının kendi aracı girişte %96.6
# güvenle okundu ama genel eşik (varsayılan %97) altında kaldığı için o geçiş
# HİÇ KAYDEDİLMEDİ; aynı araç çıkışta (daha yüksek güvenle) kaydedildiği için
# giriş/çıkış kayıtları tutarsız hale geldi. Genel eşik, sistemde HİÇ KAYITLI
# OLMAYAN (yanlış okunmuş) plakaların kayıtları şişirmesini önlemek içindir;
# sahada zaten kayıtlı bir plakayla TAM/çok yakın eşleşen düşük güvenli bir
# tespit için artık daha düşük, ikinci bir "bilinen araç" eşiği uygulanıyor
# (bkz. main.py::_bilinen_plakaya_yakin_mi).

def _mevcut_otomatik_kayit_esigi_bilinen_arac(client, yetkili_header) -> float:
    r = client.get("/sistem/ayarlar", headers=yetkili_header)
    assert r.status_code == 200, r.text
    return float(r.json().get("otomatik_kayit_min_guven_skoru_bilinen_arac", 0.80))


def test_otomatik_kayit_bilinen_arac_esigi_varsayilan_080(client, yetkili_header):
    assert _mevcut_otomatik_kayit_esigi_bilinen_arac(client, yetkili_header) == pytest.approx(0.80)


def test_kayit_ekle_otomatik_bilinen_arac_dusuk_guvenle_de_kaydedilir(client, yetkili_header):
    """Genel eşiğin altında ama bilinen-araç eşiğinin üzerinde, sistemde
    kayıtlı bir plakayla TAM eşleşen bir tespit -- artık atlanmıyor, aynen
    kullanıcının bildirdiği canlı senaryodaki gibi."""
    r = client.post("/kisiler", json={
        "ad_soyad": "Bilinen Arac Testi", "plaka_no": "39 BG 262", "tip": "abone",
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text

    genel_esik = _mevcut_otomatik_kayit_esigi(client, yetkili_header)
    bilinen_esik = _mevcut_otomatik_kayit_esigi_bilinen_arac(client, yetkili_header)
    assert bilinen_esik < genel_esik, "Test, bilinen-araç eşiğinin genel eşikten düşük olduğunu varsayıyor"
    guven = (genel_esik + bilinen_esik) / 2

    r2 = client.post("/kayitlar/otomatik", data={
        "plaka_no": "39 BG 262", "kamera_id": "TEST-BILINEN-ARAC", "yon": "giris",
        "guven_skoru": guven,
    })
    assert r2.status_code == 200, r2.text
    veri = r2.json()
    assert "atlandi" not in veri, "Bilinen araç olduğu halde tespit yanlışlıkla atlandı"
    assert veri["plaka_no"] == "39 BG 262"

    r3 = client.get("/kayitlar", params={"plaka": "39 BG 262"}, headers=yetkili_header)
    assert r3.status_code == 200, r3.text
    assert len(r3.json()) == 1, "Bilinen araç, genel eşiğin altında kalsa bile kaydedilmeliydi"


def test_kayit_ekle_otomatik_bilinmeyen_arac_dusuk_guvenle_yine_atlanir(client, yetkili_header):
    """Bilinen-araç istisnası yalnızca SİSTEMDE KAYITLI bir plakayla eşleşen
    tespitler için geçerli -- alakasız/bilinmeyen bir plaka için genel eşik
    hâlâ olduğu gibi uygulanmalı, aksi halde istisna filtreyi anlamsız
    kılardı."""
    genel_esik = _mevcut_otomatik_kayit_esigi(client, yetkili_header)
    bilinen_esik = _mevcut_otomatik_kayit_esigi_bilinen_arac(client, yetkili_header)
    guven = (genel_esik + bilinen_esik) / 2

    r = client.post("/kayitlar/otomatik", data={
        "plaka_no": "06 XYZ 999", "kamera_id": "TEST-BILINMEYEN-ARAC", "yon": "giris",
        "guven_skoru": guven,
    })
    assert r.status_code == 200, r.text
    veri = r.json()
    assert veri["atlandi"] is True
    assert veri["sebep"] == "dusuk_guven_skoru"


def test_kayit_ekle_otomatik_bilinen_arac_esiginin_de_altinda_atlanir(client, yetkili_header):
    """Bilinen araç istisnası bir güvenlik tabanını atlamaz -- bilinen-araç
    eşiğinin de altındaki (çok düşük güvenli, muhtemelen gerçekten hatalı)
    bir tespit, plaka bilinen bir araca eşleşse bile yine atlanmalı."""
    r = client.post("/kisiler", json={
        "ad_soyad": "Cok Dusuk Guven Testi", "plaka_no": "34 CDG 555", "tip": "abone",
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text

    bilinen_esik = _mevcut_otomatik_kayit_esigi_bilinen_arac(client, yetkili_header)
    cok_dusuk = max(0.0, bilinen_esik - 0.30)

    r2 = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 CDG 555", "kamera_id": "TEST-COK-DUSUK-BILINEN", "yon": "giris",
        "guven_skoru": cok_dusuk,
    })
    assert r2.status_code == 200, r2.text
    veri = r2.json()
    assert veri["atlandi"] is True
    assert veri["sebep"] == "dusuk_guven_skoru"


def test_dusuk_guven_kayitlarini_temizle_yalnizca_otomatik_ve_dusuk_olanlari_siler(client, yetkili_header):
    """Geriye dönük temizlik uç noktası (bkz. main.py::dusuk_guven_kayitlarini_temizle):
    yalnızca guven_skoru dolu VE eşiğin altında VE manuel_giris=False olan
    kayıtları hedef almalı; yüksek güvenli ve elle girilmiş kayıtlara
    dokunmamalı."""
    from backend.database import SessionLocal
    from backend import models

    dusuk_dosya = _test_goruntu_dosyasi_olustur(client, "test-dusuk-guven.jpg")
    yuksek_dosya = _test_goruntu_dosyasi_olustur(client, "test-yuksek-guven.jpg")

    db = SessionLocal()
    try:
        dusuk = models.Kayit(
            plaka_no="34 TMZ 05", kamera_id="TEST-TEMIZLE", guven_skoru=0.3,
            goruntu_yolu=dusuk_dosya, manuel_giris=False,
        )
        yuksek = models.Kayit(
            plaka_no="34 TMZ 06", kamera_id="TEST-TEMIZLE", guven_skoru=0.99,
            goruntu_yolu=yuksek_dosya, manuel_giris=False,
        )
        manuel_dusuk = models.Kayit(
            plaka_no="34 TMZ 07", kamera_id="TEST-TEMIZLE", guven_skoru=0.1,
            manuel_giris=True,  # elle girilmiş -- ASLA silinmemeli
        )
        db.add_all([dusuk, yuksek, manuel_dusuk])
        db.commit()
        dusuk_id, yuksek_id, manuel_id = dusuk.id, yuksek.id, manuel_dusuk.id
    finally:
        db.close()

    try:
        r = client.post("/sistem/dusuk-guven-temizle", params={"esik": 0.97}, headers=yetkili_header)
        assert r.status_code == 200, r.text
        veri = r.json()
        assert veri["silinen_kayit"] >= 1
        assert veri["silinen_goruntu"] >= 1

        db2 = SessionLocal()
        try:
            assert db2.query(models.Kayit).filter(models.Kayit.id == dusuk_id).first() is None
            assert db2.query(models.Kayit).filter(models.Kayit.id == yuksek_id).first() is not None
            assert db2.query(models.Kayit).filter(models.Kayit.id == manuel_id).first() is not None
        finally:
            db2.close()
        assert not os.path.isfile(dusuk_dosya)
        assert os.path.isfile(yuksek_dosya)
    finally:
        for yol in (dusuk_dosya, yuksek_dosya):
            try:
                os.remove(yol)
            except OSError:
                pass
        db3 = SessionLocal()
        try:
            db3.query(models.Kayit).filter(
                models.Kayit.plaka_no.in_(["34 TMZ 05", "34 TMZ 06", "34 TMZ 07"])
            ).delete(synchronize_session=False)
            db3.commit()
        finally:
            db3.close()


def test_dusuk_guven_temizle_operator_yetkisiz_403_doner(client, operator_header):
    """Bu toplu/kalıcı silme işlemi kayit_sil ile aynı sıkı kısıtlamayı
    taşır: yalnızca yönetici çalıştırabilir (operatör dahi değil)."""
    r = client.post("/sistem/dusuk-guven-temizle", headers=operator_header)
    assert r.status_code == 403, r.text


def test_dusuk_guven_temizle_izleyici_yetkisiz_403_doner(client, izleyici_header):
    r = client.post("/sistem/dusuk-guven-temizle", headers=izleyici_header)
    assert r.status_code == 403, r.text


# ------------------------------------------------------------------
# Dışa aktarma (Excel/PDF) uç noktaları -- kimlik doğrulama (2026-09-18)
# ------------------------------------------------------------------
# GÜVENLİK KÖK NEDEN DÜZELTMESİ: bu uç noktaların DÖRDÜ de (artı yeni eklenen
# toplu içe aktarma şablonu) önceden HİÇBİR kimlik doğrulama gerektirmiyordu
# -- frontend bunları `window.open()` ile açtığı için Authorization başlığı
# taşıyamıyordu, bu yüzden hiç eklenmemişti. Sonuç: KVKK kapsamındaki plaka/
# ad-soyad/daire/görsel gibi kişisel verileri girişsiz HERHANGİ bir istemciye
# açık ediyordu. Artık `_giris_gerekli` hem Authorization başlığını hem de
# (SADECE başlık yoksa, geriye dönük uyumluluk için değil -- window.open()
# yolunun TEK çalışma biçimi olarak) bir `?token=` sorgu parametresini kabul
# ediyor (bkz. frontend/app.js::indirmeUrlOlustur). Aşağıdaki testler hem
# tokensiz isteğin reddedildiğini hem de her iki yolun (başlık VE sorgu
# parametresi) çalıştığını doğrular.

_DISA_AKTAR_SABIT_YOLLAR = [
    "/disa-aktar/excel/kayitlar",
    "/disa-aktar/pdf/kayitlar",
    "/disa-aktar/excel/kisiler",
    "/kisiler/toplu-import/sablon",
]


def test_disa_aktar_uc_noktalari_tokensiz_401_doner(client):
    for yol in _DISA_AKTAR_SABIT_YOLLAR:
        r = client.get(yol)
        assert r.status_code == 401, f"GET {yol}: token/başlık olmadan 401 beklenirdi, {r.status_code} alındı ({r.text})"


def test_disa_aktar_uc_noktalari_authorization_basligiyla_calisir(client, yetkili_header):
    for yol in _DISA_AKTAR_SABIT_YOLLAR:
        r = client.get(yol, headers=yetkili_header)
        assert r.status_code == 200, f"GET {yol}: Authorization başlığıyla 200 beklenirdi ({r.text})"


def test_disa_aktar_kayitlar_offset_query_nesnesi_olarak_sizmaz(client, yetkili_header):
    """KÖK NEDEN (2026-09-21, kullanıcının paylaştığı GERÇEK hata iziyle
    bulundu): `kayitlari_excel_indir`/`kayitlari_pdf_indir`,
    `kayitlari_listele()`'yi FastAPI'nin DI mekanizması ÜZERİNDEN DEĞİL
    doğrudan bir Python fonksiyonu olarak çağırıyor (bkz. her ikisindeki
    "kullanici AÇIKÇA geçirilmeli" notu) -- ama `offset` parametresi
    AÇIKÇA geçirilmiyordu. `kayitlari_listele`'nin imzasındaki
    `offset: int = Query(0, ge=0)` yalnızca FastAPI'nin kendi routing
    katmanından çağrıldığında gerçek bir tam sayıya çözülür; düz bir
    Python çağrısında `offset`, FastAPI'nin `Query` sınıfının bir örneği
    olarak KALIR. Bu, `.offset(offset)` satırında SQLAlchemy içinde
    yakalanmamış bir `TypeError` fırlatıp HER Excel/PDF dışa aktarma
    isteğini (istek kendisi -- kimlik doğrulaması, veri, hepsi -- tamamen
    geçerli olsa bile) genel 500'e düşürüyordu. Bu test özellikle EN AZ
    bir kayıt varken bu iki uç noktayı çağırır (`.offset()`/`.limit()`
    HER çağrıda -- kayıt sayısından bağımsız olarak -- işletilir, bu
    yüzden aslında sıfır kayıtla da tetiklenirdi, ama gerçekçi bir veri
    kümesiyle test etmek daha az kırılgan)."""
    olusturulan = client.post("/kayitlar", json={
        "plaka_no": "34 DAO 001", "kamera_id": "TEST-DISA-AKTAR-OFFSET", "yon": "giris",
    }, headers=yetkili_header)
    assert olusturulan.status_code == 200, olusturulan.text

    for yol in ("/disa-aktar/excel/kayitlar", "/disa-aktar/pdf/kayitlar"):
        r = client.get(yol, headers=yetkili_header)
        assert r.status_code == 200, (
            f"GET {yol}: 200 beklenirdi ama {r.status_code} döndü -- `offset` parametresi "
            f"kayitlari_listele()'ye açıkça geçirilmiyor olabilir (bkz. kök neden notu), {r.text}"
        )


def test_disa_aktar_uc_noktalari_sorgu_token_ile_calisir(client, admin_token):
    """window.open() ile açılan indirme bağlantıları Authorization başlığı
    TAŞIYAMAZ -- bu yüzden frontend token'ı ?token= sorgu parametresi olarak
    ekliyor. Bu istekler HİÇ Authorization başlığı olmadan da çalışmalı."""
    for yol in _DISA_AKTAR_SABIT_YOLLAR:
        r = client.get(yol, params={"token": admin_token})
        assert r.status_code == 200, f"GET {yol}: sorgu token ile 200 beklenirdi ({r.text})"


def test_disa_aktar_pdf_kayit_detay_tokensiz_401_ve_token_ile_calisir(client, yetkili_header, admin_token):
    """Tek kayıt detay PDF'i, yola gömülü bir id parametresi aldığı için
    yukarıdaki sabit yol listesine uymuyor -- ayrı test edilir."""
    olusturulan = client.post("/kayitlar", json={
        "plaka_no": "34 EXP 001", "kamera_id": "TEST-DISA-AKTAR", "yon": "giris",
    }, headers=yetkili_header)
    assert olusturulan.status_code == 200, olusturulan.text
    kayit_id = olusturulan.json()["id"]

    r = client.get(f"/disa-aktar/pdf/kayit/{kayit_id}")
    assert r.status_code == 401, r.text
    r2 = client.get(f"/disa-aktar/pdf/kayit/{kayit_id}", headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    r3 = client.get(f"/disa-aktar/pdf/kayit/{kayit_id}", params={"token": admin_token})
    assert r3.status_code == 200, r3.text


def test_kisi_ice_aktarma_sablonu_gecerli_ve_dogru_basliklara_sahip_xlsx_doner(client, yetkili_header):
    """Şablon uç noktasının döndürdüğü dosya gerçekten açılabilir bir .xlsx
    olmalı ve toplu_kisi_import'un beklediği sütunlarla eşleşmeli (bkz.
    tests/test_excel_export.py -- burada aynı doğrulama, birim testinden
    farklı olarak GERÇEK HTTP uç noktası üzerinden yapılır)."""
    import io
    import openpyxl

    r = client.get("/kisiler/toplu-import/sablon", headers=yetkili_header)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert "Kişiler" in wb.sheetnames
    assert "Açıklama" in wb.sheetnames
    basliklar_ham = [h.value for h in wb["Kişiler"][1]]
    basliklar = [str(b).strip().lower().replace(" ", "_") if b else "" for b in basliklar_ham]
    for zorunlu in ("ad_soyad", "plaka_no", "tip"):
        assert zorunlu in basliklar, f"'{zorunlu}' şablonda eksik -- içe aktarma her zaman başarısız olurdu"


# ------------------------------------------------------------------
# Kayıtlar dışa aktarma raporu -- Site/Daire/Araç Tipi zenginleştirmesi
# (2026-09-18, bkz. backend/main.py::_kayitlari_rapor_satirlari)
# ------------------------------------------------------------------
# Kullanıcının paylaştığı referans "GEÇİŞ RAPORU" biçimine yaklaştırma kararı
# (bkz. README'deki 2026-09-18 notu). Bu testler gerçek bir Site + Nokta +
# personel/abone Kişi + kara liste kaydı oluşturup, dışa aktarılan Excel'in
# Site/Daire/Araç Tipi sütunlarının belgelenen sınıflandırma kurallarıyla
# eşleştiğini UÇTAN UCA (gerçek HTTP + gerçek DB) doğrular.
#
# NOT: aşağıdaki plakalar kasıtlı olarak birbirinden çok FARKLI seçildi (tek
# karakter farkıyla eşleşmiyorlar) -- aksi halde _bilinen_plakaya_yakinlik_
# duzelt (OCR yakınlık düzeltmesi, guven_skoru verilmediği için burada da
# devreye girer) "tanımsız araç" olması gereken bir plakayı yanlışlıkla
# bilinen bir plakaya düzeltip testi anlamsızlaştırabilirdi.

_RAPOR_TEST_PLAKA_PERSONEL = "77RAP9001"
_RAPOR_TEST_PLAKA_ABONE = "88RAP9002"
_RAPOR_TEST_PLAKA_TANIMSIZ = "12TANIMSIZ99"
_RAPOR_TEST_PLAKA_KARALISTE = "34KARALISTE01"


@pytest.fixture
def rapor_test_kurulumu(client, yetkili_header):
    """Site + Nokta'yı doğrudan DB'ye yazar, ama Nokta.kamera_id ve
    Kayit.kamera_id için GERÇEK bir kamera üzerinden GERÇEK "id"/"ad"
    çiftini kullanır (bkz. kamera_erisim_test_kameralari'nin docstring'i,
    2026-09-21 "id vs ad" hata sınıfı): Nokta.kamera_id kamera "id"si
    (cameras.json), Kayit.kamera_id ise kameranın "ad"ı olmalı -- ikisi
    KASITLI olarak farklı tutulur. Bu test eskiden "KAM-RAPTEST" adında
    UYDURMA bir dizeyi HEM Nokta.kamera_id HEM DE Kayit.kamera_id için
    kullanıyordu -- id ile ad'ın TESADÜFEN aynı olduğu bu senaryo,
    _kayitlari_rapor_satirlari'ndaki gerçek id/ad karışıklığı hatasını
    (Nokta/Site sütunlarının HER ZAMAN boş kalması) maskeliyordu."""
    from backend.database import SessionLocal
    from backend import models

    anahtar = lisans.uret("Rapor Test Site", kamera_limiti=10, gun=30)
    r = client.post("/lisans/aktive-et", json={"anahtar": anahtar}, headers=yetkili_header)
    assert r.status_code == 200, r.text
    kamera_ad = "RAPTEST KAMERA"
    r = client.post("/kameralar", json={
        "ad": kamera_ad, "rtsp_url": "rtsp://127.0.0.1/raptest", "yon": "giris",
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text
    kamera_id = r.json()["id"]
    assert kamera_id != kamera_ad, (
        "id, UUID olduğu için 'ad' ile ASLA aynı olmamalı -- aksi halde bu fixture "
        "id/ad karışıklığı hatasını maskeler"
    )

    db = SessionLocal()
    try:
        site = models.Site(ad="RAPTEST SİTESİ")
        db.add(site)
        db.commit()
        db.refresh(site)
        nokta = models.Nokta(site_id=site.id, ad="RAPTEST GİRİŞ", yon="giris", kamera_id=kamera_id)
        db.add(nokta)
        db.commit()
    finally:
        db.close()

    r = client.post("/kisiler", json={
        "ad_soyad": "Ahmet Yılmaz", "plaka_no": _RAPOR_TEST_PLAKA_PERSONEL, "tip": "personel",
        "daire_departman": "GÜVENLİK ŞEFLİĞİ",
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text

    r = client.post("/kisiler", json={
        "ad_soyad": "Mehmet Demir", "plaka_no": _RAPOR_TEST_PLAKA_ABONE, "tip": "abone",
        "daire_departman": "A BLOK NO 5",
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text

    r = client.post("/kara-listesi", json={
        "plaka_no": _RAPOR_TEST_PLAKA_KARALISTE, "sebep": "test",
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text

    for plaka in (
        _RAPOR_TEST_PLAKA_PERSONEL, _RAPOR_TEST_PLAKA_ABONE,
        _RAPOR_TEST_PLAKA_TANIMSIZ, _RAPOR_TEST_PLAKA_KARALISTE,
    ):
        # ÖNEMLİ: gerçek pipeline'ın damgaladığı gibi (kamera_id=kamera["ad"])
        # burada da kamera_id "ad" değeri olarak gönderiliyor, id DEĞİL.
        r = client.post("/kayitlar", json={
            "plaka_no": plaka, "kamera_id": kamera_ad, "yon": "giris",
        }, headers=yetkili_header)
        assert r.status_code == 200, r.text


def _rapor_satirini_getir(client, yetkili_header, plaka: str) -> dict:
    """Kayıtlar Excel raporunu tek bir plakayla filtreleyip tek satırı, sütun
    adı -> değer sözlüğü olarak döner."""
    import io
    import openpyxl

    r = client.get("/disa-aktar/excel/kayitlar", params={"plaka": plaka}, headers=yetkili_header)
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    basliklar = [h.value for h in ws[1]]
    satirlar = [dict(zip(basliklar, satir)) for satir in ws.iter_rows(min_row=2, values_only=True)]
    assert len(satirlar) == 1, f"'{plaka}' için tam olarak 1 satır bekleniyordu, {len(satirlar)} bulundu"
    return satirlar[0]


def test_kayitlar_raporu_personel_satiri_daire_ve_arac_tipini_dogru_doldurur(client, yetkili_header, rapor_test_kurulumu):
    satir = _rapor_satirini_getir(client, yetkili_header, _RAPOR_TEST_PLAKA_PERSONEL)
    assert satir["Site"] == "RAPTEST SİTESİ"
    assert satir["Nokta"] == "RAPTEST GİRİŞ"
    assert satir["Adı"] == "Ahmet"
    assert satir["Soyadı"] == "Yılmaz"
    assert satir["Daire"] == "PERSONEL", "personel tipi için Daire sütunu her zaman sabit 'PERSONEL' olmalı"
    assert satir["Araç Tipi"] == "GÜVENLİK ŞEFLİĞİ", "personel için Araç Tipi kişinin departman adını göstermeli"
    assert satir["Geçiş Tipi"] == "Giriş"
    assert satir["Blok"] is None, "Blok kavramı bu sistemde yok -- UYDURULMAMALI, her zaman boş kalmalı"
    assert satir["Otopark"] is None, "Otopark kavramı bu sistemde yok -- UYDURULMAMALI, her zaman boş kalmalı"


def test_kayitlar_raporu_abone_satiri_daireyi_ve_tanimli_arac_tipini_dogru_doldurur(client, yetkili_header, rapor_test_kurulumu):
    satir = _rapor_satirini_getir(client, yetkili_header, _RAPOR_TEST_PLAKA_ABONE)
    assert satir["Site"] == "RAPTEST SİTESİ"
    assert satir["Daire"] == "A BLOK NO 5", "abone tipi için Daire, kişinin kendi daire_departman değeri olmalı"
    assert satir["Araç Tipi"] == "Tanımlı"


def test_kayitlar_raporu_taninmayan_plaka_tanimsiz_arac_olarak_isaretlenir(client, yetkili_header, rapor_test_kurulumu):
    satir = _rapor_satirini_getir(client, yetkili_header, _RAPOR_TEST_PLAKA_TANIMSIZ)
    assert satir["Araç Tipi"] == "Tanımsız Araç"
    assert satir["Daire"] is None
    assert satir["Adı"] is None
    assert satir["Soyadı"] is None


def test_kayitlar_raporu_kara_liste_plakasi_ayrica_isaretlenir(client, yetkili_header, rapor_test_kurulumu):
    """Referans raporda yok ama sistemin zaten tuttuğu bir bilgi -- 'Tanımsız
    Araç' içinde gizlenmek yerine ayrıca 'Kara Liste' olarak gösterilir."""
    satir = _rapor_satirini_getir(client, yetkili_header, _RAPOR_TEST_PLAKA_KARALISTE)
    assert satir["Araç Tipi"] == "Kara Liste"
    assert satir["Daire"] is None


# ------------------------------------------------------------------
# Geçmiş kayıtları yeni eklenen/düzenlenen kişiye bağlama (2026-09-18)
# ------------------------------------------------------------------
# Kullanıcı bildirimi: "39 AEZ 645" aracını gerçek kamerayla test etti (araç
# o an sistemde tanımlı değildi -> kayıtlar "yetkisiz" düştü), SONRA aracı
# personel olarak kaydetti. Yeni tespitler doğru şekilde "yetkili"
# gösteriliyordu (bu zaten çalışıyordu -- ayrı bir araştırmayla doğrulandı,
# kamera pipeline'ında hiçbir "bilinen plaka" önbelleği yok, her tespit canlı
# DB sorgusu kullanıyor), ama kayıttan ÖNCEKİ eski tespitler "yetkisiz"
# olarak donmuş kalıyordu. Kullanıcı bunların da düzeltilebilmesini istedi.

def _plaka_ile_manuel_kayit_olustur(client, yetkili_header, plaka: str, kamera_id: str = "TEST-GECMIS") -> dict:
    r = client.post("/kayitlar", json={"plaka_no": plaka, "kamera_id": kamera_id, "yon": "giris"}, headers=yetkili_header)
    assert r.status_code == 200, r.text
    return r.json()


def test_kisi_eklenince_gecmis_yetkisiz_kayitlar_otomatik_yetkiliye_donusur(client, yetkili_header):
    """Kök senaryo: plaka önce bilinmiyorken kayıt oluşur (yetkisiz), SONRA
    aynı plaka personel olarak eklenir -- kisi_ekle bu eski kaydı OTOMATİK
    olarak günceller (bkz. backend/main.py::_gecmis_kayitlari_kisiye_bagla)."""
    plaka = "77GECMIS01"
    eski_kayit = _plaka_ile_manuel_kayit_olustur(client, yetkili_header, plaka)
    assert eski_kayit["yetki_durumu"] == "yetkisiz"
    assert eski_kayit["kisi_id"] is None

    r = client.post("/kisiler", json={
        "ad_soyad": "Test Personel", "plaka_no": plaka, "tip": "personel",
        "daire_departman": "TEST DEPARTMANI",
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text
    kisi_id = r.json()["id"]

    r2 = client.get("/kayitlar", params={"plaka": plaka}, headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    guncellenmis = next(k for k in r2.json() if k["id"] == eski_kayit["id"])
    assert guncellenmis["yetki_durumu"] == "yetkili", "kişi eklendikten sonra eski kayıt otomatik güncellenmedi"
    assert guncellenmis["kisi_id"] == kisi_id
    assert guncellenmis["kisi_tip_anlik"] == "personel"
    assert guncellenmis["duzenleyen"] == "admin", "otomatik düzeltme de denetlenebilirlik için 'duzenleyen' alanını doldurmalı"


def test_gecmis_kayitlari_guncelle_uc_noktasi_elle_tetiklenebilir_ve_idempotenttir(client, yetkili_header):
    """Bu özellik eklenmeden ÖNCE kaydedilmiş kişiler için panelden elle
    tetiklenen uç nokta (bkz. frontend/app.js::kisiGecmisKayitlariGuncelle)."""
    plaka = "77GECMIS02"
    eski_kayit = _plaka_ile_manuel_kayit_olustur(client, yetkili_header, plaka)

    r = client.post("/kisiler", json={
        "ad_soyad": "Test Abone", "plaka_no": plaka, "tip": "abone",
    }, headers=yetkili_header)
    kisi_id = r.json()["id"]

    # kisi_ekle zaten otomatik güncellemiş olmalı (yukarıdaki test bunu
    # doğruluyor) -- bu yüzden burada elle tetiklenen çağrı 0 dönmeli
    # (idempotentlik: zaten bağlı bir kaydı ikinci kez "güncellenen" olarak
    # saymamalı).
    r2 = client.post(f"/kisiler/{kisi_id}/gecmis-kayitlari-guncelle", headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    assert r2.json()["guncellenen_kayit_sayisi"] == 0


def test_gecmis_kayitlari_guncelle_kara_listedeki_plakayi_yetkiliye_cevirmez(client, yetkili_header):
    """GÜVENLİK: bir plaka HEM kara listede HEM de (çelişkili biçimde) bir
    Kişi'ye bağlıysa, kara liste her zaman ÖNCELİKLİDİR (bkz.
    _plaka_yetki_kontrol). Geçmiş kayıtları yeniden değerlendirme özelliği
    bu önceliği BOZMAMALI -- yani bir güvenlik engelini sessizce "yetkili"ye
    çevirmemeli."""
    plaka = "77GECMIS03"
    eski_kayit = _plaka_ile_manuel_kayit_olustur(client, yetkili_header, plaka)
    assert eski_kayit["yetki_durumu"] == "yetkisiz"

    r = client.post("/kara-listesi", json={"plaka_no": plaka, "sebep": "test"}, headers=yetkili_header)
    assert r.status_code == 200, r.text

    r2 = client.post("/kisiler", json={
        "ad_soyad": "Çelişkili Kayıt", "plaka_no": plaka, "tip": "personel",
    }, headers=yetkili_header)
    assert r2.status_code == 200, r2.text

    r3 = client.get("/kayitlar", params={"plaka": plaka}, headers=yetkili_header)
    hala_eski = next(k for k in r3.json() if k["id"] == eski_kayit["id"])
    assert hala_eski["yetki_durumu"] == "yetkisiz", (
        "kara listedeki bir plaka, kişi eklendi diye yanlışlıkla yetkiliye çevrildi -- güvenlik regresyonu"
    )
    assert hala_eski["kisi_id"] is None


def test_gecmis_kayitlari_guncelle_saat_kisitlamasina_uymayan_eski_kaydi_atlar(client, yetkili_header):
    """DOĞRULUK: saat kısıtlaması olan bir personel eklendiğinde, geçmişte bu
    kısıtlamaya UYMAYAN bir saatte oluşmuş eski kayıt "şimdi"ye göre değil
    KENDİ ORİJİNAL saatine göre değerlendirilmeli -- aksi halde örn. yalnızca
    08:00-18:00 arası yetkili bir personelin GECE geçmiş eski bir kaydı,
    güncelleme "şimdi" gündüzse yanlışlıkla yetkili işaretlenirdi."""
    from backend.database import SessionLocal
    from backend import models

    plaka = "77GECMIS04"
    db = SessionLocal()
    try:
        gece_kaydi = models.Kayit(
            plaka_no=plaka, kamera_id="TEST-GECMIS", yon="giris",
            yetki_durumu="yetkisiz", kisi_id=None,
            tarih_saat=datetime(2026, 9, 1, 23, 0, 0),  # 23:00 -- 08-18 dışında
        )
        db.add(gece_kaydi)
        db.commit()
        gece_kaydi_id = gece_kaydi.id
    finally:
        db.close()

    r = client.post("/kisiler", json={
        "ad_soyad": "Saatli Personel", "plaka_no": plaka, "tip": "personel",
        "giris_saati_baslangic": "08:00", "giris_saati_bitis": "18:00",
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text

    db2 = SessionLocal()
    try:
        guncel = db2.query(models.Kayit).filter(models.Kayit.id == gece_kaydi_id).first()
        assert guncel.yetki_durumu == "yetkisiz", (
            "saat kısıtlamasına uymayan eski bir kayıt yanlışlıkla yetkili yapıldı -- "
            "'referans_zaman' yerine 'şimdi' kullanılmış olabilir"
        )
        assert guncel.kisi_id is None
    finally:
        db2.close()


def test_gecmis_kayitlari_guncelle_zaten_baska_kisiye_bagli_satiri_ezmiyor(client, yetkili_header):
    """Bir Kayıt satırı zaten (yetki_durumu ne olursa olsun) BAŞKA bir
    kisi_id'ye bağlıysa (ör. daha önce elle düzenlenmiş), yeni eklenen bir
    kişi bu bağlantıyı sessizce ÜZERİNE YAZMAMALI -- yalnızca kisi_id IS NULL
    olan satırlar aday olarak değerlendirilir."""
    from backend.database import SessionLocal
    from backend import models

    plaka = "77GECMIS05"
    r = client.post("/kisiler", json={
        "ad_soyad": "Diğer Kişi", "plaka_no": "77GECMISDIGER", "tip": "abone",
    }, headers=yetkili_header)
    diger_kisi_id = r.json()["id"]

    db = SessionLocal()
    try:
        bagli_kayit = models.Kayit(
            plaka_no=plaka, kamera_id="TEST-GECMIS", yon="giris",
            yetki_durumu="yetkisiz", kisi_id=diger_kisi_id,
        )
        db.add(bagli_kayit)
        db.commit()
        bagli_kayit_id = bagli_kayit.id
    finally:
        db.close()

    r2 = client.post("/kisiler", json={
        "ad_soyad": "Yeni Personel", "plaka_no": plaka, "tip": "personel",
    }, headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    yeni_kisi_id = r2.json()["id"]

    db2 = SessionLocal()
    try:
        guncel = db2.query(models.Kayit).filter(models.Kayit.id == bagli_kayit_id).first()
        assert guncel.kisi_id == diger_kisi_id, "zaten başka bir kişiye bağlı kayıt sessizce üzerine yazıldı"
    finally:
        db2.close()
    assert yeni_kisi_id != diger_kisi_id


def test_gecmis_kayitlari_guncelle_izleyici_yetkisiz_403_doner(client, izleyici_header):
    r = client.post("/kisiler/1/gecmis-kayitlari-guncelle", headers=izleyici_header)
    assert r.status_code == 403, r.text


def test_gecmis_kayitlari_guncelle_olmayan_kisi_404_doner(client, yetkili_header):
    r = client.post("/kisiler/999999/gecmis-kayitlari-guncelle", headers=yetkili_header)
    assert r.status_code == 404, r.text


# ------------------------------------------------------------------
# Güvenlik Personeli Vardiya Filtresi (2026-09-18)
# ------------------------------------------------------------------
# Kullanıcı talebi (özetle): güvenlik personeli "operatör" ile AYNI ekranlara
# erişebilmeli ama Kayıtlar listesi/raporları yalnızca KENDİ vardiyasında
# geçen araçları göstermeli; vardiya saatleri günden güne değişebildiği için
# (4 vardiya/vardiya amiri rotasyonu) GÜN BAZLI atanır; çakışan vardiyalarda
# bir kayıt İKİ kullanıcının da listesinde ayrı ayrı görünebilir (dışlayıcı
# değil). Aşağıdaki testler bu üç kararı doğrudan doğrular.
#
# Zamanlamayla ilgili testler (özellikle gece yarısını geçen vardiya testi),
# testin çalıştığı GERÇEK saatten bağımsız olarak HER ZAMAN doğru sonucu
# üretecek şekilde (ör. "00:00 -> 00:00" tam gün penceresi, ya da "-25 saat"
# geri tarihleme) bilinçli olarak kurgulandı -- bkz. her testin kendi notu.

def _rbac_guvenlik_kullanici_olustur(client, yetkili_header) -> tuple:
    """Yeni, benzersiz adlı bir 'güvenlik' rolünde kullanıcı oluşturur, giriş
    yapar ve (kullanici_id, header) döner. Fonksiyon-scope'lu bir fixture
    yerine yardımcı fonksiyon olarak tanımlandı ki her çağrı BAĞIMSIZ bir
    kullanıcı versin (testler birbirini etkilemesin).

    NOT (2026-09-20, öz-hizmet vardiya sistemine geçiş): bu çağrı, `/auth/
    giris`'in bir yan etkisi olarak ARTIK OTOMATİK olarak TAM OLARAK BİR açık
    (cikis_zamani=None, giris_zamani=çağrı anı) VardiyaOturumu satırı da
    oluşturur (bkz. main.py::_guvenlik_oturum_baslat) -- yani dönen kullanıcı
    "sıfır vardiyalı" DEĞİLDİR, çağrı anından itibaren AÇIK bir vardiyası
    vardır. Bu ANDAN ÖNCEKİ kayıtlar (bkz. test_guvenlik_oturum_acilmadan_
    onceki_kaydi_gormez) hâlâ görünmez."""
    import uuid
    kullanici_adi = f"rbac-guvenlik-{uuid.uuid4().hex[:10]}"
    r = client.post("/kullanicilar", json={
        "kullanici_adi": kullanici_adi, "parola": "GucluParola123!", "rol": "güvenlik",
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text
    kullanici_id = r.json()["id"]
    r2 = client.post("/auth/giris", json={"kullanici_adi": kullanici_adi, "parola": "GucluParola123!"})
    assert r2.status_code == 200, r2.text
    return kullanici_id, {"Authorization": f"Bearer {r2.json()['token']}"}


def test_guvenlik_kullanicisi_operator_yetkilerine_sahip_ama_yonetici_islemi_yapamaz(client, yetkili_header):
    _, guvenlik_header = _rbac_guvenlik_kullanici_olustur(client, yetkili_header)
    # operatör-eşdeğeri: manuel kayıt ekleyebilmeli (_rol_dogrula(YONETICI, OPERATOR) gerektirir).
    r1 = client.post("/kayitlar", json={
        "plaka_no": "34 GVN 01", "kamera_id": "PANEL-MANUEL", "yon": "giris",
    }, headers=guvenlik_header)
    assert r1.status_code == 200, r1.text
    # yönetici-only bir işlemi YAPAMAMALI (kullanıcı oluşturma).
    r2 = client.post("/kullanicilar", json={
        "kullanici_adi": "guvenlik-baskasini-olusturamaz", "parola": "GucluParola123!", "rol": "izleyici",
    }, headers=guvenlik_header)
    assert r2.status_code == 403, r2.text


def test_vardiya_oturumu_giriste_otomatik_acilir_ve_yalnizca_yonetici_gorup_sonlandirabilir(
    client, yetkili_header, operator_header
):
    """Öz-hizmet sistemde vardiya oturumu ELLE oluşturulmaz -- güvenlik
    rolündeki bir kullanıcı `/auth/giris` ile giriş yapar yapmaz OTOMATİK
    olarak AÇIK (cikis_zamani=None) bir VardiyaOturumu satırı oluşturulmalı
    (bkz. main.py::_guvenlik_oturum_baslat). Bu oturumları yalnızca yönetici
    listeleyebilir/sonlandırabilir (bkz. /vardiya-oturumlari)."""
    guvenlik_id, _ = _rbac_guvenlik_kullanici_olustur(client, yetkili_header)

    assert client.get("/vardiya-oturumlari", headers=operator_header).status_code == 403

    r1 = client.get("/vardiya-oturumlari", params={"kullanici_id": guvenlik_id}, headers=yetkili_header)
    assert r1.status_code == 200, r1.text
    oturumlar = r1.json()
    assert len(oturumlar) == 1, "giriş otomatik olarak TAM OLARAK bir oturum açmalı"
    assert oturumlar[0]["cikis_zamani"] is None, "yeni açılan oturum henüz kapalı olmamalı"
    oturum_id = oturumlar[0]["id"]

    assert client.post(f"/vardiya-oturumlari/{oturum_id}/sonlandir", headers=operator_header).status_code == 403

    r2 = client.post(f"/vardiya-oturumlari/{oturum_id}/sonlandir", headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    r3 = client.get("/vardiya-oturumlari", params={"kullanici_id": guvenlik_id}, headers=yetkili_header)
    assert r3.json()[0]["cikis_zamani"] is not None

    # Zaten kapalı bir oturumu tekrar sonlandırmaya çalışmak 400 dönmeli.
    r4 = client.post(f"/vardiya-oturumlari/{oturum_id}/sonlandir", headers=yetkili_header)
    assert r4.status_code == 400, r4.text


def test_guvenlik_oturum_acilmadan_onceki_kaydi_gormez(client, yetkili_header):
    """Fail-closed'ın öz-hizmet sistemdeki karşılığı: vardiya penceresinin
    ALT SINIRI giriş anıdır -- güvenlik kullanıcısının oturumu açılmadan
    (giriş yapmadan) ÖNCE gerçekleşmiş bir kayıt, o oturumun penceresine
    denk düşmediği için görünmemelidir."""
    r1 = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 GVN 02", "kamera_id": "GIRIS-KAM", "yon": "giris", "guven_skoru": 0.99,
    })
    assert r1.status_code == 200, r1.text
    # Kayıt oluşturulduktan SONRA güvenlik kullanıcısı giriş yapıyor -- bu
    # yüzden yeni açılan oturumun penceresi bu kayıttan SONRA başlar.
    _, guvenlik_header = _rbac_guvenlik_kullanici_olustur(client, yetkili_header)

    r2 = client.get("/kayitlar", params={"plaka": "34 GVN 02"}, headers=guvenlik_header)
    assert r2.status_code == 200, r2.text
    assert r2.json() == []

    # Aynı kayıt, yönetici için (filtresiz) normal şekilde görünmeye devam etmeli.
    r3 = client.get("/kayitlar", params={"plaka": "34 GVN 02"}, headers=yetkili_header)
    assert len(r3.json()) == 1


def test_guvenlik_hic_oturumu_olmayan_kullanici_hicbir_kayit_goremez_fail_closed(client, yetkili_header):
    """Normal akışta bir güvenlik kullanıcısının en az bir vardiya oturumu
    HER ZAMAN vardır (giriş yaptığı an otomatik açılır) -- ama savunma
    amaçlı "sıfır pencere -> sıfır kayıt" varsayılanının (bkz.
    _guvenlik_kayit_filtresi_uygula) hâlâ doğru çalıştığını, oturum
    kaydının veritabanından (ör. elle bir bakım işlemiyle) silinmiş olması
    durumunu simüle ederek doğrular."""
    from backend.database import SessionLocal
    from backend import models as _models

    guvenlik_id, guvenlik_header = _rbac_guvenlik_kullanici_olustur(client, yetkili_header)
    db = SessionLocal()
    try:
        db.query(_models.VardiyaOturumu).filter(_models.VardiyaOturumu.kullanici_id == guvenlik_id).delete()
        db.commit()
    finally:
        db.close()

    r1 = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 GVN 02B", "kamera_id": "GIRIS-KAM", "yon": "giris", "guven_skoru": 0.99,
    })
    assert r1.status_code == 200, r1.text
    r2 = client.get("/kayitlar", params={"plaka": "34 GVN 02B"}, headers=guvenlik_header)
    assert r2.status_code == 200 and r2.json() == [], r2.text


def test_guvenlik_giriste_acilan_oturum_su_anki_kaydi_gorur_ve_istatistiklere_yansir(client, yetkili_header):
    """Giriş yapan bir güvenlik kullanıcısının otomatik açılan oturumu,
    giriş ANINDAN itibaren gerçekleşen kayıtları hemen görünür kılmalı --
    ayrıca panel istatistiklerinin kayıt-bazlı alanlarına da yansımalı
    (bkz. istatistikler()'deki not)."""
    _, guvenlik_header = _rbac_guvenlik_kullanici_olustur(client, yetkili_header)

    r1 = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 GVN 03", "kamera_id": "GIRIS-KAM", "yon": "giris", "guven_skoru": 0.99,
    })
    assert r1.status_code == 200, r1.text

    r2 = client.get("/kayitlar", params={"plaka": "34 GVN 03"}, headers=guvenlik_header)
    assert r2.status_code == 200 and len(r2.json()) == 1, r2.text

    r3 = client.get("/kayitlar/istatistik", headers=guvenlik_header)
    assert r3.status_code == 200, r3.text
    assert r3.json()["bugunku_kayit"] >= 1


def test_cikis_yapinca_vardiya_kapanir_ve_sonraki_kayitlar_gorunmez(client, yetkili_header):
    """`/auth/cikis` çağrısı (bkz. frontend/app.js::oturumKapat), o anki
    açık VardiyaOturumu'nu kapatmalı -- kapandıktan SONRA gerçekleşen bir
    kayıt artık bu kullanıcıya görünmemeli (vardiyası bitti). Token'ın
    kendisi sunucu tarafında iptal edilmediği için (bkz. cikis_yap'ın
    docstring'i) aynı header ile istek yapmaya devam edilebiliyor olması
    kasıtlı -- yalnızca vardiya penceresi kapandığı için kayıt görünmüyor."""
    guvenlik_id, guvenlik_header = _rbac_guvenlik_kullanici_olustur(client, yetkili_header)

    r0 = client.post("/auth/cikis", headers=guvenlik_header)
    assert r0.status_code == 200, r0.text

    r1 = client.get("/vardiya-oturumlari", params={"kullanici_id": guvenlik_id}, headers=yetkili_header)
    assert r1.json()[0]["cikis_zamani"] is not None, "çıkış sonrası oturum kapanmış olmalı"

    r2 = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 GVN 03B", "kamera_id": "GIRIS-KAM", "yon": "giris", "guven_skoru": 0.99,
    })
    assert r2.status_code == 200, r2.text

    r3 = client.get("/kayitlar", params={"plaka": "34 GVN 03B"}, headers=guvenlik_header)
    assert r3.status_code == 200 and r3.json() == [], r3.text


def test_tekrar_giris_unutulmus_acik_oturumu_kapatip_yenisini_acar(client, yetkili_header):
    """Personel çıkış yapmadan (tarayıcıyı/bilgisayarı kapatıp) tekrar giriş
    yaparsa, ÖNCEKİ açık oturum bu anda kapatılmalı ve YENİ bir oturum
    açılmalı -- böylece unutulan bir çıkış, bir SONRAKİ vardiyanın
    kayıtlarına asla karışmaz (bkz. main.py::_guvenlik_oturum_baslat)."""
    import uuid
    kullanici_adi = f"rbac-guvenlik-{uuid.uuid4().hex[:10]}"
    r = client.post("/kullanicilar", json={
        "kullanici_adi": kullanici_adi, "parola": "GucluParola123!", "rol": "güvenlik",
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text
    kullanici_id = r.json()["id"]

    r1 = client.post("/auth/giris", json={"kullanici_adi": kullanici_adi, "parola": "GucluParola123!"})
    assert r1.status_code == 200, r1.text

    # ÇIKIŞ YAPMADAN tekrar giriş yapıyor (unutulmuş oturum senaryosu).
    r2 = client.post("/auth/giris", json={"kullanici_adi": kullanici_adi, "parola": "GucluParola123!"})
    assert r2.status_code == 200, r2.text

    r3 = client.get("/vardiya-oturumlari", params={"kullanici_id": kullanici_id}, headers=yetkili_header)
    assert r3.status_code == 200, r3.text
    oturumlar = sorted(r3.json(), key=lambda o: o["id"])
    assert len(oturumlar) == 2, "iki giriş, İKİ ayrı oturum satırı üretmeli"
    assert oturumlar[0]["cikis_zamani"] is not None, "birinci (unutulmuş) oturum ikinci girişte otomatik kapanmalı"
    assert oturumlar[1]["cikis_zamani"] is None, "ikinci (güncel) oturum hâlâ açık olmalı"


def test_gece_yarisini_gecen_vardiya_oturumu_dogru_filtrelenir(client, yetkili_header):
    """Öz-hizmet oturumları HH:MM string'ler yerine gerçek datetime'lar
    kullandığı için (bkz. models.VardiyaOturumu) gece yarısını geçen bir
    vardiya artık ÖZEL bir hesaplama gerektirmiyor -- oturum satırı ne kadar
    sürerse pencere de o kadar sürer. Bu test, otomatik açılan oturumu
    (giriş dün 23:00, çıkış bugün 07:00 gibi) DOĞRUDAN veritabanında
    güncelleyip bunun doğru filtrelendiğini kanıtlar (mutlak takvim
    tarihleriyle, testin çalıştığı GERÇEK saatten bağımsız)."""
    from backend.database import engine
    from sqlalchemy import text as sqltext

    guvenlik_id, guvenlik_header = _rbac_guvenlik_kullanici_olustur(client, yetkili_header)
    dun = datetime.now().date() - timedelta(days=1)
    bugun = datetime.now().date()
    giris_ts = f"{dun.isoformat()} 23:00:00"
    cikis_ts = f"{bugun.isoformat()} 07:00:00"
    icerde_ts = f"{bugun.isoformat()} 00:30:00"
    disarda_ts = f"{bugun.isoformat()} 12:00:00"

    with engine.connect() as conn:
        # Giriş sırasında otomatik açılan oturumu, testin ihtiyacı olan bilinen
        # bir giriş/çıkışla GÜNCELLİYORUZ (silmek yerine -- ID sabit kalsın diye).
        conn.execute(sqltext(
            "UPDATE vardiya_oturumlari SET giris_zamani=:g, cikis_zamani=:c WHERE kullanici_id=:kid"
        ), {"g": giris_ts, "c": cikis_ts, "kid": guvenlik_id})
        for ts in (icerde_ts, disarda_ts):
            conn.execute(sqltext(
                "INSERT INTO plaka_kayitlari (plaka_no, kamera_id, yon, yetki_durumu, tarih_saat, manuel_giris) "
                "VALUES ('34 GVN 05', 'GIRIS-KAM', 'giris', 'yetkisiz', :ts, 0)"
            ), {"ts": ts})
        conn.commit()

    r1 = client.get("/kayitlar", params={"plaka": "34 GVN 05"}, headers=guvenlik_header)
    assert r1.status_code == 200, r1.text
    kayitlar = r1.json()
    assert len(kayitlar) == 1, "gece yarısını geçen vardiya penceresi yanlış filtrelendi"
    assert "00:30" in kayitlar[0]["tarih_saat"], (
        f"beklenen 'içerde' kayıt (bugün 00:30) yerine farklı bir kayıt döndü: {kayitlar[0]['tarih_saat']}"
    )


def test_ayni_gun_icinde_iki_ayri_oturum_bolunmus_vardiya_dogru_filtrelenir(client, yetkili_header):
    """2026-09-20 kullanıcı sorusu: "Eser sürekli 15:00-23:00'de değil, başka
    vardiyalar da olabiliyor" -- bunun bir uzantısı olarak, aynı personelin
    AYNI GÜN İÇİNDE İKİ AYRI (giriş->çıkış->giriş->çıkış) vardiya oturumu
    olması ("bölünmüş vardiya") senaryosu doğrulanıyor. Öz-hizmet sistemde
    bu, personelin gün içinde iki kez giriş/çıkış yapmasıyla DOĞAL olarak
    elde edilir -- `VardiyaOturumu`'nda (kullanici_id, tarih) üzerinde bir
    TEKİLLİK KISITLAMASI yoktur ve filtre TÜM oturum satırlarını OR ile
    birleştirir (bkz. _guvenlik_kayit_filtresi_uygula). Testin çalıştığı
    GERÇEK saatten bağımsız kalması için oturumlar dünün MUTLAK takvim
    tarihiyle doğrudan veritabanına yazılıyor (gerçek login akışını
    beklemek yerine)."""
    from backend.database import engine
    from sqlalchemy import text as sqltext

    guvenlik_id, guvenlik_header = _rbac_guvenlik_kullanici_olustur(client, yetkili_header)
    dun = datetime.now().date() - timedelta(days=1)
    pencere1_ici_ts = f"{dun.isoformat()} 09:00:00"   # 07:00-11:00 penceresi içinde
    bosluk_ts = f"{dun.isoformat()} 15:00:00"          # iki pencere ARASINDAKİ boşlukta
    pencere2_ici_ts = f"{dun.isoformat()} 21:00:00"   # 19:00-23:00 penceresi içinde

    with engine.connect() as conn:
        for ts in (pencere1_ici_ts, bosluk_ts, pencere2_ici_ts):
            conn.execute(sqltext(
                "INSERT INTO plaka_kayitlari (plaka_no, kamera_id, yon, yetki_durumu, tarih_saat, manuel_giris) "
                "VALUES ('34 GVN 08', 'GIRIS-KAM', 'giris', 'yetkisiz', :ts, 0)"
            ), {"ts": ts})
        # Giriş sırasında otomatik açılan oturumu BİRİNCİ pencereye taşı...
        conn.execute(sqltext(
            "UPDATE vardiya_oturumlari SET giris_zamani=:g, cikis_zamani=:c WHERE kullanici_id=:kid"
        ), {"g": f"{dun.isoformat()} 07:00:00", "c": f"{dun.isoformat()} 11:00:00", "kid": guvenlik_id})
        # ...ve İKİNCİ pencereyi ayrı bir oturum satırı olarak ekle (bölünmüş vardiya).
        conn.execute(sqltext(
            "INSERT INTO vardiya_oturumlari (kullanici_id, giris_zamani, cikis_zamani) VALUES (:kid, :g, :c)"
        ), {"kid": guvenlik_id, "g": f"{dun.isoformat()} 19:00:00", "c": f"{dun.isoformat()} 23:00:00"})
        conn.commit()

    r1 = client.get("/kayitlar", params={"plaka": "34 GVN 08"}, headers=guvenlik_header)
    assert r1.status_code == 200, r1.text
    saatler = sorted(k["tarih_saat"] for k in r1.json())
    assert len(saatler) == 2, (
        f"bölünmüş vardiyanın İKİ oturumundaki kayıtlar da görünmeli, boşluktaki görünmemeli: {saatler}"
    )
    assert any("09:00" in s for s in saatler), "1. oturum (07:00-11:00) içindeki kayıt kayıp"
    assert any("21:00" in s for s in saatler), "2. oturum (19:00-23:00) içindeki kayıt kayıp"
    assert not any("15:00" in s for s in saatler), "iki oturum arasındaki boşluktaki kayıt YANLIŞLIKLA görünüyor"


def test_cakisan_acik_oturumlarda_ayni_kayit_iki_guvenlik_kullanicisinda_da_gorunur(client, yetkili_header):
    """Kullanıcı talebi: aynı anda birden fazla güvenlik personelinin
    vardiyası (öz-hizmet sistemde: açık oturumu) çakışıyorsa, bir kayıt
    DIŞLAYICI bir şekilde tek bir kullanıcıya atanmaz -- ikisinin de
    listesinde ayrı ayrı görünür. İki güvenlik kullanıcısının GİRİŞ yapması
    zaten örtüşen açık oturumlar üretir, ayrıca bir işlem GEREKMEZ."""
    _, guvenlik1_header = _rbac_guvenlik_kullanici_olustur(client, yetkili_header)
    _, guvenlik2_header = _rbac_guvenlik_kullanici_olustur(client, yetkili_header)

    r1 = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 GVN 06", "kamera_id": "GIRIS-KAM", "yon": "giris", "guven_skoru": 0.99,
    })
    assert r1.status_code == 200, r1.text

    for header in (guvenlik1_header, guvenlik2_header):
        r2 = client.get("/kayitlar", params={"plaka": "34 GVN 06"}, headers=header)
        assert r2.status_code == 200 and len(r2.json()) == 1, r2.text


def test_disa_aktar_uc_noktalari_guvenlik_kullanicisiyla_calisir(client, yetkili_header):
    """kayitlari_listele() /disa-aktar/... uçlarından FastAPI DI'ı olmadan
    doğrudan çağrıldığı için (bkz. kayitlari_excel_indir/kayitlari_pdf_indir
    içindeki not), `kullanici` parametresinin açıkça geçirilmemesi bir
    TypeError'a (ve dolayısıyla 500'e) yol açardı -- bu test asıl olarak bu
    kablolamanın doğru yapıldığını doğrular."""
    _, guvenlik_header = _rbac_guvenlik_kullanici_olustur(client, yetkili_header)
    assert client.get("/disa-aktar/excel/kayitlar", headers=guvenlik_header).status_code == 200
    assert client.get("/disa-aktar/pdf/kayitlar", headers=guvenlik_header).status_code == 200


def test_disa_aktar_pdf_kayit_detay_vardiya_disinda_403_doner(client, yetkili_header):
    r1 = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 GVN 07", "kamera_id": "GIRIS-KAM", "yon": "giris", "guven_skoru": 0.99,
    })
    assert r1.status_code == 200, r1.text
    r2 = client.get("/kayitlar", params={"plaka": "34 GVN 07"}, headers=yetkili_header)
    kayit_id = r2.json()[0]["id"]

    # Güvenlik kullanıcısı KAYITTAN SONRA giriş yapıyor -> oturumu bu kayıttan
    # SONRA başlıyor, kayıt onun vardiyasına ait DEĞİL.
    _, guvenlik_header = _rbac_guvenlik_kullanici_olustur(client, yetkili_header)

    r3 = client.get(f"/disa-aktar/pdf/kayit/{kayit_id}", headers=guvenlik_header)
    assert r3.status_code == 403, r3.text


def test_vardiya_oturumu_durumum_guvenlik_disi_rol_icin_zararsiz_yanit_doner(client, yetkili_header):
    """/vardiya-oturumlari/durumum rol kontrolü YAPMAZ (bkz. uç noktanın
    docstring'i) -- yönetici/operatör gibi güvenlik-dışı roller için de 200
    döner, ama yalnızca rol_guvenlik_mi: false ve sunucu saatini içerir."""
    r = client.get("/vardiya-oturumlari/durumum", headers=yetkili_header)
    assert r.status_code == 200, r.text
    gövde = r.json()
    assert gövde["rol_guvenlik_mi"] is False
    assert "sunucu_simdiki_zaman" in gövde
    assert "oturumlar" not in gövde


def test_vardiya_oturumu_durumum_yeni_girisin_acik_oturumunu_dogru_bildirir(client, yetkili_header):
    """Yeni giriş yapmış bir güvenlik kullanıcısı için `su_an_aktif_vardiya_
    var_mi` TRUE ve tek oturumun `devam_ediyor` alanı da TRUE olmalı."""
    _, guvenlik_header = _rbac_guvenlik_kullanici_olustur(client, yetkili_header)
    r = client.get("/vardiya-oturumlari/durumum", headers=guvenlik_header)
    assert r.status_code == 200, r.text
    gövde = r.json()
    assert gövde["rol_guvenlik_mi"] is True
    assert gövde["su_an_aktif_vardiya_var_mi"] is True
    assert gövde["toplam_oturum_sayisi"] == 1
    assert len(gövde["oturumlar"]) == 1
    assert gövde["oturumlar"][0]["devam_ediyor"] is True
    assert gövde["oturumlar"][0]["cikis_zamani"] is None


def test_vardiya_oturumu_durumum_cikis_sonrasi_aktif_degil(client, yetkili_header):
    _, guvenlik_header = _rbac_guvenlik_kullanici_olustur(client, yetkili_header)
    r0 = client.post("/auth/cikis", headers=guvenlik_header)
    assert r0.status_code == 200, r0.text
    r = client.get("/vardiya-oturumlari/durumum", headers=guvenlik_header)
    assert r.status_code == 200, r.text
    gövde = r.json()
    assert gövde["su_an_aktif_vardiya_var_mi"] is False
    assert gövde["oturumlar"][0]["devam_ediyor"] is False
    assert gövde["oturumlar"][0]["cikis_zamani"] is not None


def test_rapor_vardiya_sutunu_acik_oturumdaki_kaydi_dogru_etiketler(client, yetkili_header):
    """2026-09-20 kullanıcı talebi: "her vardiya için kendi geçiş raporları
    olsun ... raporda bir sütun tanımlansın" -- dışa aktarma (Excel/PDF)
    satırlarına main.py::_kayitlari_rapor_satirlari tarafından eklenen
    "vardiya" alanının, kaydın gerçekleştiği anda AÇIK olan güvenlik
    oturumunu (kullanıcı adı + giriş saati) doğru etiketlediğini,
    oturumdan ÖNCEKİ bir kaydı ise BOŞ bıraktığını doğrular."""
    from backend.database import SessionLocal
    from backend import models as _models

    r0 = client.get("/auth/me", headers=yetkili_header)
    assert r0.status_code == 200, r0.text

    r_eski = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 GVN 09A", "kamera_id": "GIRIS-KAM", "yon": "giris", "guven_skoru": 0.99,
    })
    assert r_eski.status_code == 200, r_eski.text

    guvenlik_id, guvenlik_header = _rbac_guvenlik_kullanici_olustur(client, yetkili_header)
    r1 = client.get("/auth/me", headers=guvenlik_header)
    assert r1.status_code == 200, r1.text
    guvenlik_adi = r1.json()["kullanici_adi"]

    r_yeni = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 GVN 09B", "kamera_id": "GIRIS-KAM", "yon": "giris", "guven_skoru": 0.99,
    })
    assert r_yeni.status_code == 200, r_yeni.text

    db = SessionLocal()
    try:
        eski_kayit = db.query(_models.Kayit).filter(_models.Kayit.plaka_no == "34 GVN 09A").first()
        yeni_kayit = db.query(_models.Kayit).filter(_models.Kayit.plaka_no == "34 GVN 09B").first()
        satirlar = pts_main._kayitlari_rapor_satirlari([eski_kayit, yeni_kayit], db)
    finally:
        db.close()

    satir_by_plaka = {s["plaka_no"]: s for s in satirlar}
    assert satir_by_plaka["34 GVN 09A"]["vardiya"] == "", (
        "güvenlik kullanıcısının oturumu açılmadan ÖNCEKİ kayıt hiçbir vardiyaya etiketlenmemeli"
    )
    assert guvenlik_adi in satir_by_plaka["34 GVN 09B"]["vardiya"], (
        f"oturum AÇIKKEN oluşan kayıt, o oturumun sahibiyle etiketlenmeli: {satir_by_plaka['34 GVN 09B']['vardiya']}"
    )


# ================================================================
# "VARDİYA GRUPLARI" (2026-09-21) -- kullanıcı talebi: "kayıtlar ekranına
# yeni bir sütun ekleyebilir miyiz. A B C D Vardiyaları olacak şekilde ...
# LOJMAN A Vardiyası Bülent ile aynı zaman aralığında çalışacağı için ...
# Ana nizamiyeden bülent kontrol ettiğinde Lojman A geçişlerini de
# görebilecek ... A B C D vardiyaları 8 saat bazlı çalışmakta ... giriş
# saatinden 8 saat sonra otomatik çıkış yapılsın." Aşağıdaki testler:
#  1) kullanici_ekle/guncelle'nin vardiya_adi'nı doğru normalize/atayıp
#     kaldırdığını,
#  2) AYNI vardiya adını paylaşan (farklı hesap/nokta) iki güvenlik
#     kullanıcısının birbirinin kayıtlarını görebildiğini,
#  3) Kayıtlar ekranındaki "Vardiya" (A/B/C/D) filtresini,
#  4) Plaka Analizi'ndeki "vardiya_adi" etiketleme alanını,
#  5) 8 saati aşan açık oturumların otomatik kapandığını
# doğrudan doğrular.
# ================================================================

def _rbac_guvenlik_kullanici_olustur_vardiyali(client, yetkili_header, vardiya_adi: str) -> tuple:
    """`_rbac_guvenlik_kullanici_olustur` ile AYNI ama oluşturulan hesaba
    ADLANDIRILMIŞ bir vardiya grubu (bkz. models.Kullanici.vardiya_adi)
    atar."""
    import uuid
    kullanici_adi = f"rbac-guvenlik-vardiyali-{uuid.uuid4().hex[:10]}"
    r = client.post("/kullanicilar", json={
        "kullanici_adi": kullanici_adi, "parola": "GucluParola123!", "rol": "güvenlik",
        "vardiya_adi": vardiya_adi,
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text
    assert r.json()["vardiya_adi"] == vardiya_adi
    kullanici_id = r.json()["id"]
    r2 = client.post("/auth/giris", json={"kullanici_adi": kullanici_adi, "parola": "GucluParola123!"})
    assert r2.status_code == 200, r2.text
    return kullanici_id, {"Authorization": f"Bearer {r2.json()['token']}"}


def test_kullanici_ekle_vardiya_adi_normalize_edilir(client, yetkili_header):
    """Baş/son boşluk temizlenmeli, büyük harfe çevrilmeli (bkz.
    main.py::_vardiya_adi_normalize) -- arayüz A/B/C/D önerir ama alan
    serbest metindir."""
    import uuid
    kullanici_adi = f"vardiya-normalize-{uuid.uuid4().hex[:10]}"
    r = client.post("/kullanicilar", json={
        "kullanici_adi": kullanici_adi, "parola": "GucluParola123!", "rol": "güvenlik",
        "vardiya_adi": "  a  ",
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text
    assert r.json()["vardiya_adi"] == "A"


def test_kullanici_ekle_vardiya_adi_gonderilmezse_none_doner(client, yetkili_header):
    """Varsayılan: vardiya_adi hiç gönderilmezse hesap bağımsız kalır (eski/
    varsayılan davranış -- yalnızca KENDİ oturumlarını görür)."""
    import uuid
    kullanici_adi = f"vardiya-yok-{uuid.uuid4().hex[:10]}"
    r = client.post("/kullanicilar", json={
        "kullanici_adi": kullanici_adi, "parola": "GucluParola123!", "rol": "güvenlik",
    }, headers=yetkili_header)
    assert r.status_code == 200, r.text
    assert r.json()["vardiya_adi"] is None


def test_kullanici_guncelle_vardiya_adi_atanir_ve_kaldirilir(client, yetkili_header):
    """`kullanici_guncelle`: None = değiştirme; boş string (normalize
    sonrası) = kaldır; dolu string = ata (bkz. schemas.KullaniciGuncelle.
    vardiya_adi docstring'i)."""
    guvenlik_id, _ = _rbac_guvenlik_kullanici_olustur(client, yetkili_header)

    # Boş PUT (vardiya_adi göndermeden) -- None kalmalı.
    r0 = client.put(f"/kullanicilar/{guvenlik_id}", json={}, headers=yetkili_header)
    assert r0.status_code == 200, r0.text
    assert r0.json()["vardiya_adi"] is None

    r1 = client.put(f"/kullanicilar/{guvenlik_id}", json={"vardiya_adi": " b "}, headers=yetkili_header)
    assert r1.status_code == 200, r1.text
    assert r1.json()["vardiya_adi"] == "B"

    r2 = client.put(f"/kullanicilar/{guvenlik_id}", json={"vardiya_adi": "   "}, headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    assert r2.json()["vardiya_adi"] is None, "boş string (normalize sonrası) ataması KALDIRMALI"


def test_vardiya_grubu_paylasimli_gorunurluk_farkli_hesaplar_arasinda(client, yetkili_header):
    """Ana özellik: AYNI vardiya adına ("A") atanmış, birbirinden BAĞIMSIZ
    iki güvenlik hesabı (ör. Ana Nizamiye + Lojman Nizamiye) birbirinin
    vardiya penceresindeki kayıtları görebilmeli -- kullanıcı talebi:
    "Ana nizamiyeden bülent kontrol ettiğinde Lojman A geçişlerini de
    görebilecek." Farklı ("B") bir gruba atanmış üçüncü bir hesap ise
    GÖREMEMELİ."""
    _, ana_a_header = _rbac_guvenlik_kullanici_olustur_vardiyali(client, yetkili_header, "A")
    _, lojman_a_header = _rbac_guvenlik_kullanici_olustur_vardiyali(client, yetkili_header, "A")
    _, b_header = _rbac_guvenlik_kullanici_olustur_vardiyali(client, yetkili_header, "B")

    # "lojman-A" hesabının oturumu AÇIKKEN oluşan bir kayıt.
    r = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 VG 01", "kamera_id": "LOJMAN-KAM", "yon": "giris", "guven_skoru": 0.99,
    })
    assert r.status_code == 200, r.text

    # AYNI ("A") gruptaki diğer hesap (farklı fiziksel nokta) bu kaydı görmeli.
    r_ana = client.get("/kayitlar", params={"plaka": "34 VG 01"}, headers=ana_a_header)
    assert r_ana.status_code == 200, r_ana.text
    assert len(r_ana.json()) == 1, "aynı vardiya grubundaki BAŞKA bir hesap bu kaydı görebilmeli"

    # FARKLI ("B") gruptaki hesap GÖREMEMELİ.
    r_b = client.get("/kayitlar", params={"plaka": "34 VG 01"}, headers=b_header)
    assert r_b.status_code == 200, r_b.text
    assert r_b.json() == [], "farklı vardiya grubundaki bir hesap bu kaydı GÖRMEMELİ"


def test_kayitlar_vardiya_adi_filtresi_dogru_kayitlari_getirir(client, yetkili_header):
    """Kayıtlar ekranındaki "Vardiya" filtresi (rol ne olursa olsun
    çağrılabilir, bkz. main.py::_vardiya_adi_filtresi_uygula) -- "Tüm
    Güvenlik Personeli kayıtlar ekranından A B C D Vardiyalarında geçen
    araçları filtreleyip ... arayabilsin" talebini doğrular."""
    _, a_header = _rbac_guvenlik_kullanici_olustur_vardiyali(client, yetkili_header, "A")

    r1 = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 VG 02", "kamera_id": "KAM-1", "yon": "giris", "guven_skoru": 0.99,
    })
    assert r1.status_code == 200, r1.text

    # Yönetici (filtresiz görünürlük) "A" filtresini uygularsa yalnızca bu kaydı görmeli.
    r_a = client.get("/kayitlar", params={"plaka": "34 VG 02", "vardiya_adi": "a"}, headers=yetkili_header)
    assert r_a.status_code == 200, r_a.text
    assert len(r_a.json()) == 1, "küçük harfle gönderilse bile normalize edilip eşleşmeli"
    assert r_a.json()[0]["vardiya_adi"] == "A"

    # Hiç kullanılmayan bir vardiya adı -- GÜVENLİ TARAF: boş liste döner.
    r_c = client.get("/kayitlar", params={"plaka": "34 VG 02", "vardiya_adi": "C"}, headers=yetkili_header)
    assert r_c.status_code == 200, r_c.text
    assert r_c.json() == []

    # Filtre uygulanmadan (vardiya_adi yok) da kayıt görünmeli VE "Vardiya" sütunu dolu olmalı.
    r_tumu = client.get("/kayitlar", params={"plaka": "34 VG 02"}, headers=yetkili_header)
    assert r_tumu.status_code == 200, r_tumu.text
    assert r_tumu.json()[0]["vardiya_adi"] == "A"
    del a_header  # yalnızca hesabı oluşturmak için gerekliydi


def test_kayitlar_vardiya_adi_filtresi_guvenlik_hesabi_kendi_disindaki_vardiyayi_da_bulur(client, yetkili_header):
    """GERÇEK ÜRETİMDE BULUNAN HATA (2026-09-22, ekran görüntüleriyle
    bildirildi): yukarıdaki test (test_kayitlar_vardiya_adi_filtresi_dogru_
    kayitlari_getirir) bu filtreyi YALNIZCA kısıtlamasız bir `yetkili_header`
    (yönetici) ile test ediyordu -- bu rol için _guvenlik_kayit_filtresi_
    uygula zaten no-op olduğundan, testin kendisi asıl üretim hatasını hiç
    yakalayamıyordu. Gerçek hata SADECE bir GÜVENLİK hesabı KENDİ vardiyası
    DIŞINDA bir vardiyayı bu filtreyle ararsa ortaya çıkıyordu: kendi vardiya
    penceresi kısıtlaması ile "Vardiya" filtresi ANDlanıyor, iki farklı
    vardiyanın pencereleri çakışmadığı için sonuç HER ZAMAN boş dönüyordu --
    kullanıcı raporu tam olarak buydu ("A" hesabı "D" filtresini seçince
    "Kayıt bulunamadı", plaka yazsa bile).

    Bu test doğrudan o senaryoyu kurar: "A" vardiyasındaki bir güvenlik
    hesabı, "D" vardiyasının oturumu SIRASINDA oluşmuş bir kaydı, hem
    plakasız HEM DE plaka ile (kullanıcının bildirdiği ikinci belirti) "D"
    filtresiyle bulabilmeli; kendi ("A") vardiyasını filtrelerken ise o
    kaydı GÖRMEMELİ (çapraz filtre sınırsız bir görünürlük açığına
    dönüşmemeli, yalnızca SEÇİLEN vardiyayı göstermeli)."""
    _, a_header = _rbac_guvenlik_kullanici_olustur_vardiyali(client, yetkili_header, "A")
    _rbac_guvenlik_kullanici_olustur_vardiyali(client, yetkili_header, "D")  # "D" oturumunu açar

    r1 = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 XVD 01", "kamera_id": "KAM-1", "yon": "giris", "guven_skoru": 0.99,
    })
    assert r1.status_code == 200, r1.text

    # "A" hesabı, "D" filtresiyle (plakasız) bu kaydı bulabilmeli.
    r_d_plakasiz = client.get("/kayitlar", params={"vardiya_adi": "D"}, headers=a_header)
    assert r_d_plakasiz.status_code == 200, r_d_plakasiz.text
    assert any(k["plaka_no"] == "34 XVD 01" for k in r_d_plakasiz.json()), (
        "güvenlik hesabı KENDİ vardiyası dışındaki bir vardiyayı filtreleyince "
        "boş dönmemeli -- iki filtre birbirini ANDlamamalı"
    )

    # AYNI şey plaka filtresiyle BİRLİKTE de çalışmalı (kullanıcının ikinci belirtisi).
    r_d_plakali = client.get(
        "/kayitlar", params={"plaka": "34 XVD 01", "vardiya_adi": "D"}, headers=a_header,
    )
    assert r_d_plakali.status_code == 200, r_d_plakali.text
    assert len(r_d_plakali.json()) == 1
    assert r_d_plakali.json()[0]["vardiya_adi"] == "D"

    # Sayfalama sayacı (/kayitlar/sayfa-bilgisi) da AYNI şekilde tutarlı olmalı.
    r_sayfa = client.get(
        "/kayitlar/sayfa-bilgisi", params={"plaka": "34 XVD 01", "vardiya_adi": "D"}, headers=a_header,
    )
    assert r_sayfa.status_code == 200, r_sayfa.text
    assert r_sayfa.json()["toplam"] == 1

    # Sınırsız bir görünürlük açığı DEĞİL: kendi ("A") vardiyasıyla filtrelerken bu kayıt GÖRÜNMEMELİ.
    r_a = client.get(
        "/kayitlar", params={"plaka": "34 XVD 01", "vardiya_adi": "A"}, headers=a_header,
    )
    assert r_a.status_code == 200, r_a.text
    assert r_a.json() == []


def test_disa_aktar_pdf_kayit_detay_vardiya_adi_filtresiyle_baska_vardiyadan_izin_verir(client, yetkili_header):
    """İkinci, ilişkili GERÇEK ÜRETİM hatası: yukarıdaki listede "D" filtresiyle
    görünür hale gelen bir kaydı, kullanıcı satırdaki "PDF indir" düğmesiyle
    indirmeye çalışınca (bkz. app.js::kayitPdfIndir, artık aynı filtre
    değerini bu uca da iletiyor) HÂLÂ 403 alıyordu -- çünkü tekil kayıt
    görünürlüğü (_guvenlik_kayit_gorunur_mu) `vardiya_adi` parametresini hiç
    bilmiyor, yalnızca kullanıcının KENDİ vardiya penceresine bakıyordu.
    `test_disa_aktar_pdf_kayit_detay_vardiya_disinda_403_doner` bu ucun
    varsayılan (parametre olmadan) davranışının BOZULMADIĞINI zaten
    doğruluyor; bu test yeni `vardiya_adi` parametresinin gerçekten izin
    verdiğini doğrular."""
    _, a_header = _rbac_guvenlik_kullanici_olustur_vardiyali(client, yetkili_header, "A")
    _rbac_guvenlik_kullanici_olustur_vardiyali(client, yetkili_header, "D")

    r1 = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 XVD 02", "kamera_id": "KAM-1", "yon": "giris", "guven_skoru": 0.99,
    })
    assert r1.status_code == 200, r1.text
    r2 = client.get("/kayitlar", params={"plaka": "34 XVD 02", "vardiya_adi": "D"}, headers=a_header)
    kayit_id = r2.json()[0]["id"]

    # `vardiya_adi` OLMADAN hâlâ 403 (varsayılan/eski davranış korunuyor).
    r3 = client.get(f"/disa-aktar/pdf/kayit/{kayit_id}", headers=a_header)
    assert r3.status_code == 403, r3.text

    # `vardiya_adi=D` İLE artık izin verilmeli (listede zaten görünüyordu).
    r4 = client.get(f"/disa-aktar/pdf/kayit/{kayit_id}", params={"vardiya_adi": "D"}, headers=a_header)
    assert r4.status_code == 200, r4.text


def test_plaka_analizinde_vardiya_adi_alani_dolar(client, yetkili_header):
    """"plaka arayınca karşısına kimin vardiyasında girip çıktığı
    gözükebilsin" talebi -- bkz. main.py::plaka_analiz, son_kayitlar[].
    vardiya_adi."""
    _rbac_guvenlik_kullanici_olustur_vardiyali(client, yetkili_header, "D")
    r1 = client.post("/kayitlar/otomatik", data={
        "plaka_no": "34 VG 03", "kamera_id": "KAM-1", "yon": "giris", "guven_skoru": 0.99,
    })
    assert r1.status_code == 200, r1.text
    r2 = client.get("/kayitlar/analiz/34 VG 03", headers=yetkili_header)
    assert r2.status_code == 200, r2.text
    son_kayitlar = r2.json()["son_kayitlar"]
    assert len(son_kayitlar) == 1
    assert son_kayitlar[0]["vardiya_adi"] == "D"


def test_vardiya_otomatik_kapama_8_saati_asan_acik_oturumu_kapatir(client, yetkili_header):
    """"A B C D vardiyaları 8 saat bazlı çalışmakta yani vardiya amiri
    çıkış yapmayı unutsa bile giriş saatinden 8 saat sonra otomatik çıkış
    yapılsın" -- bkz. main.py::_vardiya_otomatik_kapama_calistir. Giriş
    zamanı elle 9 saat geriye alınarak "unutulmuş, süresi geçmiş" bir açık
    oturum simüle edilir; 7 saat önce açılmış bir oturum ise HENÜZ
    kapanmamalı."""
    from backend.database import SessionLocal
    from backend import models as _models

    guvenlik_id, _ = _rbac_guvenlik_kullanici_olustur(client, yetkili_header)
    db = SessionLocal()
    try:
        oturum = db.query(_models.VardiyaOturumu).filter(
            _models.VardiyaOturumu.kullanici_id == guvenlik_id
        ).first()
        assert oturum is not None and oturum.cikis_zamani is None
        gercek_giris = datetime.now() - timedelta(hours=9)
        oturum.giris_zamani = gercek_giris
        db.commit()

        guvenlik_id2, _ = _rbac_guvenlik_kullanici_olustur(client, yetkili_header)
        oturum2 = db.query(_models.VardiyaOturumu).filter(
            _models.VardiyaOturumu.kullanici_id == guvenlik_id2
        ).first()
        oturum2.giris_zamani = datetime.now() - timedelta(hours=7)
        db.commit()

        kapatilan = pts_main._vardiya_otomatik_kapama_calistir(db)
        assert kapatilan == 1, "yalnızca 8 saati AŞAN oturum kapatılmalı"

        db.refresh(oturum)
        db.refresh(oturum2)
        assert oturum.cikis_zamani is not None, "9 saat önce açılmış oturum 8 saat sınırını aştığı için kapanmalı"
        assert oturum.cikis_zamani == gercek_giris + timedelta(hours=8), (
            "çıkış zamanı ŞİMDİ değil, giriş+8 saat olarak ayarlanmalı"
        )
        assert oturum2.cikis_zamani is None, "7 saat önce açılmış oturum HENÜZ 8 saati aşmadığı için kapanmamalı"
    finally:
        db.close()
