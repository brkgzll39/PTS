"""backend/excel_export.py için testler.

Bu modül (main.py'nin aksine) fastapi/sqlalchemy'ye bağımlı DEĞİL --
yalnızca openpyxl kullanıyor. Bu yüzden (depodaki diğer birçok testin aksine)
bu dosya GERÇEKTEN çalıştırılıp doğrulanabildi (bkz. README'deki
"Kalıcı Test Altyapısı" bölümündeki, sandbox'ta fastapi/sqlalchemy'nin
kurulu olmadığına dair not -- bu iki modül o kısıtlamanın DIŞINDA kalıyor).
"""
from datetime import datetime

import openpyxl
import pytest

from backend import excel_export


def _ornek_satir(**gecersizler):
    satir = {
        "id": 1,
        "plaka_no": "34 ABC 123",
        "ad": "Ahmet",
        "soyad": "Yılmaz",
        "site": "LOJMAN",
        "blok": "",
        "daire": "PERSONEL",
        "otopark": "",
        "nokta": "LOJMAN GİRİŞ",
        "gecis_tipi": "Giriş",
        "arac_tipi": "GÜVENLİK ŞEFLİĞİ",
        "tarih_saat": datetime(2026, 9, 17, 8, 4, 7),
        "notlar": "test notu",
        "goruntu_yolu": None,
        "vardiya": "",
    }
    satir.update(gecersizler)
    return satir


# ------------------------------------------------------------------
# "GEÇİŞ RAPORU" Excel çıktısı (2026-09-18) -- kullanıcının paylaştığı
# referans ürünün rapor biçimine yaklaştırma kararı.
# ------------------------------------------------------------------

def test_kayitlar_excel_basliklari_referans_rapor_ile_eslesir(tmp_path):
    """Kullanıcının paylaştığı referans "GEÇİŞ RAPORU" Excel çıktısının
    başlık satırı: Plaka, Adı, Soyadı, Site, Blok, Daire, Otopark, Nokta,
    Geçiş Tipi, Araç Tipi, Tarih, Notlar. Bizim raporumuz bunu birebir
    içerir; başa "ID", sona da (2026-09-20) "Vardiya" -- ikisi de kendi
    sistemimize özgü ek sütunlar."""
    dosya = tmp_path / "kayitlar.xlsx"
    excel_export.kayitlar_excel_olustur([_ornek_satir()], str(dosya))

    wb = openpyxl.load_workbook(str(dosya))
    ws = wb.active
    basliklar = [h.value for h in ws[1]]
    assert basliklar == [
        "ID", "Plaka", "Adı", "Soyadı", "Site", "Blok", "Daire", "Otopark",
        "Nokta", "Geçiş Tipi", "Araç Tipi", "Tarih", "Notlar", "Vardiya",
    ]
    assert ws.title == "Geçiş Raporu"


def test_kayitlar_excel_veri_satiri_dogru_yazilir(tmp_path):
    dosya = tmp_path / "kayitlar.xlsx"
    excel_export.kayitlar_excel_olustur(
        [_ornek_satir(vardiya="Eser Akar (15:00-23:10)")], str(dosya)
    )

    wb = openpyxl.load_workbook(str(dosya))
    ws = wb.active
    satir = [h.value for h in ws[2]]
    assert satir == [
        1, "34 ABC 123", "Ahmet", "Yılmaz", "LOJMAN", None, "PERSONEL", None,
        "LOJMAN GİRİŞ", "Giriş", "GÜVENLİK ŞEFLİĞİ", "17.09.2026 08:04:07", "test notu",
        "Eser Akar (15:00-23:10)",
    ]


def test_kayitlar_excel_formul_enjeksiyonu_engellenir(tmp_path):
    """Kamera/OCR ya da kişi kaydından gelen bir metin '=', '+', '-' veya '@'
    ile başlıyorsa Excel'de formül olarak YORUMLANMAMALI (bkz.
    excel_export._guvenli_hucre) -- Vardiya sütunu için de aynı koruma
    geçerli olmalı, çünkü içeriği (kullanıcı adı) da nihayetinde kullanıcı
    girdisinden türetiliyor."""
    dosya = tmp_path / "kayitlar.xlsx"
    excel_export.kayitlar_excel_olustur(
        [_ornek_satir(ad="=CMD('zararli')", notlar="+1+1", vardiya="=ALSO_BAD()")], str(dosya)
    )
    wb = openpyxl.load_workbook(str(dosya))
    ws = wb.active
    satir = [h.value for h in ws[2]]
    assert satir[2] == "'=CMD('zararli')"
    assert satir[-2] == "'+1+1"
    assert satir[-1] == "'=ALSO_BAD()"


def test_kayitlar_excel_kontrol_karakteri_icermez_cokmez(tmp_path):
    """KÖK NEDEN (2026-09-21, kullanıcı bildirdi -- toplu "GEÇİŞ RAPORU"
    Excel'i, kimliği doğrulanmış geçerli bir istekte bile genel 500'e
    düşüyordu): OOXML biçimi XML 1.0 gereği belirli KONTROL
    KARAKTERLERİNİ (NUL, backspace vb. -- sekme/satır sonu HARİÇ) hücre
    metninde barındıramaz. `not_metni` gibi serbest metin alanlarından
    birine (kopyala/yapıştır ya da bozuk bir kaynaktan) böyle bir karakter
    karışırsa, önceden openpyxl `wb.save()` sırasında yakalanmamış bir
    `IllegalCharacterError` fırlatırdı. Artık `_guvenli_hucre` bu
    karakterleri hücreye yazmadan önce temizliyor."""
    dosya = tmp_path / "kontrol_karakteri.xlsx"
    tehlikeli = "not: \x00\x01\x02 arıza bildirildi \x0b\x0c\x1f"
    # Çökmemeli:
    excel_export.kayitlar_excel_olustur([_ornek_satir(notlar=tehlikeli)], str(dosya))
    wb = openpyxl.load_workbook(str(dosya))
    ws = wb.active
    notlar_hucresi = ws[2][12].value  # "Notlar" sütunu (0-indeksli 12. sütun)
    assert "\x00" not in notlar_hucresi and "\x01" not in notlar_hucresi
    assert "arıza bildirildi" in notlar_hucresi, "kontrol karakterleri temizlenirken ÇEVRESİNDEKİ düz metin de silinmemeli"


def test_kayitlar_excel_bos_liste_ile_de_calisir(tmp_path):
    """Filtreye uyan hiçbir kayıt yoksa da rapor (yalnızca başlık satırıyla)
    hatasız üretilmeli -- boş bir sonuç, dışa aktarma özelliğini bozmamalı."""
    dosya = tmp_path / "kayitlar.xlsx"
    excel_export.kayitlar_excel_olustur([], str(dosya))
    wb = openpyxl.load_workbook(str(dosya))
    ws = wb.active
    assert ws.max_row == 1


# ------------------------------------------------------------------
# Personel/abone/ziyaretçi toplu içe aktarma ŞABLONU (2026-09-18)
# ------------------------------------------------------------------

def _basliklari_normalize_et(basliklar):
    # main.py::toplu_kisi_import'un YAPTIĞI normalizasyonun birebir aynısı --
    # şablonun bu ayrı, bağımsız kopyayla hâlâ eşleştiğini doğrulamak için.
    return [str(b).strip().lower().replace(" ", "_") if b else "" for b in basliklar]


def test_ice_aktarma_sablonu_basliklari_toplu_import_ile_eslesir(tmp_path):
    """Şablonun ürettiği başlıklar, main.py::toplu_kisi_import'un normalize
    ettikten sonra BEKLEDİĞİ 'ad_soyad', 'plaka_no', 'tip', 'telefon',
    'daire_departman' sütunlarıyla birebir eşleşmeli -- aksi halde kullanıcı
    şablonu hiç değiştirmeden geri yüklese bile içe aktarma başarısız olurdu."""
    dosya = tmp_path / "sablon.xlsx"
    excel_export.kisi_ice_aktarma_sablonu_olustur(str(dosya))

    wb = openpyxl.load_workbook(str(dosya))
    ws = wb["Kişiler"]
    basliklar_ham = [h.value for h in ws[1]]
    basliklar = _basliklari_normalize_et(basliklar_ham)
    for zorunlu in ("ad_soyad", "plaka_no", "tip"):
        assert zorunlu in basliklar, f"'{zorunlu}' şablonda eksik -- içe aktarma her zaman başarısız olurdu"
    assert "telefon" in basliklar
    assert "daire_departman" in basliklar


def test_ice_aktarma_sablonu_ornek_satirlari_gecerli_tip_kullanir(tmp_path):
    """Şablondaki örnek satırların 'tip' değeri, toplu_kisi_import'un kabul
    ettiği GEÇERLİ değerlerden biri olmalı -- kullanıcı örneği hiç
    değiştirmeden yanlışlıkla yüklerse bile sessizce 'abone'ya
    düşürülmemeli (bkz. toplu_kisi_import'un aynı davranışı)."""
    dosya = tmp_path / "sablon.xlsx"
    excel_export.kisi_ice_aktarma_sablonu_olustur(str(dosya))
    wb = openpyxl.load_workbook(str(dosya))
    ws = wb["Kişiler"]
    basliklar = [h.value for h in ws[1]]
    tip_idx = basliklar.index("Tip")
    for satir in ws.iter_rows(min_row=2, values_only=True):
        assert satir[tip_idx] in ("abone", "personel", "ziyaretci")


def test_ice_aktarma_sablonu_coklu_plaka_ornegi_icerir(tmp_path):
    """2026-09-23 kullanıcı isteği: birden fazla aracı olan bir kişi/birim
    TEK satırda virgülle ayrılmış birden fazla plakayla içe aktarılabilmeli
    -- şablonda bunu gösteren, virgül içeren en az bir örnek satır olmalı ki
    kullanıcı bu söz dizimini panelden ayrıca sormadan görebilsin (bkz.
    metin_araclari.py::plaka_hucresini_ayir)."""
    dosya = tmp_path / "sablon.xlsx"
    excel_export.kisi_ice_aktarma_sablonu_olustur(str(dosya))
    wb = openpyxl.load_workbook(str(dosya))
    ws = wb["Kişiler"]
    basliklar = [h.value for h in ws[1]]
    plaka_idx = basliklar.index("Plaka No")
    plaka_degerleri = [satir[plaka_idx] for satir in ws.iter_rows(min_row=2, values_only=True)]
    assert any("," in (deger or "") for deger in plaka_degerleri), (
        "Şablonda çoklu plaka söz dizimini (virgülle ayırma) gösteren bir örnek satır yok"
    )


def test_ice_aktarma_sablonu_aciklama_sayfasi_var(tmp_path):
    dosya = tmp_path / "sablon.xlsx"
    excel_export.kisi_ice_aktarma_sablonu_olustur(str(dosya))
    wb = openpyxl.load_workbook(str(dosya))
    assert "Açıklama" in wb.sheetnames
    aciklama = wb["Açıklama"]
    metin = "\n".join(str(c.value) for row in aciklama.iter_rows() for c in row if c.value)
    assert "ad_soyad" not in metin.lower() or "Ad Soyad" in metin  # başlıklar okunabilir Türkçe biçimde açıklanır
    assert "abone" in metin and "personel" in metin and "ziyaretci" in metin


def test_ice_aktarma_sablonu_tip_sutununda_veri_dogrulamasi_var(tmp_path):
    """'Tip' sütununa açılır liste (data validation) eklenmiş olmalı --
    aksi halde kullanıcı serbest metin yazıp typo yapabilir ve bu, sessizce
    'abone'ya düşer (bkz. toplu_kisi_import)."""
    dosya = tmp_path / "sablon.xlsx"
    excel_export.kisi_ice_aktarma_sablonu_olustur(str(dosya))
    wb = openpyxl.load_workbook(str(dosya))
    ws = wb["Kişiler"]
    assert len(ws.data_validations.dataValidation) >= 1
    dogrulama = ws.data_validations.dataValidation[0]
    assert dogrulama.type == "list"
    for tip in ("abone", "personel", "ziyaretci"):
        assert tip in dogrulama.formula1
